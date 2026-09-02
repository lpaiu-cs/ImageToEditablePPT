from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from image_to_editable_ppt.ml.phase16.artifacts import (
    bound_file_set_digest,
    pptx_semantic_digest,
    pptx_shape_counts,
)


def _deck(path: Path, *, text: str) -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = presentation.slides[0].shapes.add_textbox(
        Inches(1), Inches(1), Inches(2), Inches(1)
    )
    shape.name = "node:test"
    shape.text = text
    presentation.save(path)


def test_pptx_semantic_digest_tracks_editable_content_not_package_bytes(tmp_path: Path) -> None:
    first = tmp_path / "first.pptx"
    second = tmp_path / "second.pptx"
    changed = tmp_path / "changed.pptx"
    _deck(first, text="same")
    _deck(second, text="same")
    _deck(changed, text="different")

    assert pptx_semantic_digest(first) == pptx_semantic_digest(second)
    assert pptx_semantic_digest(first) != pptx_semantic_digest(changed)
    assert pptx_shape_counts(first) == {
        "total": 1,
        "nodes": 1,
        "containers": 0,
        "connector_segments": 0,
        "semantic_connectors": 0,
        "textboxes": 1,
        "objects_with_text": 1,
        "blank": False,
    }


def test_bound_file_set_digest_binds_relative_paths_and_bytes(tmp_path: Path) -> None:
    first = tmp_path / "a.json"
    second = tmp_path / "nested" / "b.json"
    second.parent.mkdir()
    first.write_text(json.dumps({"a": 1}), encoding="utf-8")
    second.write_text(json.dumps({"b": 2}), encoding="utf-8")

    forward = bound_file_set_digest([first, second], root=tmp_path)
    reverse = bound_file_set_digest([second, first], root=tmp_path)
    assert forward == reverse

    second.write_text(json.dumps({"b": 3}), encoding="utf-8")
    assert bound_file_set_digest([first, second], root=tmp_path) != forward

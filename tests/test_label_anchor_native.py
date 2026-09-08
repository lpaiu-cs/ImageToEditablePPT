from dataclasses import replace

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from image_to_editable_ppt.region.emitter import write_native
from image_to_editable_ppt.research.label_anchor_native import write_label_anchor_native
from image_to_editable_ppt.research.preset_emit import preset_vertices
from image_to_editable_ppt.v3.core.enums import NodeKind, PortOwnerKind, ConnectorKind, TextRegionRole
from image_to_editable_ppt.v3.core.types import BBox, ImageSize, Point
from image_to_editable_ppt.v3.emit.models import EmitScene, EmitShapePrimitive, EmitConnectorPrimitive, EmitTextPrimitive
from image_to_editable_ppt.v3.emit.pptx_writer import write_pptx
from image_to_editable_ppt.v3.emit.style import ShapeVisualStyle


def fixture(kind=NodeKind.LABEL_ANCHOR, label='A long label\nSecond line'):
    shapes = tuple(EmitShapePrimitive(key, PortOwnerKind.NODE, kind, BBox(*bounds), .9, label=label)
                   for key, bounds in [('a', (10, 10, 70, 50)), ('b', (150, 110, 220, 160))])
    edge = EmitConnectorPrimitive('c', ConnectorKind.ARROW, .9, 'a', PortOwnerKind.NODE,
        'b', PortOwnerKind.NODE, 'p1', 'p2', path_points=(Point(70, 22), Point(183, 22), Point(183, 110)),
        arrowhead_start=True, arrowhead_end=True)
    return EmitScene(ImageSize(260, 200), shapes=shapes, connectors=(edge,))


def objects(path):
    return {s.name: s for s in Presentation(path).slides[0].shapes}


@pytest.mark.parametrize('reverse', [False, True])
@pytest.mark.parametrize('mixed', [False, True])
def test_real_sites_paths_arrows_and_semantics(tmp_path, reverse, mixed):
    scene = fixture()
    if mixed:
        scene = replace(scene, shapes=(replace(scene.shapes[0], shape_kind=NodeKind.BOX), scene.shapes[1]))
    if reverse:
        c = scene.connectors[0]
        scene = replace(scene, connectors=(replace(c, source_owner_id='b', target_owner_id='a',
                                                   path_points=tuple(reversed(c.path_points))),))
    emitted, audit = write_label_anchor_native(scene, tmp_path/'out.pptx',
        styles={'b': ShapeVisualStyle(fill=(255, 0, 0), stroke=(0, 0, 0))})
    assert emitted == scene
    assert not audit['rejected_connectors']
    items = objects(tmp_path/'out.pptx')
    assert set(items) == {'a', 'b', 'c'}
    assert items['b']._element.spPr.find(qn('a:noFill')) is not None
    assert items['b']._element.spPr.find(qn('a:ln')).find(qn('a:noFill')) is not None
    for tag, owner, point in [('stCxn', scene.connectors[0].source_owner_id, scene.connectors[0].path_points[0]),
                              ('endCxn', scene.connectors[0].target_owner_id, scene.connectors[0].path_points[-1])]:
        binding = items['c']._element.xpath('.//a:'+tag)[0]
        assert int(binding.get('id')) == items[owner].shape_id
        geom = items[owner]._element.spPr.find(qn('a:custGeom'))
        index = int(binding.get('idx'))
        site = geom.find(qn('a:cxnLst'))[index]
        guides = {g.get('name'): g.get('fmla') for g in geom.find(qn('a:gdLst'))}
        pos = site.find(qn('a:pos'))
        shape = next(s for s in scene.shapes if s.id == owner)
        for axis, expected in [('x', (point.x-shape.bbox.x0)/shape.bbox.width),
                               ('y', (point.y-shape.bbox.y0)/shape.bbox.height)]:
            assert int(guides[pos.get(axis)].split()[2])/1e6 == pytest.approx(expected, abs=1e-6)
    assert [x.get('type') for x in items['c']._element.xpath('.//a:headEnd | .//a:tailEnd')] == ['triangle', 'triangle']
    vertices = preset_vertices(items['c'])
    for actual, expected in zip(vertices, audit['fits']['c']['route_px'], strict=True):
        assert actual == pytest.approx((expected[0]*.75, expected[1]*.75), abs=.001)
    assert audit['fits']['c']['error_px'] == pytest.approx(0)
    assert not Presentation(tmp_path/'out.pptx').slides[0]._element.xpath('.//p:pic')


@pytest.mark.parametrize('label', [None, '', '   ', 'Readable\nText'])
@pytest.mark.parametrize('owned', [False, True])
def test_exact_text_format_and_blank_policy(tmp_path, label, owned):
    scene = fixture(label=label)
    if owned:
        text = EmitTextPrimitive('ocr', TextRegionRole.LABEL, BBox(15, 15, 60, 30), .99,
                                 text='Owned text', owner_ids=('a',))
        scene = replace(scene, texts=(text,))
    write_pptx(replace(scene, connectors=()), tmp_path/'reference.pptx')
    emitted, audit = write_label_anchor_native(scene, tmp_path/'out.pptx')
    assert emitted == scene
    actual, reference = objects(tmp_path/'out.pptx'), objects(tmp_path/'reference.pptx')
    for key in ('a', 'b'):
        assert etree.tostring(actual[key]._element.find(qn('p:txBody'))) == etree.tostring(reference[key]._element.find(qn('p:txBody')))
    assert audit['label_anchor_adapter']['visible_label_count'] == sum(bool(reference[k].text.strip()) for k in ('a', 'b'))


@pytest.mark.parametrize('kind', [NodeKind.BOX, NodeKind.ROUNDED_BOX, NodeKind.ELLIPSE])
def test_ordinary_output_unchanged(tmp_path, kind):
    scene = fixture(kind)
    original, original_audit = write_native(scene, tmp_path/'original.pptx')
    emitted, audit = write_label_anchor_native(scene, tmp_path/'out.pptx')
    assert emitted == original
    assert audit['fits'] == original_audit['fits']
    assert audit['rejected_connectors'] == original_audit['rejected_connectors']
    assert Presentation(tmp_path/'original.pptx').slides[0]._element.xml == Presentation(tmp_path/'out.pptx').slides[0]._element.xml


def test_rejection_and_failure_preserve_output(tmp_path):
    scene = fixture()
    scene = replace(scene, connectors=(replace(scene.connectors[0], source_owner_id='missing'),))
    emitted, audit = write_label_anchor_native(scene, tmp_path/'out.pptx')
    assert not emitted.connectors
    assert audit['rejected_connectors'] == {'c': 'unknown endpoint owner'}
    before = (tmp_path/'out.pptx').read_bytes()
    duplicate = replace(scene, shapes=(scene.shapes[0], scene.shapes[0]))
    with pytest.raises(ValueError, match='duplicate shape IDs'):
        write_label_anchor_native(duplicate, tmp_path/'out.pptx')
    assert (tmp_path/'out.pptx').read_bytes() == before
    assert not list(tmp_path.glob('.label-anchor-*'))


def test_overlapping_owners_are_not_merged(tmp_path):
    scene = fixture()
    duplicate = replace(scene.shapes[0], id='overlap')
    scene = replace(scene, shapes=(*scene.shapes, duplicate))
    emitted, audit = write_label_anchor_native(scene, tmp_path/'out.pptx')
    assert emitted == scene
    assert set(objects(tmp_path/'out.pptx')) == {'a', 'b', 'overlap', 'c'}
    assert audit['label_anchor_adapter']['visible_label_count'] == 3
    # Count denotes nonblank emitted text, NOT correctness or unique useful labels.
    assert audit['label_anchor_adapter']['owners'][2]['attached_connector_ids'] == []


def test_clockwise_and_filtered_owner_fail_closed(tmp_path):
    scene = fixture()
    scene = replace(scene, shapes=(replace(scene.shapes[0], provenance=('region:text_clockwise',)), scene.shapes[1]))
    write_label_anchor_native(scene, tmp_path/'out.pptx')
    assert objects(tmp_path/'out.pptx')['a'].text_frame._txBody.bodyPr.get('vert') == 'vert'
    before = (tmp_path/'out.pptx').read_bytes()
    scene = replace(scene, texts=(EmitTextPrimitive('ocr', TextRegionRole.LABEL,
        BBox(15, 15, 60, 30), .99, text='Legend', owner_ids=('a',)),))
    with pytest.raises(ValueError, match='omitted by base writer'):
        write_label_anchor_native(scene, tmp_path/'out.pptx')
    assert (tmp_path/'out.pptx').read_bytes() == before


@pytest.mark.parametrize('start,end', [(False, False), (True, False), (False, True)])
def test_straight_arrow_flags_and_unconnected_anchor(tmp_path, start, end):
    scene = fixture()
    route = replace(scene.connectors[0], path_points=(Point(70, 22), Point(150, 122)),
                    arrowhead_start=start, arrowhead_end=end)
    spare = replace(scene.shapes[0], id='spare', label=None)
    scene = replace(scene, shapes=(*scene.shapes, spare), connectors=(route,))
    emitted, audit = write_label_anchor_native(scene, tmp_path/'out.pptx')
    assert emitted == scene
    items = objects(tmp_path/'out.pptx')
    for tag, active in [('headEnd', start), ('tailEnd', end)]:
        assert items['c']._element.xpath('.//a:'+tag)[0].get('type') == ('triangle' if active else 'none')
    assert len(items) == 4
    assert audit['label_anchor_adapter']['owners'][-1]['has_nonblank_text'] is False


def test_unsupported_path_stays_rejected(tmp_path):
    scene = fixture()
    route = replace(scene.connectors[0], path_points=tuple(Point(*p) for p in
        [(70, 22), (100, 22), (100, 180), (240, 180), (240, 130), (220, 130)]))
    emitted, audit = write_label_anchor_native(replace(scene, connectors=(route,)), tmp_path/'out.pptx')
    assert not emitted.connectors
    assert 'no supported preset' in audit['rejected_connectors']['c']
    assert set(objects(tmp_path/'out.pptx')) == {'a', 'b'}

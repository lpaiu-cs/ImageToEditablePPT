from copy import deepcopy
import numpy as np
import pytest
from image_to_editable_ppt.research.source_evidence_pipeline import detector_scene


def obj():
    return dict(id='opaque',role='node',semantic_kind='label_anchor',bbox_xyxy=[10,10,30,30],confidence=.9)


def test_only_detector_fields_and_no_geometry_changes():
    o=obj();s=detector_scene(np.zeros((100,100,3),np.uint8),[o])
    assert s.shapes[0].label is None and s.shapes[0].bbox.x0==10 and not s.connectors


@pytest.mark.parametrize('mode',['label','gt','edges','duplicate','negative','outside','nan','confidence','kind'])
def test_reject_untrusted_object_contract(mode):
    o=obj();objects=[o]
    if mode in ('label','gt','edges'):o[mode]='forbidden'
    if mode=='duplicate':objects.append(deepcopy(o))
    if mode=='negative':o['bbox_xyxy'][0]=-1
    if mode=='outside':o['bbox_xyxy'][2]=101
    if mode=='nan':o['bbox_xyxy'][0]=float('nan')
    if mode=='confidence':o['confidence']=1.1
    if mode=='kind':o['semantic_kind']='photo'
    with pytest.raises(ValueError):detector_scene(np.zeros((100,100,3),np.uint8),objects)


def test_native_pipeline_accepts_no_oracle_inputs(tmp_path):
    from image_to_editable_ppt.research.source_evidence_pipeline import reconstruct
    from pptx import Presentation
    calls=[]
    def empty_engine(image,**kwargs):
        calls.append(image.shape)
        return None,None
    out=tmp_path/'native.pptx'
    report=reconstruct(np.full((100,100,3),255,np.uint8),[obj()],out,engine=empty_engine)
    assert report['counts']['nodes']==1 and report['counts']['labels']==0 and calls
    shapes=Presentation(out).slides[0].shapes
    assert len(shapes)==1 and shapes[0].name=='opaque' and not shapes[0].text

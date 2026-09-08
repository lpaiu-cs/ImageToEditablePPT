from dataclasses import replace

import pytest
from pptx import Presentation

from image_to_editable_ppt.research.predicted_owners import prediction_scene, supported_owner_connections
from image_to_editable_ppt.region.emitter import write_native
from image_to_editable_ppt.v3.core.enums import ConnectorKind, NodeKind, PortOwnerKind
from image_to_editable_ppt.v3.core.types import Point
from image_to_editable_ppt.v3.emit.models import EmitConnectorPrimitive


def sample():
    return dict(image_size=[640, 360], objects=[
        dict(id='rgb-a:det:0', role='node', semantic_kind='label_anchor',
             bbox_xyxy=[10.5, 20, 40, 40], confidence=.91),
        dict(id='rgb-a:det:1', role='node', semantic_kind='box',
             bbox_xyxy=[100, 20, 140, 40], confidence=.92)])


def test_adapter_preserves_identity_geometry_and_semantics():
    data = sample()
    scene = prediction_scene(data)
    assert scene.shapes[0].shape_kind is NodeKind.LABEL_ANCHOR
    assert scene.shapes[0].bbox.x0 == 10.5
    assert [(s.id, s.confidence) for s in scene.shapes] == [(o['id'], o['confidence']) for o in data['objects']]
    assert all(s.label is None for s in scene.shapes)


def test_duplicate_identity_rejected():
    data = sample()
    data['objects'].append(data['objects'][0])
    with pytest.raises(ValueError, match='duplicate'):
        prediction_scene(data)


def test_unsupported_attachment_is_explicit_and_text_stays_editable(tmp_path):
    scene = prediction_scene(sample())
    scene = replace(scene, shapes=(replace(scene.shapes[0], label='A'), scene.shapes[1]))
    link = EmitConnectorPrimitive(id='e0', kind=ConnectorKind.ARROW, confidence=.9,
        source_owner_id=scene.shapes[0].id, source_owner_kind=PortOwnerKind.NODE,
        target_owner_id=scene.shapes[1].id, target_owner_kind=PortOwnerKind.NODE,
        source_port_id='a', target_port_id='b', path_points=(Point(40, 30), Point(100, 30)), arrowhead_end=True)
    scene = replace(scene, connectors=(link,))
    supported, rejected = supported_owner_connections(scene)
    assert rejected == {'e0': 'label_anchor_native_attachment_unsupported'}
    assert len(scene.connectors) == 1 and not supported.connectors
    output = tmp_path / 'result.pptx'
    write_native(supported, output)
    deck = Presentation(output)
    shapes = {s.name: s for s in deck.slides[0].shapes}
    assert len(shapes) == 2 and shapes['rgb-a:det:0'].text == 'A'
    assert shapes['rgb-a:det:0']._element.xpath('./p:spPr/a:noFill')
    assert not deck.slides[0]._element.xpath('.//p:pic')

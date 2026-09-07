from dataclasses import replace
from pptx import Presentation
from image_to_editable_ppt.research.path_emit import write_path_sites,path_audit
from image_to_editable_ppt.v3.core.types import Point,ImageSize,BBox
from image_to_editable_ppt.v3.core.enums import NodeKind,PortOwnerKind,ConnectorKind
from image_to_editable_ppt.v3.emit.models import EmitScene,EmitShapePrimitive,EmitConnectorPrimitive


def test_one_bound_native_connector_retains_every_polyline_vertex(tmp_path):
    a=EmitShapePrimitive('a',PortOwnerKind.NODE,NodeKind.BOX,BBox(20,20,120,80),1,label='Keep')
    b=replace(a,id='b',bbox=BBox(180,180,280,240))
    c=EmitConnectorPrimitive('edge',ConnectorKind.ARROW,1,'a',PortOwnerKind.NODE,'b',PortOwnerKind.NODE,'a:p','b:p',
        path_points=(Point(120,50),Point(210,50),Point(230,70),Point(230,180)),arrowhead_end=True)
    s=EmitScene(image_size=ImageSize(400,300),shapes=(a,b),connectors=(c,))
    path=tmp_path/'path.pptx'; write_path_sites(s,path)
    audit=path_audit(s,path)[0]
    assert audit['max_vertex_error_pt']<.001
    assert audit['vertices']==4 and audit['bound_endpoints']==2
    slide=Presentation(path).slides[0]
    assert len(slide.shapes)==3
    assert len(slide._element.xpath('.//p:cxnSp'))==1
    assert len(slide._element.xpath('.//p:pic'))==0


def test_zero_width_path_remains_vertical(tmp_path):
    a=EmitShapePrimitive('a',PortOwnerKind.NODE,NodeKind.BOX,BBox(20,20,120,80),1)
    b=replace(a,id='b',bbox=BBox(20,180,120,240))
    c=EmitConnectorPrimitive('edge',ConnectorKind.LINE,1,'a',PortOwnerKind.NODE,'b',PortOwnerKind.NODE,'a:p','b:p',
        path_points=(Point(70,80),Point(70,180)))
    s=EmitScene(image_size=ImageSize(400,300),shapes=(a,b),connectors=(c,))
    path=tmp_path/'vertical.pptx'; write_path_sites(s,path)
    assert path_audit(s,path)[0]['max_vertex_error_pt']<.001

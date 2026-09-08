import cv2
import numpy as np
from pptx import Presentation
from image_to_editable_ppt.research.closed_owner_regions import closed_regions,owner_region,place_owned_text
from image_to_editable_ppt.research.closed_rect_proposals import observe_text
from image_to_editable_ppt.research.source_evidence_pipeline import detector_scene
from image_to_editable_ppt.research.label_anchor_native import write_label_anchor_native
from image_to_editable_ppt.region.geometry import normalize_scene
from dataclasses import replace
import pytest


@pytest.mark.parametrize('scale',[1,2,3])
def test_thin_closed_child_still_reserves_pixels(scale):
    rgb=np.full((360*scale,640*scale,3),255,np.uint8)
    cv2.rectangle(rgb,(40*scale,30*scale),(310*scale,280*scale),(0,0,0),scale)
    cv2.rectangle(rgb,(65*scale,95*scale),(250*scale,105*scale),(0,0,0),scale)
    regions=closed_regions(rgb)
    parent=max(regions,key=lambda r:cv2.contourArea(r['contour']))
    own,excluded,fill,audit=owner_region(rgb,parent,regions)
    assert excluded[100*scale,150*scale]
    assert not own[100*scale,150*scale]


def test_unselected_narrow_child_cannot_supply_parent_color_or_text():
    rgb=np.full((360,640,3),255,np.uint8)
    cv2.rectangle(rgb,(40,30),(310,280),(220,240,210),-1)
    cv2.rectangle(rgb,(40,30),(310,280),(0,0,0),1)
    cv2.rectangle(rgb,(65,95),(285,260),(255,190,190),-1)
    cv2.rectangle(rgb,(65,95),(285,260),(0,0,0),1)
    regions=closed_regions(rgb)
    parent=max(regions,key=lambda r:cv2.contourArea(r['contour']))
    own,excluded,fill,audit=owner_region(rgb,parent,regions)
    assert fill==(220,240,210)
    assert excluded[150,150] and not own[150,150]
    assert not excluded[55,100]
    x,y,_,_=parent['bbox']
    def engine(image,**kwargs):
        s=image.shape[0]/(parent['bbox'][3]-y)
        return [(np.array([[a-x,b-y],[c-x,b-y],[c-x,d-y],[a-x,d-y]])*s,text,.99)
                for a,b,c,d,text in [(75,45,200,65,'Parent'),(80,120,250,145,'Child')]],None
    label,info=observe_text(rgb,parent['bbox'],engine,excluded=excluded)
    assert label=='Parent'
    assert all(word['text']=='Parent' for arm in info['arms'] for word in arm)


def test_parent_label_uses_native_owner_insets_and_no_extra_shape(tmp_path):
    rgb=np.full((360,640,3),255,np.uint8)
    scene=detector_scene(rgb,[dict(id='parent',role='node',semantic_kind='box',bbox_xyxy=[40,30,310,280],confidence=.9)])
    scene=replace(scene,shapes=(replace(scene.shapes[0],label='Parent'),))
    normalized,scale=normalize_scene(scene)
    path=tmp_path/'test.pptx'
    write_label_anchor_native(normalized,path)
    placement=place_owned_text(path,scene,{'parent':{'accepted':[{'bbox':[35,15,160,35],'text':'Parent'}]}},scale)
    deck=Presentation(path);obj=deck.slides[0].shapes[0]
    assert len(deck.slides[0].shapes)==1
    assert obj.text=='Parent'
    assert abs(obj.text_frame.margin_top-15*scale*9525)<=1
    assert not obj.text_frame.word_wrap
    assert placement[0]['source_text_bbox']==[35,15,160,35]

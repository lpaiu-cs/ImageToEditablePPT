from dataclasses import replace
import cv2
import numpy as np
import pytest
from image_to_editable_ppt.region.identity import filter_identity
from image_to_editable_ppt.region.text import ordered_text, recognize_consistent
from image_to_editable_ppt.region.geometry import validate_roi
from image_to_editable_ppt.v3.core.enums import NodeKind,ContainerKind,PortOwnerKind,ConnectorKind
from image_to_editable_ppt.v3.core.types import BBox,ImageSize,Point
from image_to_editable_ppt.v3.emit.models import EmitScene,EmitShapePrimitive,EmitConnectorPrimitive


def box(name,b,*,container=False,fill=False,label=None):
    return EmitShapePrimitive(name,PortOwnerKind.CONTAINER if container else PortOwnerKind.NODE,
        ContainerKind.FLOW_CLUSTER if container else NodeKind.BOX,BBox(*b),.9,label=label,
        source='visual_geometry:'+('uniform_fill' if fill else 'closed_outline'))


def edge(name,source,target,points):
    return EmitConnectorPrimitive(name,ConnectorKind.ARROW,.9,source,PortOwnerKind.NODE,target,PortOwnerKind.NODE,
        name+':start',name+':end',path_points=tuple(Point(*p) for p in points),arrowhead_end=True)


def test_gap_requires_two_independently_owned_paths():
    rgb=np.full((200,240,3),255,np.uint8)
    shapes=(box('top',(10,10,230,50)),box('bottom',(10,100,230,180),container=True),box('gap',(80,50,150,100)))
    for s in shapes: cv2.rectangle(rgb,tuple(map(int,(s.bbox.x0,s.bbox.y0))),tuple(map(int,(s.bbox.x1,s.bbox.y1))),(0,0,0),1)
    shapes+= (box('a',(70,120,90,140)),box('b',(140,120,160,140)))
    connections=(edge('c1','a','top',[(80,120),(80,50)]),edge('c2','b','top',[(150,120),(150,50)]))
    scene=EmitScene(ImageSize(240,200),shapes=shapes,connectors=connections)
    out,audit=filter_identity(scene,rgb,(0,0,1,1))
    assert 'gap' in audit['removed'] and len(out.shapes)==len(shapes)-1
    # An ordinary blank shared-boundary box without BOTH independent paths survives.
    for cs in ((),connections[:1],(edge('c','a','gap',[(80,120),(80,100)]),)):
        out,_=filter_identity(replace(scene,connectors=cs),rgb,(0,0,1,1))
        assert len(out.shapes)==len(shapes)
    labeled=replace(scene,shapes=tuple(replace(s,label='Legitimate') if s.id=='gap' else s for s in shapes))
    assert len(filter_identity(labeled,rgb,(0,0,1,1))[0].shapes)==len(shapes)


def test_fill_fragment_and_independent_small_box():
    rgb=np.full((200,240,3),(210,230,245),np.uint8)
    parent=box('parent',(5,5,235,195),container=True)
    fragment=box('fragment',(50,50,90,80),fill=True)
    scene=EmitScene(ImageSize(240,200),shapes=(parent,fragment))
    assert len(filter_identity(scene,rgb,(0,0,1,1))[0].shapes)==1
    for color in ((255,255,255),(230,30,30)):
        changed=rgb.copy(); changed[50:80,50:90]=color
        assert len(filter_identity(scene,changed,(0,0,1,1))[0].shapes)==2
    cv2.rectangle(rgb,(50,50),(90,80),(0,0,0),2)
    assert len(filter_identity(scene,rgb,(0,0,1,1))[0].shapes)==2


def snippet(text,x,y,score=.99):
    return [[[x,y],[x+18,y],[x+18,y+10],[x,y+10]],text,score]


def test_ocr_x_order_with_y_jitter():
    assert ordered_text([snippet('right',40,10),snippet('left',10,11)])[0]=='left right'
    assert ordered_text([snippet('line2',10,35),snippet('line1',10,10)])[0]=='line1\nline2'


def test_ocr_omits_entire_uncertain_label():
    assert ordered_text([snippet('Good',0,0),snippet('bad',30,0,.89)])[0] is None
    values=iter([([snippet('one',0,0)],None),([snippet('two',0,0)],None)])
    assert recognize_consistent(lambda crop:next(values),np.zeros((40,60,3),np.uint8))[1]['reason']=='scale_disagreement'
    assert recognize_consistent(None,np.zeros((10,10,3),np.uint8))[1]['reason']=='backend_unavailable'
    assert ordered_text([])[0] is None


@pytest.mark.parametrize('roi',[(0,0,0,1),(-.1,0,1,1),(0,0,float('nan'),1),(0,0,.01,.01),(0,0,1,float('inf'))])
def test_bad_roi(roi):
    with pytest.raises(ValueError): validate_roi(roi,(200,200))


def test_native_unsupported_route_is_reported(tmp_path):
    from image_to_editable_ppt.region.emitter import write_native
    from pptx import Presentation
    shapes=(box('a',(10,10,50,50)),box('b',(180,10,220,50)))
    route=edge('c','a','b',[(50,30),(80,30),(80,140),(160,140),(160,30),(180,30)])
    scene=EmitScene(ImageSize(240,200),shapes=shapes,connectors=(route,))
    emitted,audit=write_native(scene,tmp_path/'out.pptx')
    assert len(emitted.shapes)==2 and not emitted.connectors and 'c' in audit['rejected_connectors']
    assert len(Presentation(tmp_path/'out.pptx').slides[0].shapes)==2


def test_native_one_bound_connector_no_effects(tmp_path):
    from image_to_editable_ppt.region.emitter import write_native
    from image_to_editable_ppt.research.preset_emit import preset_vertices
    from pptx import Presentation
    shapes=(box('a',(10,10,50,50)),box('b',(150,110,200,150)))
    route=edge('c','a','b',[(50,30),(175,30),(175,110)])
    scene=EmitScene(ImageSize(240,200),shapes=shapes,connectors=(route,))
    emitted,audit=write_native(scene,tmp_path/'out.pptx')
    assert len(emitted.connectors)==1 and not audit['rejected_connectors']
    deck=Presentation(tmp_path/'out.pptx'); objects={s.name:s for s in deck.slides[0].shapes}
    assert set(objects)=={'a','b','c'}
    assert len(objects['c']._element.xpath('.//a:stCxn | .//a:endCxn'))==2
    assert objects['c']._element.xpath('.//a:effectLst')
    assert not deck.slides[0]._element.xpath('.//p:pic')
    points=preset_vertices(objects['c'])
    assert points[0]==pytest.approx((37.5,22.5),abs=.001)
    assert points[-1]==pytest.approx((131.25,82.5),abs=.001)


def test_cli_end_to_end_and_no_overwrite(tmp_path):
    import json
    from PIL import Image
    from image_to_editable_ppt.ml.convert_to_pptx import main
    rgb=np.full((240,400,3),255,np.uint8)
    cv2.rectangle(rgb,(30,60),(140,170),(0,0,0),2)
    cv2.rectangle(rgb,(260,60),(370,170),(0,0,0),2)
    cv2.arrowedLine(rgb,(140,115),(260,115),(0,0,0),2,tipLength=.15)
    source=tmp_path/'input.png'; Image.fromarray(rgb).save(source)
    output=tmp_path/'out.pptx'
    argv=[str(source),'-o',str(output),'--region','0','0','1','1','--no-ocr']
    assert main(argv)==0
    report=json.loads(output.with_suffix('.region.json').read_text())
    assert report['counts']['nodes']==2 and report['counts']['connectors']==1
    before=output.read_bytes()
    assert main(argv)==1 and output.read_bytes()==before
    assert not list(tmp_path.glob('.region-*'))


def test_region_blank_and_invalid_flags(tmp_path):
    import json
    from PIL import Image
    from image_to_editable_ppt.ml.convert_to_pptx import main
    source=tmp_path/'blank.png'; Image.new('RGB',(200,200),'white').save(source)
    output=tmp_path/'blank.pptx'
    argv=[str(source),'-o',str(output),'--region','0','0','1','1','--no-ocr']
    assert main(argv+['--no-gate'])==2 and not output.exists()
    assert main(argv)==0
    assert json.loads(output.with_suffix('.region.json').read_text())['status']=='empty'


def test_atomic_publication_race_rolls_back_only_own_report(tmp_path,monkeypatch):
    import os
    from PIL import Image
    from image_to_editable_ppt.region.convert import convert_region
    source=tmp_path/'blank.png'; Image.new('RGB',(200,200),'white').save(source)
    output=tmp_path/'out.pptx'; link=os.link
    def racing_link(src,dst):
        if str(dst).endswith('.pptx'):
            output.write_bytes(b'user-created-file')
        return link(src,dst)
    monkeypatch.setattr(os,'link',racing_link)
    with pytest.raises(FileExistsError): convert_region(source,output,region=(0,0,1,1),recover_text=False)
    assert output.read_bytes()==b'user-created-file'
    assert not output.with_suffix('.region.json').exists()


def test_runtime_has_no_research_imports():
    from pathlib import Path
    import ast
    directory=Path(__file__).parents[1]/'src/image_to_editable_ppt/region'
    for file in directory.glob('*.py'):
        for node in ast.walk(ast.parse(file.read_text())):
            if isinstance(node,ast.ImportFrom): assert '.research' not in (node.module or '')

from copy import deepcopy
import numpy as np
import pytest
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from image_to_editable_ppt.research.portrait_text import observe_portrait,source_line,place_portrait


@pytest.mark.parametrize('orientation,expected',[('clockwise',[30,10,50,100]),('counterclockwise',[20,100,40,190])])
def test_line_and_word_return_to_original_pixels(orientation,expected):
    line=dict(bbox=[10,20,100,40],text='Text',words=[dict(bbox=[10,20,100,40],text='Text')])
    selected=dict(orientation=orientation,spatial=dict(accepted=[deepcopy(line)]))
    observed=source_line(selected,70,200)
    assert observed['bbox']==observed['words'][0]['bbox']==expected
    assert selected['spatial']['accepted'][0]==line


@pytest.mark.parametrize('fill,text_color',[((255,255,255),(31,41,55)),((0,0,0),(255,255,255))])
@pytest.mark.parametrize('orientation,value',[('clockwise','vert'),('counterclockwise','vert270')])
def test_native_text_has_explicit_visible_color_and_no_owner_rotation(tmp_path,fill,text_color,orientation,value):
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
    shape=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,0,0,Emu(70*9525),Emu(200*9525))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(*fill)
    place_portrait(shape,dict(text='Text',bbox=[20,10,45,140]),70,200,orientation,1.)
    path=tmp_path/'native.pptx';deck.save(path);restored=Presentation(path).slides[0].shapes[0]
    assert len(slide.shapes)==1 and restored.rotation==0
    assert restored.text=='Text'
    assert restored.text_frame._txBody.bodyPr.get('vert')==value
    assert tuple(restored.text_frame.paragraphs[0].runs[0].font.color.rgb)==text_color


@pytest.mark.parametrize('score',[.97,1.])
def test_confidence_and_orientation_ambiguity_cannot_be_bypassed(score):
    def engine(image,**kwargs):
        if kwargs.get('use_det') is False:return [('Text',score)],None
        h,w=image.shape[:2]
        return [([[.1*w,.2*h],[.8*w,.2*h],[.8*w,.8*h],[.1*w,.8*h]],'Text',1.)],None
    selected,audit=observe_portrait(np.full((300,60,3),255,np.uint8),[0,0,60,300],np.zeros((300,60),np.uint8),engine)
    assert selected is None
    assert sum(a['approved'] for a in audit['arms'])==(0 if score<.98 else 2)

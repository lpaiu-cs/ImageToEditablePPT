from dataclasses import replace

from pptx import Presentation
import pytest

from image_to_editable_ppt.research.native_emit import connection_sites,normalize_scene,write_native_sites,endpoint_audit,custom_geometry
from image_to_editable_ppt.research.region_first import inspect_pptx
from image_to_editable_ppt.v3.core.enums import NodeKind,ConnectorKind,PortOwnerKind
from image_to_editable_ppt.v3.core.types import BBox,Point,ImageSize
from image_to_editable_ppt.v3.emit.models import EmitScene,EmitShapePrimitive,EmitConnectorPrimitive


def fan_in():
    target=EmitShapePrimitive("target",PortOwnerKind.NODE,NodeKind.BOX,BBox(100,20,500,120),1,label="Keep exact OCR string")
    sources=[EmitShapePrimitive(f"s{i}",PortOwnerKind.NODE,NodeKind.ELLIPSE,BBox(x-15,220,x+15,250),1) for i,x in enumerate((140,460))]
    connectors=tuple(EmitConnectorPrimitive(f"c{i}",ConnectorKind.ARROW,1,s.id,PortOwnerKind.NODE,"target",PortOwnerKind.NODE,
        f"s{i}:port","target:port",path_points=(Point(s.bbox.center.x,220),Point(s.bbox.center.x,120)),arrowhead_end=True) for i,s in enumerate(sources))
    return EmitScene(image_size=ImageSize(600,300),shapes=(target,*sources),connectors=connectors)


def test_normalization_is_independent_of_source_pixel_resolution():
    scene=fan_in()
    doubled=replace(scene,image_size=ImageSize(1200,600),
        shapes=tuple(replace(s,bbox=BBox(s.bbox.x0*2,s.bbox.y0*2,s.bbox.x1*2,s.bbox.y1*2)) for s in scene.shapes),
        connectors=tuple(replace(c,path_points=tuple(Point(p.x*2,p.y*2) for p in c.path_points)) for c in scene.connectors))
    assert normalize_scene(scene)==normalize_scene(doubled)


def test_distinct_observed_target_sites_remain_distinct():
    sites,bindings=connection_sites(fan_in())
    assert bindings[("c0","end")]!=bindings[("c1","end")]
    assert sites["target"][bindings[("c0","end")]][:2]==(.1,1)
    assert sites["target"][bindings[("c1","end")]][:2]==(.9,1)


def test_custom_emission_has_exact_endpoints_and_no_proxy_shapes(tmp_path):
    scene=normalize_scene(fan_in())
    path=tmp_path/"sites.pptx"
    write_native_sites(scene,path)
    assert max(a["max_coordinate_error_px"] for a in endpoint_audit(path,scene))<.001
    native=inspect_pptx(path)
    assert native["native_shapes"]==5
    assert native["raster_images"]==0
    assert native["bound_endpoints"]==4
    ppt=Presentation(path)
    target=next(s for s in ppt.slides[0].shapes if s.name=="target")
    assert target.text=="Keep exact OCR string"
    assert len(target._element.xpath(".//a:custGeom"))==1
    assert len(target._element.xpath(".//a:cxnLst/a:cxn"))==6
    assert len(target._element.xpath(".//a:pathLst/a:path/a:close"))==1


@pytest.mark.parametrize("kind",[NodeKind.BOX,NodeKind.ROUNDED_BOX,NodeKind.ELLIPSE])
def test_simple_path_is_single_closed_native_object(kind):
    shape=replace(fan_in().shapes[0],shape_kind=kind)
    geom=custom_geometry(shape,[(.2,1,5400000)])
    assert len(geom.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}pathLst"))==1
    assert b"*/ w 200000 1000000" in __import__("lxml.etree",fromlist=["tostring"]).tostring(geom)


def test_unsupported_path_is_rejected_before_output(tmp_path):
    scene=fan_in()
    edge=replace(scene.connectors[0],path_points=(Point(140,220),Point(140,160),Point(200,160)))
    path=tmp_path/"bad.pptx"
    with pytest.raises(ValueError,match="two-point"):
        write_native_sites(replace(scene,connectors=(edge,)),path)
    assert not path.exists()


def test_unknown_endpoint_is_not_silently_unbound():
    scene=fan_in()
    edge=replace(scene.connectors[0],target_owner_id="missing")
    with pytest.raises(ValueError,match="unknown"):
        connection_sites(replace(scene,connectors=(edge,)))

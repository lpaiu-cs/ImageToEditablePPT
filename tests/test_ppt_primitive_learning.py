import numpy as np
import pytest
import torch
from pptx import Presentation
from image_to_editable_ppt.research.ppt_primitive_learning import (
    PrimitiveUNet,PALETTE,palette_target,components,object_counts,write_shapes)


def test_palette_antialias_and_rgb_boundary():
    assert palette_target(PALETTE[None]).tolist()==[list(range(7))]
    assert palette_target(np.array([[[128,0,0]]],np.uint8))[0,0]==255
    with pytest.raises(ValueError):palette_target(np.zeros((8,8),np.uint8))


def test_prediction_components_and_one_to_one_matching(tmp_path):
    p=np.zeros((7,32,40),np.float32);p[0]=1;p[0,4:20,6:30]=0;p[4,4:20,6:30]=1
    predicted=components(p,(80,64),.95)
    assert len(predicted)==1 and predicted[0]['bbox']==[12,8,60,40] and predicted[0]['kind']==4
    assert object_counts(predicted*2,predicted)==dict(tp=1,fp=1,fn=0)
    assert object_counts(predicted,[])==dict(tp=0,fp=1,fn=0)
    output=tmp_path/'result.pptx';write_shapes(predicted,(160,128),output)
    assert len(Presentation(output).slides[0].shapes)==1
    with pytest.raises(FileExistsError):write_shapes(predicted,(160,128),output)


def test_reused_unet_supports_six_presets():
    torch.set_num_threads(1);model=PrimitiveUNet()
    assert model(torch.zeros(1,3,32,40)).shape==(1,7,32,40)


def test_object_matching_maximizes_valid_count(monkeypatch):
    from image_to_editable_ppt.research import closed_rect_proposals
    scores=[[.99,.51],[.51,.49]]
    monkeypatch.setattr(closed_rect_proposals,'_iou',lambda a,b:scores[a[0]][b[0]])
    a=[dict(kind=1,bbox=[i]) for i in range(2)]
    assert object_counts(a,a)==dict(tp=2,fp=0,fn=0)

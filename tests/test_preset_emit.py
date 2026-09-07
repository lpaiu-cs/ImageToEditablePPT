from pptx import Presentation
from image_to_editable_ppt.research.preset_emit import fit_path,preset_vertices,write_preset_paths
from image_to_editable_ppt.v3.core.types import Point,ImageSize,BBox
from image_to_editable_ppt.v3.core.enums import NodeKind,PortOwnerKind,ConnectorKind
from image_to_editable_ppt.v3.emit.models import EmitScene,EmitShapePrimitive,EmitConnectorPrimitive
import pytest


def test_native_single_bend_and_bound_endpoints(tmp_path):
    a=EmitShapePrimitive('a',PortOwnerKind.NODE,NodeKind.BOX,BBox(20,20,120,80),1)
    b=EmitShapePrimitive('b',PortOwnerKind.NODE,NodeKind.BOX,BBox(180,180,280,240),1)
    c=EmitConnectorPrimitive('edge',ConnectorKind.ARROW,1,'a',PortOwnerKind.NODE,'b',PortOwnerKind.NODE,'a:p','b:p',
        path_points=(Point(120,50),Point(230,50),Point(230,180)),arrowhead_end=True)
    s=EmitScene(image_size=ImageSize(400,300),shapes=(a,b),connectors=(c,))
    path=tmp_path/'preset.pptx'; fits=write_preset_paths(s,path)
    assert fits['edge']['error_px']==0
    deck=Presentation(path); obj=next(s for s in deck.slides[0].shapes if s.name=='edge')
    points=preset_vertices(obj)
    assert points[0]==(90,37.5) and points[-1]==(172.5,135)
    assert len(deck.slides[0].shapes)==3


def test_long_unsupported_loop_fails_before_output():
    with pytest.raises(ValueError,match='no supported preset'):
        fit_path((Point(0,0),Point(0,100),Point(100,100),Point(100,0)))

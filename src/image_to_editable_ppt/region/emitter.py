"""One supported preset per connection; explicit sites, neutral style, no helper objects."""
from dataclasses import replace
import math
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.util import Pt
from image_to_editable_ppt.region._native_geometry import custom_geometry, element
from image_to_editable_ppt.region._path_fit import fit_path
from image_to_editable_ppt.v3.emit.pptx_writer import write_pptx, EMU_PER_PIXEL


def supported_scene(scene):
    owners={s.id:s for s in scene.shapes}; kept=[]; fits={}; rejected={}
    if len(owners)!=len(scene.shapes): raise ValueError('duplicate shape IDs')
    ids=set(owners)
    for c in scene.connectors:
        if c.id in ids: raise ValueError('duplicate primitive ID')
        ids.add(c.id)
        try:
            if len(c.path_points)<2 or any(not math.isfinite(v) for p in c.path_points for v in (p.x,p.y)):
                raise ValueError('invalid path')
            for owner,p in ((c.source_owner_id,c.path_points[0]),(c.target_owner_id,c.path_points[-1])):
                if owner not in owners: raise ValueError('unknown endpoint owner')
                b=owners[owner].bbox
                if b.width<=0 or b.height<=0: raise ValueError('degenerate owner')
                u,v=(p.x-b.x0)/b.width,(p.y-b.y0)/b.height
                if not (-1e-6<=u<=1+1e-6 and -1e-6<=v<=1+1e-6): raise ValueError('endpoint outside owner')
                if min(abs(u),abs(v),abs(1-u),abs(1-v))>1e-5: raise ValueError('endpoint not on owner boundary')
            fits[c.id]=fit_path(c.path_points); kept.append(c)
        except ValueError as error: rejected[c.id]=str(error)
    return replace(scene,connectors=tuple(kept)),fits,rejected


def write_native(scene,path,*,styles=None):
    scene,fits,rejected=supported_scene(scene)
    owners={s.id:s for s in scene.shapes}; sites={}; bindings={}
    for c in scene.connectors:
        for end,owner,p in [('start',c.source_owner_id,c.path_points[0]),('end',c.target_owner_id,c.path_points[-1])]:
            b=owners[owner].bbox; u=max(0,min(1,(p.x-b.x0)/b.width)); v=max(0,min(1,(p.y-b.y0)/b.height))
            angle=min(((v,270),(u,180),(1-v,90),(1-u,0)))[1]*60000
            site=(round(u,6),round(v,6),angle)
            current=sites.setdefault(owner,[(.5,0,16200000),(0,.5,10800000),(.5,1,5400000),(1,.5,0)])
            if site not in current: current.append(site)
            bindings[c.id,end]=current.index(site)
    geometries={key:custom_geometry(owners[key],value) for key,value in sites.items()}
    write_pptx(scene,path,styles=styles)
    deck=Presentation(path); objects={s.name:s for s in deck.slides[0].shapes}
    for owner,geom in geometries.items():
        props=objects[owner]._element.spPr; props.replace(props.find(qn('a:prstGeom')),geom)
    for c in scene.connectors:
        shape=objects[c.id]; props=shape._element.spPr; fit=fits[c.id]
        cnv=shape._element.find(qn('p:nvCxnSpPr')).find(qn('p:cNvCxnSpPr'))
        for tag,end,owner in [('stCxn','start',c.source_owner_id),('endCxn','end',c.target_owner_id)]:
            old=cnv.find(qn('a:'+tag))
            if old is not None: cnv.remove(old)
            cnv.append(element(tag,id=objects[owner].shape_id,idx=bindings[c.id,end]))
        a,b=c.path_points[0],c.path_points[-1]; xf=props.xfrm
        xf.off.x=round(min(a.x,b.x)*EMU_PER_PIXEL); xf.off.y=round(min(a.y,b.y)*EMU_PER_PIXEL)
        xf.ext.cx=max(1,round(abs(a.x-b.x)*EMU_PER_PIXEL)); xf.ext.cy=max(1,round(abs(a.y-b.y)*EMU_PER_PIXEL))
        xf.attrib.pop('rot',None)
        for name,condition in [('flipH',a.x>b.x),('flipV',a.y>b.y)]:
            if condition: xf.set(name,'1')
            else: xf.attrib.pop(name,None)
        geom=element('prstGeom',prst=fit['preset']); av=element('avLst'); geom.append(av)
        if fit['adj'] is not None: av.append(element('gd',name='adj1',fmla=f"val {round(fit['adj']*100000)}"))
        props.replace(props.find(qn('a:prstGeom')),geom)
        shape.shadow.inherit=False
        shape.line.color.rgb=RGBColor(31,41,55); shape.line.width=Pt(1)
        for tag,active in [('headEnd',c.arrowhead_start),('tailEnd',c.arrowhead_end)]:
            old=props.ln.find(qn('a:'+tag))
            if old is not None: props.ln.remove(old)
            props.ln.append(element(tag,type='triangle' if active else 'none',w='sm',len='sm'))
    deck.save(path)
    return scene,dict(fits=fits,rejected_connectors=rejected)

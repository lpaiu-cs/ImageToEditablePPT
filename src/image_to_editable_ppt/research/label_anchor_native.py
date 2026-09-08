"""Research-only text-owner attachment; semantic kinds never leave this boundary changed."""
from copy import deepcopy
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.oxml.ns import qn

from image_to_editable_ppt.region.emitter import write_native
from image_to_editable_ppt.v3.core.enums import NodeKind
from image_to_editable_ppt.v3.emit.pptx_writer import write_pptx
from image_to_editable_ppt.v3.emit.style import ShapeVisualStyle


def write_label_anchor_native(scene, path, *, styles=None):
    """Emit native sites on existing borderless owners, returning semantic scene/audit.

    Blank owners can have native attachments but are NEVER visible-label successes.
    No OCR, route inference, extra objects, or changes to product writer behavior.
    The output is replaced only after both temporary decks have been validated.
    """
    anchors = {s.id: s for s in scene.shapes if s.shape_kind is NodeKind.LABEL_ANCHOR}
    if not anchors:
        emitted, audit = write_native(scene, path, styles=styles)
        return emitted, {**audit, 'label_anchor_adapter': {'owners': [], 'visible_label_count': 0}}
    physical = replace(scene, shapes=tuple(
        replace(s, shape_kind=NodeKind.BOX) if s.id in anchors else s for s in scene.shapes))
    physical_styles = dict(styles or {})
    physical_styles.update({key: ShapeVisualStyle() for key in anchors})
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix='.label-anchor-', dir=output.parent) as temporary:
        reference_path = Path(temporary) / 'text-reference.pptx'
        native_path = Path(temporary) / 'native.pptx'
        # Reference writer retains owned-text selection, exact font sizing and no-wrap.
        write_pptx(replace(scene, connectors=()), reference_path, styles=styles)
        reference = {s.name: s for s in Presentation(reference_path).slides[0].shapes}
        missing = anchors.keys() - reference.keys()
        if missing:
            raise ValueError(f'label anchors omitted by base writer: {sorted(missing)}')
        emitted, audit = write_native(physical, native_path, styles=physical_styles)
        deck = Presentation(native_path)
        objects = {s.name: s for s in deck.slides[0].shapes}
        owners = []
        for key, semantic in anchors.items():
            target = objects[key]._element
            body = deepcopy(reference[key]._element.find(qn('p:txBody')))
            if semantic.label and 'region:text_clockwise' in semantic.provenance:
                body.find(qn('a:bodyPr')).set('vert', 'vert')
            target.replace(target.find(qn('p:txBody')), body)
            text = objects[key].text
            connections = [c.id for c in emitted.connectors
                           if key in (c.source_owner_id, c.target_owner_id)]
            owners.append({'id': key, 'semantic_kind': 'label_anchor',
                           'writer_geometry': 'borderless_rect',
                           'has_nonblank_text': bool(text.strip()),
                           'attached_connector_ids': connections})
        deck.save(native_path)
        native_path.replace(output)
    semantic_emitted = replace(scene, connectors=emitted.connectors)
    return semantic_emitted, {**audit, 'label_anchor_adapter': {
        'owners': owners,
        'visible_label_count': sum(owner['has_nonblank_text'] for owner in owners),
        'blank_policy': 'retain_existing_owner; native_attachment_is_not_visible_label_success',
    }}

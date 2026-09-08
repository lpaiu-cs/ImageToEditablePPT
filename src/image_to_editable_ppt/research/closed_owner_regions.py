"""P2: raw closed descendants reserve their own pixels, including omitted nodes."""
import cv2
import numpy as np


def closed_regions(rgb):
    contours, hierarchy = cv2.findContours(
        cv2.Canny(cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY), 20, 60),
        cv2.RETR_CCOMP, cv2.CHAIN_APPROX_NONE)
    scale = rgb.shape[0] / 360
    regions = []
    for contour, entry in zip(contours, [] if hierarchy is None else hierarchy[0]):
        if entry[3] < 0:
            continue
        x, y, w, h = cv2.boundingRect(contour)
        # This is an exclusion, not a new node proposal or semantic approval.
        size_ok = (min(w, h) >= 12 * scale and max(w, h) >= 20 * scale) or (
            min(w, h) >= 6 * scale and max(w, h) >= 40 * scale)
        if size_ok and cv2.contourArea(contour) >= .85 * w * h:
            regions.append(dict(bbox=[x, y, x+w, y+h], contour=contour))
    return regions


def owner_region(rgb, candidate, regions):
    x0, y0, x1, y1 = candidate['bbox']
    contour = candidate['contour']
    own = np.zeros(rgb.shape[:2], np.uint8)
    cv2.drawContours(own, [contour], -1, 1, cv2.FILLED)
    own = cv2.erode(own, np.ones((3, 3), np.uint8))
    excluded = np.zeros_like(own)
    children = []
    for child in regions:
        a, b, c, d = child['bbox']
        if not (x0 < a < c < x1 and y0 < b < d < y1):
            continue
        # The two faces of the owner's ink band are not descendants.
        if (c-a)*(d-b) >= .8*(x1-x0)*(y1-y0):
            continue
        if not all(cv2.pointPolygonTest(contour, tuple(map(float, p)), False) > 0
                   for p in child['contour'][:, 0, :]):
            continue
        cv2.drawContours(excluded, [child['contour']], -1, 1, cv2.FILLED)
        children.append(child['bbox'])
    excluded = cv2.dilate(excluded, np.ones((7, 7), np.uint8))
    own &= 1-excluded
    # Median is safe only after removing all raw descendants, not selected nodes.
    pixels = rgb[own != 0]
    fill = tuple(map(int, np.median(pixels, axis=0))) if len(pixels) >= max(12, .1*cv2.contourArea(contour)) else None
    return own, excluded, fill, dict(children=children, owned_pixels=int(own.sum()))


def place_owned_text(pptx, scene, text_audit, scale):
    """Keep labels in their native owner; source positions become text insets."""
    from pptx import Presentation
    from pptx.util import Pt, Emu
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from image_to_editable_ppt.v3.emit.pptx_writer import EMU_PER_PIXEL
    from PIL import ImageFont
    from pathlib import Path
    font_path=Path('C:/Windows/Fonts/arial.ttf')
    if not font_path.is_file():
        raise RuntimeError('Arial metrics are required for this locked Windows research renderer')
    metric_font=ImageFont.truetype(str(font_path),100)
    deck = Presentation(pptx)
    objects = {s.name: s for s in deck.slides[0].shapes}
    audit = []
    for owner in scene.shapes:
        lines = text_audit[owner.id]['accepted']
        if not lines:
            continue
        frame = objects[owner.id].text_frame
        bx = [min(l['bbox'][0] for l in lines), min(l['bbox'][1] for l in lines),
              max(l['bbox'][2] for l in lines), max(l['bbox'][3] for l in lines)]
        # OCR boxes are owner-local source pixels. Insets retain parent headings.
        emu = lambda v: Emu(round(max(0, v)*scale*EMU_PER_PIXEL))
        frame.margin_left = emu(bx[0])
        frame.margin_right = emu(owner.bbox.width-bx[2])
        frame.margin_top = emu(bx[1])
        frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.word_wrap = False
        for index,(paragraph, line) in enumerate(zip(frame.paragraphs, lines)):
            height = line['bbox'][3]-line['bbox'][1]
            # OCR glyph windows may overlap vertically. Their sum is not the
            # source paragraph height; advance by observed row origins instead.
            advance=lines[index+1]['bbox'][1]-line['bbox'][1] if index+1<len(lines) else height
            advance=max(1,min(height,advance))
            # The glyph box is not a font em box. Use a conservative metric fit.
            width=max(1,min(bx[2]-bx[0],owner.bbox.width-bx[0]-1))
            width_fit=width*100/max(1,metric_font.getlength(line['text']))
            size = max(1, min(advance*.85,width_fit*.95)*scale*.75)
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = Pt(advance*scale*.75)
            for run in paragraph.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(size)
        audit.append(dict(owner=owner.id, source_text_bbox=bx, native_text=objects[owner.id].text))
    deck.save(pptx)
    return audit

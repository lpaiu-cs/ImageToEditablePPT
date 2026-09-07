"""R3-E one native connector per observed polyline, never helper segment objects."""
from pathlib import Path
import math
from pptx import Presentation
from pptx.oxml.ns import qn
from image_to_editable_ppt.research.native_emit import custom_geometry,element
from image_to_editable_ppt.v3.emit.pptx_writer import write_pptx,EMU_PER_PIXEL


def write_path_sites(scene,path,*,styles=None):
    owners={s.id:s for s in scene.shapes}
    sites={}; binding={}
    for c in scene.connectors:
        if len(c.path_points)<2: raise ValueError('missing source path')
        for end,owner,p in [('start',c.source_owner_id,c.path_points[0]),('end',c.target_owner_id,c.path_points[-1])]:
            b=owners[owner].bbox
            u,v=(p.x-b.x0)/b.width,(p.y-b.y0)/b.height
            if not (0<=u<=1 and 0<=v<=1): raise ValueError('endpoint outside owner')
            angle=min(((v,270),(u,180),(1-v,90),(1-u,0)))[1]*60000
            site=(round(u,6),round(v,6),angle)
            current=sites.setdefault(owner,[(.5,0,16200000),(0,.5,10800000),(.5,1,5400000),(1,.5,0)])
            if site not in current: current.append(site)
            binding[c.id,end]=current.index(site)
    geometries={s.id:custom_geometry(s,sites[s.id]) for s in scene.shapes if s.id in sites}
    write_pptx(scene,path,styles=styles)
    deck=Presentation(path); shapes={s.name:s for s in deck.slides[0].shapes}
    for owner,geom in geometries.items():
        props=shapes[owner]._element.spPr
        props.replace(props.find(qn('a:prstGeom')),geom)
    for c in scene.connectors:
        shape=shapes[c.id]; props=shape._element.spPr
        cnv=shape._element.find(qn('p:nvCxnSpPr')).find(qn('p:cNvCxnSpPr'))
        for tag in ('stCxn','endCxn'):
            old=cnv.find(qn('a:'+tag))
            if old is not None: cnv.remove(old)
        cnv.append(element('stCxn',id=shapes[c.source_owner_id].shape_id,idx=binding[c.id,'start']))
        cnv.append(element('endCxn',id=shapes[c.target_owner_id].shape_id,idx=binding[c.id,'end']))
        xs=[p.x*EMU_PER_PIXEL for p in c.path_points]; ys=[p.y*EMU_PER_PIXEL for p in c.path_points]
        x0,y0=round(min(xs)),round(min(ys))
        w,h=max(1,round(max(xs))-x0),max(1,round(max(ys))-y0)
        xf=props.xfrm
        xf.off.x=x0; xf.off.y=y0; xf.ext.cx=w; xf.ext.cy=h
        for key in ('flipH','flipV','rot'): xf.attrib.pop(key,None)
        geom=element('custGeom')
        for tag in ('avLst','gdLst','ahLst','cxnLst'): geom.append(element(tag))
        geom.append(element('rect',l='0',t='0',r='w',b='h'))
        paths=element('pathLst'); geom.append(paths)
        poly=element('path',w=100000,h=100000,fill='none'); paths.append(poly)
        for i,(x,y) in enumerate(zip(xs,ys)):
            cmd=element('moveTo' if i==0 else 'lnTo')
            cmd.append(element('pt',x=round((x-x0)/w*100000),y=round((y-y0)/h*100000)))
            poly.append(cmd)
        props.replace(props.find(qn('a:prstGeom')),geom)
    deck.save(path)
    return {k:[list(s) for s in v] for k,v in sites.items()}


def read_vertices(shape):
    """Read actual serialized polyline coordinates in pt, including native transform."""
    props=shape._element.spPr
    geom=props.find(qn('a:custGeom'))
    if geom is None:
        preset=props.find(qn('a:prstGeom'))
        if preset is not None and preset.get('prst') in ('line','straightConnector1'):
            return [(shape.begin_x/12700,shape.begin_y/12700),(shape.end_x/12700,shape.end_y/12700)]
        raise ValueError('PowerPoint replaced polyline with unsupported preset')
    paths=geom.find(qn('a:pathLst'))
    if len(paths)!=1: raise ValueError('expected one path')
    path=paths[0]; w,h=float(path.get('w')),float(path.get('h'))
    xf=props.xfrm; out=[]
    for cmd in path:
        if cmd.tag not in (qn('a:moveTo'),qn('a:lnTo')): raise ValueError('non-polyline native path')
        point=cmd[0]
        u,v=float(point.get('x'))/w,float(point.get('y'))/h
        if xf.get('flipH') in ('1','true'): u=1-u
        if xf.get('flipV') in ('1','true'): v=1-v
        x,y=(u-.5)*xf.ext.cx,(v-.5)*xf.ext.cy
        angle=float(xf.get('rot','0'))/60000*math.pi/180
        x,y=x*math.cos(angle)-y*math.sin(angle),x*math.sin(angle)+y*math.cos(angle)
        out.append(((xf.off.x+.5*xf.ext.cx+x)/12700,(xf.off.y+.5*xf.ext.cy+y)/12700))
    return out


def path_audit(scene,path):
    shapes={s.name:s for s in Presentation(path).slides[0].shapes}
    records=[]
    for c in scene.connectors:
        actual=read_vertices(shapes[c.id]); expected=[(p.x*.75,p.y*.75) for p in c.path_points]
        error=max(abs(a-b) for p,q in zip(actual,expected) for a,b in zip(p,q)) if len(actual)==len(expected) else None
        records.append(dict(id=c.id,actual_pt=actual,expected_pt=expected,max_vertex_error_pt=error,
            vertices=len(actual),expected_vertices=len(expected),bound_endpoints=len(shapes[c.id]._element.xpath('.//a:stCxn | .//a:endCxn'))))
    return records

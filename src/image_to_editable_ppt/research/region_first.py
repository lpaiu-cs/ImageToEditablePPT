"""R0 intervention harness. Manual scope/boxes are never automatic performance.

Production ML, Phase 15/16 annotations and acceptance gates are not modified.
Only the evaluator receives labels, exclusions and expected connector pairs.
"""
from __future__ import annotations

from dataclasses import asdict, replace
import hashlib
import json
from pathlib import Path
import re

import numpy as np
from PIL import Image, ImageDraw

from image_to_editable_ppt.ml.annotation_schema import AnnotationBBox, AnnotationContainer, AnnotationNode
from image_to_editable_ppt.ml.visual_geometry import detect_visual_geometry
from image_to_editable_ppt.ml.visual_connectors import visual_connectors
from image_to_editable_ppt.v3.core.enums import ContainerKind, NodeKind, PortOwnerKind
from image_to_editable_ppt.v3.core.types import BBox, ImageSize, Point
from image_to_editable_ppt.v3.emit.models import EmitConnectorPrimitive, EmitScene, EmitShapePrimitive


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pixel_box(box, size) -> tuple[int, int, int, int]:
    w, h = size
    return tuple(round(v * (w if i % 2 == 0 else h)) for i, v in enumerate(box))


def intersection(a, b) -> float:
    return max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(0, min(a[3], b[3]) - max(a[1], b[1]))


def area(a) -> float:
    return max(0, a[2] - a[0]) * max(0, a[3] - a[1])


def iou(a, b) -> float:
    shared = intersection(a, b)
    return shared / max(1, area(a) + area(b) - shared)


def validate_sample(sample, root: Path) -> Image.Image:
    if "holdout" in sample["path"].lower():
        raise ValueError("R0 must not consume holdout inputs")
    path = root / sample["path"]
    if digest(path) != sample["sha256"]:
        raise ValueError(f"source hash mismatch: {path}")
    with Image.open(path) as source:
        image = source.convert("RGB")
    if list(image.size) != sample["size"]:
        raise ValueError("source size mismatch")
    boxes = [sample["roi"], *sample["excluded"], *(o["bbox"] for o in sample["objects"])]
    for box in boxes:
        if len(box) != 4 or not all(0 <= x <= 1 for x in box) or box[0] >= box[2] or box[1] >= box[3]:
            raise ValueError(f"invalid normalized box: {box}")
    ids = [o["id"] for o in sample["objects"]]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate target ID")
    for obj in sample["objects"]:
        if obj["role"] not in ("node", "container"):
            raise ValueError("invalid role")
        if obj["role"] == "node":
            NodeKind(obj["shape"])
        if intersection(obj["bbox"], sample["roi"]) < area(obj["bbox"]) - 1e-6:
            raise ValueError(f"target outside ROI: {obj['id']}")
    if any(a not in ids or b not in ids for a, b in sample["edges"]):
        raise ValueError("unknown target endpoint")
    return image


def manual_box_inputs(sample):
    """Explicit information boundary: no label, edge, exclusion, or target importance."""
    return [{k: obj[k] for k in ("id", "bbox", "role", "shape")} for obj in sample["objects"]]


def recover_objects(rgb, *, offset=(0, 0), manual_boxes=None, source_size=None):
    nodes, containers = [], []
    if manual_boxes is not None:
        for obj in manual_boxes:
            bbox = AnnotationBBox(*pixel_box(obj["bbox"], source_size))
            args = dict(id=obj["id"], bbox=bbox, confidence=1.0, source="r0:manual_box_and_role",
                        provenance=("diagnostic:manual_input_not_automatic",))
            if obj["role"] == "node":
                nodes.append(AnnotationNode(kind=NodeKind(obj["shape"]), **args))
            else:
                containers.append(AnnotationContainer(kind=ContainerKind.FLOW_CLUSTER, **args))
        return tuple(nodes), tuple(containers)
    geometry = detect_visual_geometry(rgb)
    for role, candidates in (("node", geometry.nodes), ("container", geometry.containers)):
        for i, candidate in enumerate(candidates):
            x0, y0, x1, y1 = candidate.bbox
            bbox = AnnotationBBox(x0 + offset[0], y0 + offset[1], x1 + offset[0], y1 + offset[1])
            args = dict(id=f"{role}:r0:{i}", bbox=bbox, confidence=candidate.confidence,
                        source=f"visual_geometry:{candidate.source}", provenance=("r0:frozen_cv_candidate",))
            if role == "node":
                kind = NodeKind.ELLIPSE if candidate.source == "closed_ellipse" else (
                    NodeKind.ROUNDED_BOX if candidate.source == "uniform_fill" else NodeKind.BOX)
                nodes.append(AnnotationNode(kind=kind, **args))
            else:
                containers.append(AnnotationContainer(kind=ContainerKind.FLOW_CLUSTER, **args))
    return tuple(nodes), tuple(containers)


def recover_scene(image: Image.Image, *, roi=None, manual_boxes=None, ocr=True) -> EmitScene:
    """Source-only A, source+ROI B, source+ROI+boxes/roles C. No expected answers."""
    rgb = np.asarray(image)
    w, h = image.size
    scope = pixel_box(roi, image.size) if roi is not None else (0, 0, w, h)
    x0, y0, x1, y1 = scope
    crop = rgb[y0:y1, x0:x1]
    nodes, containers = recover_objects(crop, offset=(x0, y0), manual_boxes=manual_boxes, source_size=image.size)
    # Apply the SAME scope to every downstream stage; do not let paths outside a manual ROI leak in.
    local_nodes = tuple(replace(n, bbox=AnnotationBBox(n.bbox.x0-x0, n.bbox.y0-y0, n.bbox.x1-x0, n.bbox.y1-y0)) for n in nodes)
    local_containers = tuple(replace(n, bbox=AnnotationBBox(n.bbox.x0-x0, n.bbox.y0-y0, n.bbox.x1-x0, n.bbox.y1-y0)) for n in containers)
    candidates, _ = visual_connectors(crop, local_nodes, containers=local_containers, image_id="r0")
    engine = None
    if ocr and nodes:
        from image_to_editable_ppt.v3.text.ocr import _load_engine
        engine = _load_engine()
        if engine is None:
            raise RuntimeError("OCR unavailable: do not silently compare an OCR-disabled run")
    shapes = []
    for role, objects in ((PortOwnerKind.NODE, nodes), (PortOwnerKind.CONTAINER, containers)):
        for obj in objects:
            label = None
            if engine is not None and role is PortOwnerKind.NODE:
                from image_to_editable_ppt.v3.text.ocr import _recognize_crop
                a, b = max(x0, round(obj.bbox.x0)+3), max(y0, round(obj.bbox.y0)+3)
                c, d = min(x1, round(obj.bbox.x1)-3), min(y1, round(obj.bbox.y1)-3)
                if c > a and d > b:
                    label, _ = _recognize_crop(engine, rgb[b:d, a:c], min_snippet_confidence=0.90)
            shapes.append(EmitShapePrimitive(id=obj.id, owner_kind=role, shape_kind=obj.kind,
                bbox=BBox(obj.bbox.x0,obj.bbox.y0,obj.bbox.x1,obj.bbox.y1), confidence=obj.confidence,
                label=label, source=obj.source, provenance=obj.provenance))
    connectors = tuple(EmitConnectorPrimitive(id=c.id, kind=c.kind, confidence=c.confidence,
        source_owner_id=c.start_endpoint.owner_id, source_owner_kind=c.start_endpoint.owner_kind,
        target_owner_id=c.end_endpoint.owner_id, target_owner_kind=c.end_endpoint.owner_kind,
        source_port_id=f"{c.id}:start", target_port_id=f"{c.id}:end",
        path_points=tuple(Point(p.x+x0,p.y+y0) for p in c.path_points),
        arrowhead_start=c.arrowhead_start, arrowhead_end=c.arrowhead_end,
        source=c.source, provenance=c.provenance) for c in candidates)
    return EmitScene(image_size=ImageSize(w,h), shapes=tuple(shapes), connectors=connectors,
                     source="r0_diagnostic_not_production")


def serialize_scene(scene):
    return json.loads(json.dumps(asdict(scene)))


def scene_for_pptx(scene):
    """PowerPoint has a physical slide-size limit; only scale final emit, not inference."""
    scale = min(1.0, 3200 / max(scene.image_size.width, scene.image_size.height))
    if scale == 1:
        return scene
    def box(b):
        return BBox(b.x0*scale,b.y0*scale,b.x1*scale,b.y1*scale)
    return replace(scene, image_size=ImageSize(round(scene.image_size.width*scale),round(scene.image_size.height*scale)),
        shapes=tuple(replace(s,bbox=box(s.bbox)) for s in scene.shapes),
        connectors=tuple(replace(c,path_points=tuple(Point(p.x*scale,p.y*scale) for p in c.path_points)) for c in scene.connectors))


def match_objects(predictions, targets, size):
    """Maximum-cardinality one-to-one matching. IoU orders ties, roles are measured separately."""
    target_boxes = [pixel_box(o["bbox"], size) for o in targets]
    candidates = []
    for p in predictions:
        b = p.bbox
        scores = [(iou((b.x0,b.y0,b.x1,b.y1),t),i) for i,t in enumerate(target_boxes)]
        candidates.append([i for score,i in sorted(scores,reverse=True) if score >= .5])
    target_to_pred = {}
    def augment(p, seen):
        for t in candidates[p]:
            if t in seen:
                continue
            seen.add(t)
            if t not in target_to_pred or augment(target_to_pred[t],seen):
                target_to_pred[t] = p
                return True
        return False
    for p in range(len(predictions)):
        augment(p,set())
    return {predictions[p].id: targets[t]["id"] for t,p in target_to_pred.items()}


def evaluate(scene, sample):
    targets = {o["id"]: o for o in sample["objects"]}
    matches = match_objects(scene.shapes,sample["objects"],sample["size"])
    roi = pixel_box(sample["roi"],sample["size"])
    exclusions = [pixel_box(b,sample["size"]) for b in sample["excluded"]]
    local_unmatched, excluded_hits, role_errors, wrong_labels, missing_labels, correct_labels = [],[],[],[],[],[]
    for shape in scene.shapes:
        b = shape.bbox
        box = (b.x0,b.y0,b.x1,b.y1)
        if shape.id not in matches and intersection(box,roi)>0:
            local_unmatched.append(shape.id)
        if any(intersection(box,e)/max(1,area(box))>=.5 for e in exclusions):
            excluded_hits.append(shape.id)
        if shape.id in matches:
            target = targets[matches[shape.id]]
            if shape.owner_kind.value != target["role"]:
                role_errors.append(target["id"])
            if target["label"] is not None:
                norm = lambda x: re.sub(r"\s+"," ",x or "").strip()
                if not shape.label:
                    missing_labels.append(target["id"])
                elif norm(shape.label) == norm(target["label"]):
                    correct_labels.append(target["id"])
                else:
                    wrong_labels.append({"id":target["id"],"predicted":shape.label,"expected":target["label"]})
    # Missing boxes also imply missing labels; never improve OCR score by dropping difficult boxes.
    missing_labels += [k for k,v in targets.items() if v["label"] is not None and k not in matches.values()]
    expected = set(map(tuple,sample["edges"]))
    recovered, wrong_edges = set(),[]
    for c in scene.connectors:
        pair = (matches.get(c.source_owner_id),matches.get(c.target_owner_id))
        if c.arrowhead_end and pair in expected:
            recovered.add(pair)
        else:
            # Full-CV output outside this partial annotation cannot be called a false positive.
            owners = [s for s in scene.shapes if s.id in (c.source_owner_id,c.target_owner_id)]
            if any(intersection((s.bbox.x0,s.bbox.y0,s.bbox.x1,s.bbox.y1),roi)>0 for s in owners):
                wrong_edges.append({"id":c.id,"mapped_pair":pair,"directed":c.arrowhead_end})
    return dict(target_objects=len(targets),matched_objects=len(matches),missing_objects=sorted(set(targets)-set(matches.values())),
        object_matches=matches,role_errors=role_errors,unmatched_local_objects=local_unmatched,excluded_region_hits=excluded_hits,
        target_edges=len(expected),matched_directed_edges=len(recovered),missing_edges=sorted(expected-recovered),
        wrong_or_unresolved_local_edges=wrong_edges,raw_emitted_edges=len(scene.connectors),
        correct_labels=correct_labels,missing_labels=missing_labels,wrong_labels=wrong_labels,
        total_emitted_objects=len(scene.shapes),unassessed_outside_roi=len(scene.shapes)-len(matches)-len(local_unmatched))


def inspect_pptx(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(path)
    shapes = presentation.slides[0].shapes
    connections = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
    return dict(native_shapes=len(shapes),raster_images=sum(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in shapes),
        connectors=len(connections),bound_endpoints=sum(len(s._element.xpath(".//a:stCxn | .//a:endCxn")) for s in connections),
        nonempty_text_shapes=sum(bool(s.text.strip()) for s in shapes if s.has_text_frame))


def draw_overlay(image,sample,path,scene=None):
    overlay=image.copy()
    draw=ImageDraw.Draw(overlay)
    size=image.size
    width=max(2,round(min(size)/250))
    draw.rectangle(pixel_box(sample["roi"],size),outline="blue",width=width)
    for obj in sample["objects"]:
        box=pixel_box(obj["bbox"],size)
        draw.rectangle(box,outline="limegreen",width=width)
        draw.text((box[0]+4,box[1]+4),obj["id"],fill="green",stroke_width=1,stroke_fill="white")
    if scene is not None:
        for shape in scene.shapes:
            b=shape.bbox
            draw.rectangle((b.x0,b.y0,b.x1,b.y1),outline="magenta",width=width)
        for edge in scene.connectors:
            draw.line([(p.x,p.y) for p in edge.path_points],fill="red",width=width)
    overlay.save(path)


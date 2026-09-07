"""R3-E1 bounded path approximation with one native preset connector."""
import math
import numpy as np
from pptx import Presentation
from pptx.oxml.ns import qn
from image_to_editable_ppt.research.native_emit import element
from image_to_editable_ppt.research.path_emit import write_path_sites
from image_to_editable_ppt.v3.emit.pptx_writer import EMU_PER_PIXEL


def samples(path):
    pieces=[]
    for a,b in zip(path,path[1:]):
        a,b=np.asarray(a,float),np.asarray(b,float)
        pieces.append(np.linspace(a,b,max(2,math.ceil(np.linalg.norm(b-a)/.25)+1)))
    return np.concatenate(pieces)


def distance(points,path):
    distances=[]
    for a,b in zip(path,path[1:]):
        a,b=np.asarray(a,float),np.asarray(b,float); v=b-a
        t=np.clip((points-a)@v/max(float(v@v),1e-12),0,1)
        distances.append(np.linalg.norm(points-(a+t[:,None]*v),axis=1))
    return float(np.min(distances,axis=0).max())


def fit_path(points):
    source=[(p.x,p.y) for p in points]
    a,b=source[0],source[-1]
    candidates=[('line',None,[a,b])]
    if abs(b[0]-a[0])>1e-6 and abs(b[1]-a[1])>1e-6:
        xs=[a[0],b[0],(a[0]+b[0])/2,*[p[0] for p in source[1:-1]]]
        for x in xs:
            if min(a[0],b[0])<=x<=max(a[0],b[0]):
                candidates.append(('bentConnector3',(x-a[0])/(b[0]-a[0]),[a,(x,a[1]),(x,b[1]),b]))
    scores=[]
    for preset,adj,route in candidates:
        error=max(distance(samples(source),route),distance(samples(route),source))
        scores.append((error,preset,adj,route))
    best=min(scores,key=lambda r:r[0])
    if best[0]>4: raise ValueError(f'no supported preset within 4px: {best[0]}')
    return dict(error_px=best[0],preset=best[1],adj=best[2],route_px=best[3])


def write_preset_paths(scene,path,*,styles=None):
    fits={c.id:fit_path(c.path_points) for c in scene.connectors}
    # Reuse only its owner sites/bindings; custom connector paths are replaced before delivery.
    write_path_sites(scene,path,styles=styles)
    deck=Presentation(path); objects={s.name:s for s in deck.slides[0].shapes}
    for c in scene.connectors:
        shape=objects[c.id]; props=shape._element.spPr; fit=fits[c.id]
        a,b=c.path_points[0],c.path_points[-1]
        xf=props.xfrm
        xf.off.x=round(min(a.x,b.x)*EMU_PER_PIXEL); xf.off.y=round(min(a.y,b.y)*EMU_PER_PIXEL)
        xf.ext.cx=max(1,round(abs(a.x-b.x)*EMU_PER_PIXEL)); xf.ext.cy=max(1,round(abs(a.y-b.y)*EMU_PER_PIXEL))
        for name,condition in [('flipH',a.x>b.x),('flipV',a.y>b.y)]:
            if condition: xf.set(name,'1')
            else: xf.attrib.pop(name,None)
        geom=element('prstGeom',prst=fit['preset']); av=element('avLst'); geom.append(av)
        if fit['adj'] is not None: av.append(element('gd',name='adj1',fmla=f"val {round(fit['adj']*100000)}"))
        props.replace(props.find(qn('a:custGeom')),geom)
    deck.save(path)
    return fits


def preset_vertices(shape):
    props=shape._element.spPr; xf=props.xfrm
    geom=props.find(qn('a:prstGeom'))
    if geom is None: raise ValueError('not a native preset connector')
    preset=geom.get('prst')
    adjustments={g.get('name'):float(g.get('fmla').split()[-1])/100000 for g in geom.find(qn('a:avLst'))}
    if preset in ('line','straightConnector1'): uv=[(0,0),(1,1)]
    elif preset=='bentConnector2': uv=[(0,0),(1,0),(1,1)]
    elif preset=='bentConnector3':
        a=adjustments.get('adj1',.5); uv=[(0,0),(a,0),(a,1),(1,1)]
    else: raise ValueError(f'unsupported preset after PowerPoint: {preset}')
    out=[]
    for u,v in uv:
        if xf.get('flipH') in ('1','true'): u=1-u
        if xf.get('flipV') in ('1','true'): v=1-v
        x,y=(u-.5)*xf.ext.cx,(v-.5)*xf.ext.cy
        angle=float(xf.get('rot','0'))/60000*math.pi/180
        x,y=x*math.cos(angle)-y*math.sin(angle),x*math.sin(angle)+y*math.cos(angle)
        out.append(((xf.off.x+.5*xf.ext.cx+x)/12700,(xf.off.y+.5*xf.ext.cy+y)/12700))
    return out

"""R16 semantic-region pilot, not PGDPNet/instance or relation reproduction."""
from pathlib import Path
import cv2
import numpy as np
import torch
from torch import nn
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from image_to_editable_ppt.ml.connector_segmenter import _UNet

PALETTE=np.array([[0,0,0],[255,0,0],[0,255,0],[0,0,255],
    [255,255,0],[255,0,255],[0,255,255]],np.uint8)
PRESETS=(None,MSO_SHAPE.RECTANGLE,MSO_SHAPE.ROUNDED_RECTANGLE,MSO_SHAPE.OVAL,
         MSO_SHAPE.DIAMOND,MSO_SHAPE.PARALLELOGRAM,MSO_SHAPE.CAN)
KINDS=('background','rectangle','rounded_rectangle','ellipse','diamond','parallelogram','cylinder')
SIZE=(240,152)


class PrimitiveUNet(_UNet):
    def __init__(self):
        super().__init__()
        self.head=nn.Conv2d(16,len(KINDS),1)


def palette_target(rgb):
    if rgb.dtype!=np.uint8 or rgb.ndim!=3 or rgb.shape[2]!=3:
        raise ValueError('uint8 RGB mask required')
    target=np.full(rgb.shape[:2],255,np.uint8)
    for i,color in enumerate(PALETTE):target[np.all(rgb==color,axis=2)]=i
    return target


def components(probabilities, image_size, threshold):
    p=np.asarray(probabilities)
    if p.ndim!=3 or p.shape[0]!=len(KINDS) or not np.isfinite(p).all() or np.any((p<0)|(p>1)):
        raise ValueError('finite seven-class probabilities required')
    if not 0<=threshold<=1 or min(image_size)<=0:raise ValueError('invalid decoding contract')
    cls=p.argmax(0);conf=p.max(0);h,w=cls.shape;result=[]
    # ponytail: disconnected semantic components only; touching/overlap needs instance embeddings.
    for kind in range(1,len(KINDS)):
        n,labels,stats,_=cv2.connectedComponentsWithStats(((cls==kind)&(conf>=threshold)).astype(np.uint8),connectivity=8)
        for i in range(1,n):
            x,y,bw,bh,area=map(int,stats[i])
            if area<16 or min(bw,bh)<4:continue
            result.append(dict(id=f'pred:{len(result)}',kind=kind,
                bbox=[x*image_size[0]/w,y*image_size[1]/h,(x+bw)*image_size[0]/w,(y+bh)*image_size[1]/h],
                confidence=float(conf[labels==i].mean())))
    return result


def object_counts(predicted, expected):
    from scipy.optimize import linear_sum_assignment
    from image_to_editable_ppt.research.closed_rect_proposals import _iou
    score=np.array([[_iou(a['bbox'],b['bbox']) if a['kind']==b['kind'] else 0.
        for b in expected] for a in predicted]).reshape(len(predicted),len(expected))
    if score.size:
        # Match cardinality first; an IoU-only optimum can discard a valid second match.
        rows,cols=linear_sum_assignment(-((score>=.5).astype(float)+score/(min(score.shape)+1)))
        tp=int(np.count_nonzero(score[rows,cols]>=.5))
    else:tp=0
    return dict(tp=tp,fp=len(predicted)-tp,fn=len(expected)-tp)


def select_threshold(checks):
    eligible=[float(t) for t,c in checks.items() if c['fp']==0 and c['tp']>0]
    return (max(eligible,key=lambda t:(checks[str(t)]['recall'],t)),True) if eligible else (.95,False)


def write_shapes(predicted, size, output):
    output=Path(output)
    if output.exists():raise FileExistsError(output)
    deck=Presentation();deck.slide_width=Emu(round(size[0]*9525));deck.slide_height=Emu(round(size[1]*9525))
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    for obj in predicted:
        kind=obj['kind'];x,y,right,bottom=obj['bbox']
        if kind not in range(1,len(PRESETS)) or not 0<=x<right<=size[0] or not 0<=y<bottom<=size[1]:
            raise ValueError('invalid predicted primitive')
        shape=slide.shapes.add_shape(PRESETS[kind],*[Emu(round(v*9525)) for v in (x,y,right-x,bottom-y)])
        shape.name=obj['id'];shape.fill.background();shape.line.color.rgb=RGBColor(40,50,65)
        shape.line.width=Pt(1);shape.shadow.inherit=False
    deck.save(output)

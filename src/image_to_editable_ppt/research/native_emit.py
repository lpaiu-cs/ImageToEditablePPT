"""R1 native emitter experiment. Explicit sites, no anchors, no production patch."""
from __future__ import annotations

from dataclasses import replace
import json
from pathlib import Path

from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

from image_to_editable_ppt.v3.core.enums import ConnectorKind,ContainerKind,NodeKind,PortOwnerKind
from image_to_editable_ppt.v3.core.types import BBox,ImageSize,Point
from image_to_editable_ppt.v3.emit.models import EmitScene,EmitShapePrimitive,EmitConnectorPrimitive
from image_to_editable_ppt.v3.emit.pptx_writer import EMU_PER_PIXEL,write_pptx


def load_scene(path: Path) -> EmitScene:
    raw=json.loads(path.read_text(encoding="utf-8"))
    if raw["texts"] or raw["residuals"]:
        raise ValueError("R1 fixed scenes must have only shapes and connectors")
    shapes=[]
    for item in raw["shapes"]:
        obj=dict(item)
        obj["owner_kind"]=PortOwnerKind(obj["owner_kind"])
        enum=NodeKind if obj["owner_kind"] is PortOwnerKind.NODE else ContainerKind
        obj["shape_kind"]=enum(obj["shape_kind"])
        obj["bbox"]=BBox(**obj["bbox"])
        obj["provenance"]=tuple(obj["provenance"])
        shapes.append(EmitShapePrimitive(**obj))
    connectors=[]
    for item in raw["connectors"]:
        obj=dict(item)
        obj["kind"]=ConnectorKind(obj["kind"])
        for key in ("source_owner_kind","target_owner_kind"):
            obj[key]=PortOwnerKind(obj[key])
        obj["path_points"]=tuple(Point(**p) for p in obj["path_points"])
        obj["provenance"]=tuple(obj["provenance"])
        connectors.append(EmitConnectorPrimitive(**obj))
    return EmitScene(image_size=ImageSize(**raw["image_size"]),shapes=tuple(shapes),
        connectors=tuple(connectors),source=raw["source"],provenance=tuple(raw["provenance"]))


def normalize_scene(scene: EmitScene,width: int=768) -> EmitScene:
    scale=width/scene.image_size.width
    def box(b):
        return BBox(b.x0*scale,b.y0*scale,b.x1*scale,b.y1*scale)
    return replace(scene,image_size=ImageSize(width,max(1,round(scene.image_size.height*scale))),
        shapes=tuple(replace(s,bbox=box(s.bbox)) for s in scene.shapes),
        connectors=tuple(replace(c,path_points=tuple(Point(p.x*scale,p.y*scale) for p in c.path_points)) for c in scene.connectors))


def element(tag,**attrs):
    obj=OxmlElement("a:"+tag)
    for k,v in attrs.items():
        obj.set(k,str(v))
    return obj


def connection_sites(scene):
    """Connection coordinates come only from fixed observed endpoints, never target answers."""
    owners={s.id:s for s in scene.shapes}
    if len(owners)!=len(scene.shapes):
        raise ValueError("duplicate owner IDs")
    sites={}
    binding={}
    for c in scene.connectors:
        if len(c.path_points)!=2:
            raise ValueError("R1 only supports two-point observed paths")
        start,end=c.path_points
        if abs(start.x-end.x)>1e-5 and abs(start.y-end.y)>1e-5:
            raise ValueError("R1 only supports observed horizontal/vertical paths")
        for endpoint,owner_id,point in (("start",c.source_owner_id,start),("end",c.target_owner_id,end)):
            if owner_id not in owners:
                raise ValueError("unknown endpoint owner")
            shape=owners[owner_id]
            b=shape.bbox
            u,v=(point.x-b.x0)/b.width,(point.y-b.y0)/b.height
            if not (-1e-6<=u<=1+1e-6 and -1e-6<=v<=1+1e-6):
                raise ValueError(f"endpoint outside owner: {owner_id}")
            side=min(((abs(v),270),(abs(u),180),(abs(1-v),90),(abs(1-u),0)))[1]
            site=(round(max(0,min(1,u)),6),round(max(0,min(1,v)),6),side*60000)
            current=sites.setdefault(owner_id,[(.5,0,16200000),(0,.5,10800000),(.5,1,5400000),(1,.5,0)])
            if site not in current:
                current.append(site)
            binding[(c.id,endpoint)]=current.index(site)
    return sites,binding


def custom_geometry(primitive,sites):
    """One simple closed path plus normalized connection guides; no auxiliary shapes."""
    supported={NodeKind.BOX,NodeKind.ROUNDED_BOX,NodeKind.ELLIPSE,NodeKind.SECTION,ContainerKind.FLOW_CLUSTER,ContainerKind.PANEL}
    if primitive.shape_kind not in supported:
        raise ValueError("unsupported custom-geometry shape")
    geom=element("custGeom")
    geom.append(element("avLst"))
    guides=element("gdLst")
    geom.append(guides)
    geom.append(element("ahLst"))
    connections=element("cxnLst")
    geom.append(connections)
    for i,(u,v,angle) in enumerate(sites):
        guides.append(element("gd",name=f"r1x{i}",fmla=f"*/ w {round(u*1000000)} 1000000"))
        guides.append(element("gd",name=f"r1y{i}",fmla=f"*/ h {round(v*1000000)} 1000000"))
        site=element("cxn",ang=angle)
        site.append(element("pos",x=f"r1x{i}",y=f"r1y{i}"))
        connections.append(site)
    geom.append(element("rect",l="0",t="0",r="w",b="h"))
    paths=element("pathLst")
    path=element("path",w=100000,h=100000)
    paths.append(path)
    geom.append(paths)
    def command(name,*points):
        cmd=element(name)
        for x,y in points:
            cmd.append(element("pt",x=round(x),y=round(y)))
        path.append(cmd)
    if primitive.shape_kind is NodeKind.ELLIPSE:
        k=27614.2375
        command("moveTo",(50000,0))
        command("cubicBezTo",(50000+k,0),(100000,50000-k),(100000,50000))
        command("cubicBezTo",(100000,50000+k),(50000+k,100000),(50000,100000))
        command("cubicBezTo",(50000-k,100000),(0,50000+k),(0,50000))
        command("cubicBezTo",(0,50000-k),(50000-k,0),(50000,0))
    elif primitive.shape_kind is NodeKind.ROUNDED_BOX:
        radius=.12*min(primitive.bbox.width,primitive.bbox.height)
        rx=radius/primitive.bbox.width*100000
        ry=radius/primitive.bbox.height*100000
        k=.55228475
        command("moveTo",(rx,0))
        command("lnTo",(100000-rx,0))
        command("cubicBezTo",(100000-rx+k*rx,0),(100000,ry-k*ry),(100000,ry))
        command("lnTo",(100000,100000-ry))
        command("cubicBezTo",(100000,100000-ry+k*ry),(100000-rx+k*rx,100000),(100000-rx,100000))
        command("lnTo",(rx,100000))
        command("cubicBezTo",(rx-k*rx,100000),(0,100000-ry+k*ry),(0,100000-ry))
        command("lnTo",(0,ry))
        command("cubicBezTo",(0,ry-k*ry),(rx-k*rx,0),(rx,0))
    else:
        command("moveTo",(0,0))
        for p in ((100000,0),(100000,100000),(0,100000)):
            command("lnTo",p)
    command("close")
    return geom


def write_native_sites(scene,path,*,styles=None):
    sites,bindings=connection_sites(scene)
    # Preflight all custom paths before any write, including unsupported shape kinds.
    geometries={s.id:custom_geometry(s,sites[s.id]) for s in scene.shapes if s.id in sites}
    write_pptx(scene,path,styles=styles)
    presentation=Presentation(path)
    shapes={s.name:s for s in presentation.slides[0].shapes}
    for owner,geom in geometries.items():
        sp_pr=shapes[owner]._element.spPr
        old=sp_pr.find(qn("a:prstGeom"))
        if old is None:
            raise ValueError("expected preset shape before custom-site conversion")
        sp_pr.replace(old,geom)
    for connector in scene.connectors:
        shape=shapes[connector.id]
        cnv=shape._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
        for name in ("stCxn","endCxn"):
            old=cnv.find(qn("a:"+name))
            if old is not None:
                cnv.remove(old)
        cnv.append(element("stCxn",id=shapes[connector.source_owner_id].shape_id,idx=bindings[(connector.id,"start")]))
        cnv.append(element("endCxn",id=shapes[connector.target_owner_id].shape_id,idx=bindings[(connector.id,"end")]))
        # Do not use python-pptx begin_connect: its rect-side snap does not know custom indices.
        a,b=connector.path_points
        xfrm=shape._element.spPr.xfrm
        xfrm.off.x=round(min(a.x,b.x)*EMU_PER_PIXEL)
        xfrm.off.y=round(min(a.y,b.y)*EMU_PER_PIXEL)
        xfrm.ext.cx=round(abs(a.x-b.x)*EMU_PER_PIXEL)
        xfrm.ext.cy=round(abs(a.y-b.y)*EMU_PER_PIXEL)
        for key,needed in (("flipH",a.x>b.x),("flipV",a.y>b.y)):
            if needed:
                xfrm.set(key,"1")
            else:
                xfrm.attrib.pop(key,None)
        shape._element.spPr.find(qn("a:prstGeom")).set("prst","line")
    presentation.save(path)
    return {owner:[dict(u=u,v=v,angle=angle) for u,v,angle in points] for owner,points in sites.items()}


def endpoint_audit(path,scene):
    presentation=Presentation(path)
    shapes={s.name:s for s in presentation.slides[0].shapes}
    records=[]
    for c in scene.connectors:
        obj=shapes[c.id]
        p,q=c.path_points
        actual=[obj.begin_x/EMU_PER_PIXEL,obj.begin_y/EMU_PER_PIXEL,obj.end_x/EMU_PER_PIXEL,obj.end_y/EMU_PER_PIXEL]
        wanted=[p.x,p.y,q.x,q.y]
        records.append(dict(id=c.id,max_coordinate_error_px=max(abs(a-b) for a,b in zip(actual,wanted)),
            actual=actual,expected=wanted,bound_endpoints=len(obj._element.xpath(".//a:stCxn | .//a:endCxn"))))
    return records

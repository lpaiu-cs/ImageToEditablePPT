"""Deterministic artifact fingerprints for the isolated Phase 16 experiment."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .validation import sha256_file


def _enum_value(value: object) -> object:
    return getattr(value, "value", value)


def _connection_id(shape: object, tag: str) -> str | None:
    elements = shape._element.xpath(f".//a:{tag}")  # type: ignore[attr-defined]
    return elements[0].get("id") if elements else None


def _arrowhead_type(shape: object, tag: str) -> str:
    elements = shape._element.xpath(f".//a:{tag}")  # type: ignore[attr-defined]
    return str(elements[0].get("type", "none")) if elements else "none"


def canonical_pptx_semantics(path: Path) -> dict[str, object]:
    """Return editability-relevant PPTX semantics without package timestamps."""

    presentation = Presentation(str(path))
    slides: list[dict[str, object]] = []
    for slide_index, slide in enumerate(presentation.slides):
        shapes: list[dict[str, object]] = []
        for z_order, shape in enumerate(slide.shapes):
            record: dict[str, object] = {
                "z_order": z_order,
                "shape_id": int(shape.shape_id),
                "name": str(shape.name),
                "shape_type": _enum_value(shape.shape_type),
                "bbox_emu": [int(shape.left), int(shape.top), int(shape.width), int(shape.height)],
                "rotation": round(float(shape.rotation or 0.0), 6),
                "text": str(shape.text) if getattr(shape, "has_text_frame", False) else None,
            }
            try:
                auto_shape_type = shape.auto_shape_type
            except (AttributeError, ValueError):
                auto_shape_type = None
            if auto_shape_type is not None:
                record["auto_shape_type"] = _enum_value(auto_shape_type)
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                record["connections"] = {
                    "start_shape_id": _connection_id(shape, "stCxn"),
                    "end_shape_id": _connection_id(shape, "endCxn"),
                    "head": _arrowhead_type(shape, "headEnd"),
                    "tail": _arrowhead_type(shape, "tailEnd"),
                }
            shapes.append(record)
        slides.append({"slide_index": slide_index, "shapes": shapes})
    return {
        "slide_size_emu": [int(presentation.slide_width), int(presentation.slide_height)],
        "slides": slides,
    }


def pptx_semantic_digest(path: Path) -> str:
    payload = json.dumps(
        canonical_pptx_semantics(path),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def pptx_shape_counts(path: Path) -> dict[str, int | bool]:
    presentation = Presentation(str(path))
    if len(presentation.slides) != 1:
        raise ValueError(f"{path} must contain exactly one slide")
    shapes = list(presentation.slides[0].shapes)
    line_shapes = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    nodes = [
        shape
        for shape in shapes
        if shape.shape_type != MSO_SHAPE_TYPE.LINE
        and (shape.name.startswith("node:") or ":node:" in shape.name)
    ]
    containers = [
        shape
        for shape in shapes
        if shape.shape_type != MSO_SHAPE_TYPE.LINE
        and (shape.name.startswith("container:") or ":container:" in shape.name)
    ]
    textboxes = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    semantic_connector_names = {shape.name for shape in line_shapes}
    return {
        "total": len(shapes),
        "nodes": len(nodes),
        "containers": len(containers),
        "connector_segments": len(line_shapes),
        "semantic_connectors": len(semantic_connector_names),
        "textboxes": len(textboxes),
        "objects_with_text": sum(
            bool(getattr(shape, "has_text_frame", False) and shape.text.strip()) for shape in shapes
        ),
        "blank": len(shapes) == 0,
    }


def bound_file_set_digest(paths: Iterable[Path], *, root: Path) -> str:
    """Bind a set of files by repository-relative path and content digest."""

    resolved_root = root.resolve()
    records: list[tuple[str, str]] = []
    for path in paths:
        resolved = path.resolve()
        relative = resolved.relative_to(resolved_root).as_posix()
        records.append((relative, sha256_file(resolved)))
    digest = hashlib.sha256()
    for relative, file_digest in sorted(records):
        digest.update(relative.encode("utf-8"))
        digest.update(b"\0")
        digest.update(file_digest.encode("ascii"))
        digest.update(b"\n")
    return digest.hexdigest()

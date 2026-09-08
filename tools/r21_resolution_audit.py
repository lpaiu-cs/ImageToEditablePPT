"""Posthoc manual matching and native COM audit of the frozen resolution probe."""
import argparse
import json
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from r19_outcome_audit import ROOT, digest, save, verify, geometry, object_counts


def run(output):
    output.mkdir(parents=True, exist_ok=False)
    base = ROOT / 'artifacts/region-image-only-r21'
    result = json.loads((base / 'original-v1/results.json').read_text())
    verify(result)
    original = json.loads((ROOT / 'artifacts/region-image-only-r19/original-v1/result.json').read_text())
    for key in ('objects', 'all_candidates'):
        assert result['arms']['full600'][key] == original[key]
    manifest = base / 'com-v1/renders/powerpoint_render_manifest.json'
    receipts = {r['id']: r for r in json.loads(manifest.read_text(encoding='utf-8-sig'))['records']}
    assert set(receipts) == set(result['arms'])
    ledger_path = ROOT / 'docs/r15_visual_loss_ledger.json'
    ledger = json.loads(ledger_path.read_text())
    source = ROOT / 'artifacts/region-image-only-r10/complex-inputs/original-rgb.png'
    with Image.open(source) as im:
        source_image = im.convert('RGB')
    assert digest(source) == result['input_sha256']
    expected = [dict(kind=6 if n.get('kind') == 'cylinder' else 2,
        bbox=[v * source_image.size[i % 2] / ledger['source_display_size'][i % 2]
              for i, v in enumerate(n['box'])]) for n in ledger['nodes']]
    board = Image.new('RGB', (1380, 812), 'white')
    draw = ImageDraw.Draw(board)
    board.paste(source_image.resize((690, 376)), (0, 30))
    draw.text((4, 4), 'Exposed source RGB', fill='black')
    counts = {}
    for i, (arm, row) in enumerate(result['arms'].items(), 1):
        receipt = receipts[arm]
        assert receipt['pptx_sha256'] == row['pptx_sha256'] == digest(receipt['pptx_path'])
        assert receipt['render_sha256'] == digest(receipt['render_path'])
        slide = Presentation(receipt['pptx_path']).slides[0]
        assert len(slide.shapes) == len(row['objects'])
        assert all(s.shape_type != 13 for s in slide.shapes)
        counts[arm] = dict(kind=object_counts(row['objects'], expected),
            geometry=geometry(row['objects'], expected), candidates=geometry(row['all_candidates'], expected))
        x, y = i % 2 * 690, i // 2 * 406
        with Image.open(receipt['render_path']) as im:
            assert im.size == (1380, 752)
            board.paste(im.convert('RGB').resize((690, 376)), (x, y + 30))
        c = counts[arm]['kind']
        draw.text((x + 4, y + 4), f'{arm}: manual TP{c["tp"]} FP{c["fp"]} FN{c["fn"]}', fill='black')
    board.save(output / 'comparison.png')
    save(output / 'audit.json', dict(counts=counts, baseline_exact=True,
        result_sha256=digest(base / 'original-v1/results.json'), receipts_sha256=digest(manifest),
        ledger_sha256=digest(ledger_path), code_sha256=digest(__file__),
        scope='Posthoc exposed 17-node manual ledger; not independent or exhaustive structure GT.'))
    print(json.dumps(counts, indent=2))


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('--output', type=Path, required=True)
    run(p.parse_args().output)

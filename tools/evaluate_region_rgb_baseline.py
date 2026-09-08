"""별도 GT 평가 및 실제 native 렌더 검토. 이 파일은 추론에서 import하지 않는다."""
import argparse
from collections import Counter
import json
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from run_region_rgb_baseline import digest, save
from image_to_editable_ppt.region.geometry import intersection, area


def box(row):
    return [row['bbox'][k] for k in ('x0', 'y0', 'x1', 'y1')]


def match(shapes, truth, *, exact_kind=False):
    pairs = []
    for i, s in enumerate(shapes):
        for j, t in enumerate(truth):
            if s['owner_kind'] != t['owner_kind'] or (exact_kind and s['shape_kind'] != t['kind']):
                continue
            a, b = box(s), box(t)
            overlap = intersection(a, b)
            score = overlap / max(1, area(a)+area(b)-overlap)
            if score >= .5:
                pairs.append((score, i, j))
    used, result = set(), {}
    for _, i, j in sorted(pairs, reverse=True):
        if shapes[i]['id'] not in result and j not in used:
            result[shapes[i]['id']] = truth[j]['id']
            used.add(j)
    return result


def counts(tp, pred, ref):
    return dict(tp=tp, fp=pred-tp, fn=ref-tp, predictions=pred, references=ref)


def signature(a, b, start, end):
    if start != end:
        return ('directed', b, a) if start else ('directed', a, b)
    return ('bidirectional' if start else 'undirected', *sorted((a, b)))


def edges(scene, mapping):
    return Counter(signature(mapping.get(c['source_owner_id'], 'unmatched:'+c['source_owner_id']),
                             mapping.get(c['target_owner_id'], 'unmatched:'+c['target_owner_id']),
                             c['arrowhead_start'], c['arrowhead_end']) for c in scene['connectors'])


def main(args):
    rows = {r['id']: r for r in json.loads(args.evaluator.read_text(encoding='utf-8'))}
    results = json.loads((args.run / 'results.json').read_text(encoding='utf-8'))
    assert set(rows) == {r['id'] for r in results['records']}
    native = json.loads((args.run / 'renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
    renders = {r['id']: r for r in native['records']}
    unmonitored = getattr(args,'allow_unmonitored_replay',False)
    if unmonitored:
        assert results['mode']=='cached B/G refinement diagnostic'
        assert all(r['io']=={'mode':'cached scene diagnostic; no OS/read sandbox claim'} for r in results['records'])
    review = args.run / getattr(args,'review_name','review')
    review.mkdir(exist_ok=False)
    reports = []
    font = ImageFont.truetype('C:/Windows/Fonts/malgun.ttf', 18)
    sheets = []
    for record in results['records']:
        row = rows[record['id']]
        assert record['returncode'] == 0
        if not unmonitored:assert not record['io']['denied_paths']
        assert row['image_sha256'] == record['source_sha256']
        assert digest(row['annotation']) == row['annotation_sha256']
        reference = json.loads(args.evaluator.parent.joinpath(row['annotation']).read_text(encoding='utf-8'))
        gt = reference['primitive_scene']
        truth = [dict(t, owner_kind=role) for role, key in [('node', 'nodes'), ('container', 'containers')] for t in gt[key]]
        scene = record['report']['source_scene']
        assert [scene['image_size']['width'], scene['image_size']['height']] == [reference['image_size']['width'], reference['image_size']['height']]
        mapping = match(scene['shapes'], truth)
        strict = match(scene['shapes'], truth, exact_kind=True)
        metrics = {}
        for role in ('node', 'container'):
            pred = [s for s in scene['shapes'] if s['owner_kind'] == role]
            ref = [t for t in truth if t['owner_kind'] == role]
            metrics[role] = counts(sum(s['id'] in mapping for s in pred), len(pred), len(ref))
            metrics[role+'_exact_kind'] = counts(sum(s['id'] in strict for s in pred), len(pred), len(ref))
        truth_by_id = {t['id']: t for t in truth}
        labels = []
        for s in scene['shapes']:
            if s['owner_kind'] == 'node' and s['id'] in mapping:
                t = truth_by_id[mapping[s['id']]]
                norm = lambda value: ' '.join((value or '').split()).casefold()
                labels.append(dict(id=s['id'], actual=s['label'], expected=t.get('label'),
                                   correct=bool(s['label']) and norm(s['label']) == norm(t.get('label'))))
        expected_edges = Counter(signature(c['start_endpoint']['owner_id'], c['end_endpoint']['owner_id'],
                                           c['arrowhead_start'], c['arrowhead_end']) for c in gt['connector_candidates'])
        for label, predicted_scene in [('traced_edges', scene), ('native_edges', record['report']['emitted_scene'])]:
            predicted = edges(predicted_scene, mapping)
            metrics[label] = counts(sum((predicted & expected_edges).values()), sum(predicted.values()), sum(expected_edges.values()))
        pp = args.run / 'pptx' / (record['slot']+'.pptx')
        nr = renders[record['slot']]
        assert digest(pp) == record['pptx_sha256'] == nr['pptx_sha256']
        assert digest(nr['render_path']) == nr['render_sha256']
        assert digest(pp.with_suffix('.region.json')) == record['report_sha256']
        deck = Presentation(pp)
        objects = {s.name: s for s in deck.slides[0].shapes}
        ids = {str(s.shape_id): s.name for s in objects.values()}
        emitted = record['report']['emitted_scene']
        assert set(objects) == {s['id'] for s in emitted['shapes']} | {c['id'] for c in emitted['connectors']}
        assert not deck.slides[0]._element.xpath('.//p:pic')
        for c in emitted['connectors']:
            assert [ids[b.get('id')] for b in objects[c['id']]._element.xpath('.//a:stCxn | .//a:endCxn')] == [c['source_owner_id'], c['target_owner_id']]
        im = Image.open(args.run / 'rgb' / (record['slot']+'.png')).convert('RGB')
        assert digest(args.run / 'rgb' / (record['slot']+'.png')) == row['image_sha256']
        rendered = Image.open(nr['render_path']).convert('RGB')
        panel = Image.new('RGB', (1320, 414), '#edf1f5')
        draw = ImageDraw.Draw(panel)
        for i, (title, picture) in enumerate([('원본 RGB '+record['id'], im), (getattr(args, 'render_label', 'R6 이미지 단독 → 실제 PPT 렌더'), rendered)]):
            draw.text((i*660+10, 8), title, font=font, fill='black')
            panel.paste(ImageOps.contain(picture, (640, 360)), (i*660+10, 42))
        panel.save(review / (record['slot']+'.png'))
        sheets.append(panel)
        reports.append(dict(id=record['id'], slot=record['slot'], metrics=metrics, labels=labels,
                            source_sha256=row['image_sha256'], annotation_sha256=row['annotation_sha256'],
                            pptx_sha256=record['pptx_sha256'], render_sha256=nr['render_sha256']))
    for offset in range(0, len(sheets), 5):
        sheet = Image.new('RGB', (1320, 414*len(sheets[offset:offset+5])), 'white')
        for i, panel in enumerate(sheets[offset:offset+5]): sheet.paste(panel, (0, 414*i))
        sheet.save(review / f'contact-{offset//5}.png')
    totals = {k: {c: sum(r['metrics'][k][c] for r in reports) for c in ('tp', 'fp', 'fn', 'predictions', 'references')}
              for k in reports[0]['metrics']}
    labels = [l for r in reports for l in r['labels']]
    summary = dict(population='shared synthetic validation diagnostic; not unseen acceptance',
                   input_boundary='unmonitored cached scene replay, NOT fresh/guarded inference' if unmonitored else 'recorded Python read audit; not OS sandbox',
                   results_sha256=digest(args.run / 'results.json'), evaluator_sha256=digest(args.evaluator),
                   matching='IoU >= 0.5; emitted objects without an additional detector NMS; kind-specific and role-only separated',
                   totals=totals, labels=dict(matched_nodes=len(labels), correct=sum(l['correct'] for l in labels),
                   wrong=sum(bool(l['actual']) and not l['correct'] for l in labels), blank=sum(not l['actual'] for l in labels)), records=reports)
    save(args.run / getattr(args,'evaluation_name','evaluation.json'), summary)
    print(json.dumps(dict(totals=totals, labels=summary['labels']), ensure_ascii=False))


if __name__ == '__main__':
    from pathlib import Path
    parser = argparse.ArgumentParser()
    parser.add_argument('--run', type=Path, required=True)
    parser.add_argument('--evaluator', type=Path, required=True)
    parser.add_argument('--render-label', default='R6 이미지 단독 → 실제 PPT 렌더')
    parser.add_argument('--allow-unmonitored-replay',action='store_true')
    parser.add_argument('--review-name',default='review')
    parser.add_argument('--evaluation-name',default='evaluation.json')
    main(parser.parse_args())

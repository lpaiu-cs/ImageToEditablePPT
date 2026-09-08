"""Frozen v0.2 source-to-native complex-image validation; no tuning or holdout."""
import argparse
import json
from pathlib import Path
import subprocess
import sys
import time
from PIL import Image,ImageDraw,ImageFont,ImageOps
from pptx import Presentation
from image_to_editable_ppt.research.region_first import digest,pixel_box

CASES=[('reasoning-full','phase16-dev-p034','Dense fork/join with mixed direct/curved relations'),
    ('blue-flow','phase16-dev-p042','Blue orthogonal flow and shared branches'),
    ('nested-panels','phase16-dev-p058','Nested colored containers, many text boxes and icon exclusions'),
    ('contrastive','phase16-dev-p043','Multi-branch encoder/decoder and dotted/curved relations'),
    ('dark-transformer','phase16-dev-p011','Black background, nested stack and skip connections'),
    ('fusion','phase16-dev-p092','Nested branches, thin slanted links and fusion node')]


def freeze(root,path):
    manifest=json.loads((root/'benchmarks/phase16_dev_v1/manifest.json').read_text())
    samples={s['id']:s for s in manifest['samples']}
    items=[]
    for name,key,reason in CASES:
        s=samples[key]; source=root/s['path']; assert digest(source)==s['sha256']
        items.append(dict(id=name,path=s['path'],sha256=s['sha256'],roi=[0,0,1,1],reason=reason,population='existing development split, not holdout; prior baseline exposure'))
    for name,roi,reason in [('original-full',[0,0,1,1],'Full original complex input, including resource-bound behavior'),
        ('original-body',[0,.04,1,1],'Original body ROI excluding top 4 percent; clipped header is out of scope, not full-image success')]:
        items.append(dict(id=name,path='input.png',sha256=digest(root/'input.png'),roi=roi,reason=reason,population='previously consumed original project input'))
    for item in items:
        with Image.open(root/item['path']) as im: item['size']=list(im.size)
    data=dict(runtime_commit=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip(),
        runtime_sha256={str(p.relative_to(root)):digest(p) for p in (root/'src/image_to_editable_ppt/region').glob('*.py')},
        cli_sha256=digest(root/'src/image_to_editable_ppt/ml/convert_to_pptx.py'),
        rules='No runtime changes, resizing, retry crops or cutoff adjustment after results. No automatic full-image or human-time acceptance claim.',samples=items)
    with path.open('x',encoding='utf-8') as f: json.dump(data,f,indent=2)


def run(root,manifest,out):
    data=json.loads(manifest.read_text()); out.mkdir(parents=True,exist_ok=False)
    assert all(digest(root/p)==sha for p,sha in data['runtime_sha256'].items())
    assert digest(root/'src/image_to_editable_ppt/ml/convert_to_pptx.py')==data['cli_sha256']
    results=dict(manifest_sha256=digest(manifest),records=[])
    for item in data['samples']:
        source=root/item['path']; assert digest(source)==item['sha256']
        pp=out/f"{item['id']}.pptx"
        command=[sys.executable,'-m','image_to_editable_ppt.ml.convert_to_pptx',str(source),'-o',str(pp),'--region',*map(str,item['roi'])]
        started=time.monotonic()
        try:
            r=subprocess.run(command,cwd=root,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=180)
            record=dict(id=item['id'],elapsed_s=time.monotonic()-started,returncode=r.returncode,stdout=r.stdout,stderr=r.stderr)
        except subprocess.TimeoutExpired:
            record=dict(id=item['id'],elapsed_s=time.monotonic()-started,returncode=None,status='timeout')
        if record['returncode']==0:
            rp=pp.with_suffix('.region.json'); report=json.loads(rp.read_text(encoding='utf-8'))
            record.update(report=report,pptx_sha256=digest(pp),report_sha256=digest(rp))
            print(item['id'],report['counts'],flush=True)
        else: print(item['id'],'REJECTED',record.get('stderr',record.get('status')),flush=True)
        results['records'].append(record)
        (out/'results.json').write_text(json.dumps(results,indent=2),encoding='utf-8')


def review(root,manifest,run_dir,renders,out,*,native_label='Native PowerPoint v0.2'):
    out.mkdir(parents=True,exist_ok=False)
    data=json.loads(manifest.read_text()); result=json.loads((run_dir/'results.json').read_text())
    assert digest(manifest)==result['manifest_sha256']
    render=json.loads((renders/'powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
    for r in render['records']:
        assert digest(Path(r['pptx_path']))==r['pptx_sha256'] and digest(Path(r['render_path']))==r['render_sha256']
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',20)
    sheets=[]; audits=[]
    for item,r in zip(data['samples'],result['records']):
        assert item['id']==r['id'] and digest(root/item['path'])==item['sha256']
        if r['returncode']!=0: continue
        pp=run_dir/f"{item['id']}.pptx"; rp=pp.with_suffix('.region.json')
        assert digest(pp)==r['pptx_sha256'] and digest(rp)==r['report_sha256']
        deck=Presentation(pp); report=r['report']; objects={s.name:s for s in deck.slides[0].shapes}
        scene=report['emitted_scene']; ids={str(s.shape_id):s.name for s in objects.values()}
        bindings=all([ids[b.get('id')] for b in objects[c['id']]._element.xpath('.//a:stCxn | .//a:endCxn')]==[c['source_owner_id'],c['target_owner_id']] for c in scene['connectors'])
        check=dict(id=item['id'],ids_exact=set(objects)=={s['id'] for s in scene['shapes']}|{c['id'] for c in scene['connectors']},
            owner_bindings=bindings,raster_count=len(deck.slides[0]._element.xpath('.//p:pic')),
            ambiguous_components=sum(c.get('rejected')=='ambiguous_multi_owner' for c in report['tracing']['components']))
        assert check['ids_exact'] and check['owner_bindings'] and not check['raster_count']
        audits.append(check)
        panels=[]
        for source,label in [(root/item['path'],'Source (fixed ROI)'),(renders/f"{item['id']}.png",native_label)]:
            with Image.open(source) as im: crop=im.convert('RGB').crop(pixel_box(item['roi'],im.size))
            thumb=ImageOps.contain(crop,(780,800)); panel=Image.new('RGB',(800,850),'#eef2f6')
            panel.paste(thumb,((800-thumb.width)//2,45+(800-thumb.height)//2))
            ImageDraw.Draw(panel).text((10,10),label,font=font,fill='black'); panels.append(panel)
        sheet=Image.new('RGB',(1600,850),'white'); sheet.paste(panels[0],(0,0)); sheet.paste(panels[1],(800,0))
        path=out/f"{item['id']}-comparison.png"; sheet.save(path)
        sheets.append(dict(id=item['id'],path=str(path.relative_to(root)),sha256=digest(path)))
    evidence=dict(manifest=data,results=result,render_manifest=render,audits=audits,sheets=sheets,
        code_sha256=digest(Path(__file__)),result_sha256=digest(run_dir/'results.json'))
    (out/'evidence.json').write_text(json.dumps(evidence,indent=2),encoding='utf-8')
    print('native audits',audits)


def summarize(root):
    import cv2
    import numpy as np
    base=root/'artifacts/region-complex-v1'
    evidence=json.loads((base/'review/evidence.json').read_text())
    rows=[]
    for r in evidence['results']['records']:
        row={k:v for k,v in r.items() if k!='report'}
        if 'report' in r:
            report=r['report']
            row['counts']=report['counts']
            row['emitted_scene']=report['emitted_scene']
            row['rejected_connectors']=report['emission']['rejected_connectors']
            row['tracing']=report['tracing']
        rows.append(row)
    with Image.open(root/'data/phase16_dev_v1/phase16-dev-p042.png') as im: rgb=np.asarray(im.convert('RGB'))
    selected=(rgb[:,:,2].astype(int)-rgb[:,:,0]>70)&(rgb[:,:,1].astype(int)-rgb[:,:,0]>30)&(rgb[:,:,2]>160)
    gray=cv2.cvtColor(rgb,cv2.COLOR_RGB2GRAY)
    blue=dict(selector='B-R>70 and G-R>30 and B>160; descriptive source-pixel diagnostic, not annotated edge mask',
        count=int(selected.sum()),median_gray=float(np.median(gray[selected])),fraction_below_110=float((gray[selected]<110).mean()))
    result=dict(manifest=evidence['manifest'],records=rows,audits=evidence['audits'],sheets=evidence['sheets'],
        renders=evidence['render_manifest'],blue_pixel_diagnostic=blue,
        visual_review=json.loads((root/'benchmarks/region_complex_visual_review_v1.json').read_text()),
        provenance={str(p.relative_to(root)):digest(p) for p in [base/'run1/results.json',base/'review/evidence.json',root/'benchmarks/region_complex_inputs_v1.json',Path(__file__)]})
    with (root/'benchmarks/region_complex_evidence_v1.json').open('x',encoding='utf-8') as f: json.dump(result,f,indent=2)


if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('mode',choices=['freeze','run','review','summarize'])
    p.add_argument('--output-dir',default='artifacts/region-complex-v1/run1'); args=p.parse_args()
    root=Path(__file__).resolve().parents[1]; manifest=root/'benchmarks/region_complex_inputs_v1.json'
    if args.mode=='freeze': freeze(root,manifest)
    elif args.mode=='run': run(root,manifest,root/args.output_dir)
    elif args.mode=='review': review(root,manifest,root/'artifacts/region-complex-v1/run1',root/'artifacts/region-complex-v1/renders',root/args.output_dir)
    else: summarize(root)

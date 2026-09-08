"""Stage immutable evaluation PPTs for existing COM renderer, then audit actual renders."""
import argparse
import json
from pathlib import Path
import shutil
from PIL import Image, ImageDraw
from pptx import Presentation
from run_region_rgb_baseline import digest, save


def run(manifest, evaluation, output, mode):
    data=json.loads(manifest.read_text());records={r['id']:r for r in data['records'] if r['split']=='eval'}
    results=json.loads((evaluation/'results.json').read_text())
    assert results['manifest_sha256']==digest(manifest)
    if mode=='prepare':
        output.mkdir(parents=True,exist_ok=False);(output/'pptx').mkdir();rows=[]
        for arm,result in results['arms'].items():
            for row in result['rows']:
                source=evaluation/arm/row['id']/'result.pptx';assert digest(source)==row['pptx_sha256']
                identity=f'{arm}--{row["id"]}';dest=output/'pptx'/(identity+'.pptx');shutil.copyfile(source,dest)
                deck=Presentation(dest);shapes=deck.slides[0].shapes
                assert len(deck.slides)==1 and len(shapes)==len(row['objects'])
                assert all(s.auto_shape_type is not None for s in shapes)
                rows.append(dict(id=identity,arm=arm,case=row['id'],pptx_sha256=digest(dest),shape_count=len(shapes)))
        save(output/'inputs.json',dict(rows=rows,results_sha256=digest(evaluation/'results.json'),code_sha256=digest(__file__)))
        return
    inputs=json.loads((output/'inputs.json').read_text());assert inputs['results_sha256']==digest(evaluation/'results.json')
    assert inputs['code_sha256']==digest(__file__)
    receipts=json.loads((output/'renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
    receipts={r['id']:r for r in receipts['records']};assert len(receipts)==len(inputs['rows'])
    for row in inputs['rows']:
        receipt=receipts[row['id']]
        assert row['pptx_sha256']==receipt['pptx_sha256']==digest(receipt['pptx_path'])
        assert digest(receipt['render_path'])==receipt['render_sha256']
    boards=[]
    for identity,record in records.items():
        board=Image.new('RGB',(1280,240),'white');draw=ImageDraw.Draw(board)
        with Image.open(record['image']) as im:board.paste(im.convert('RGB'),(0,40))
        draw.text((4,4),'RGB source: '+identity,fill='black')
        for col,(arm,result) in enumerate(results['arms'].items(),1):
            row=next(r for r in result['rows'] if r['id']==identity)
            receipt=receipts[f'{arm}--{identity}']
            with Image.open(receipt['render_path']) as im:
                assert im.size==(320,200);board.paste(im.convert('RGB'),(col*320,40))
            draw.text((col*320+4,4),arm,fill='black')
            draw.text((col*320+4,19),f'TP {row["tp"]} / FP {row["fp"]} / FN {row["fn"]}',fill='black')
        path=output/(identity+'.png');assert not path.exists();board.save(path)
        boards.append(dict(id=identity,path=str(path.resolve()),sha256=digest(path)))
    save(output/'audit.json',dict(boards=boards,receipt_sha256=digest(output/'renders/powerpoint_render_manifest.json'),
        checked_pptx=len(inputs['rows']),native_shape_count=sum(r['shape_count'] for r in inputs['rows']),
        scope='Native XML and actual COM receipt audit; visual verdict recorded separately.'))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['prepare','review'])
    for name in ('manifest','evaluation','output'):p.add_argument('--'+name,type=Path,required=True)
    a=p.parse_args();run(a.manifest,a.evaluation,a.output,a.mode)

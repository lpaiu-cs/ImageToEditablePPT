"""Build R1 source/A/B/C sheets and audit real PowerPoint edited copies."""
import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.oxml.ns import qn

from image_to_editable_ppt.research.native_emit import load_scene, normalize_scene
from image_to_editable_ppt.research.region_first import digest, pixel_box


def main():
    root = Path(__file__).resolve().parents[1]
    base = root / 'artifacts/region-first-r1'
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--review-name', default='review-v1')
    parser.add_argument('--evidence-output')
    parser.add_argument('--visual-review')
    args = parser.parse_args()
    output = base / args.review_name
    output.mkdir(exist_ok=False)
    inputs = json.loads((root / 'benchmarks/region_first_r1_inputs_v1.json').read_text())['inputs']
    samples = {s['id']: s for s in json.loads((root / 'benchmarks/region_first_r0_v1.json').read_text())['samples']}
    renders = json.loads((base / 'run1-renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
    renders = {r['id']: r for r in renders['records']}
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 18)
    sheets, invariants = [], []
    for item in inputs:
        name = item['sample']
        sample = samples[name]
        old = root / f"artifacts/region-first-r0/run1-renders/{name}--{item['r0_arm']}.png"
        if digest(old) != item['baseline_render_sha256']:
            raise ValueError('R0 render drift')
        paths = [root / sample['path'], old]
        for arm in ('normalize_only', 'normalized_sites'):
            key = f'{name}--{arm}'
            path = base / f'run1-renders/{key}.png'
            if digest(path) != renders[key]['render_sha256'] or digest(base / f'run1/{key}.pptx') != renders[key]['pptx_sha256']:
                raise ValueError('R1 render drift')
            paths.append(path)
            scene = normalize_scene(load_scene(root/item['scene_path']))
            deck = Presentation(base/f'run1/{key}.pptx')
            objects = {s.name:s for s in deck.slides[0].shapes}
            by_id = {str(s.shape_id):s.name for s in objects.values()}
            baseline = {s.name:s for s in Presentation(root/item['baseline_pptx_path']).slides[0].shapes}
            pairs, arrows = [], []
            for c in scene.connectors:
                links = objects[c.id]._element.xpath('.//a:stCxn | .//a:endCxn')
                pairs.append([by_id[x.get('id')] for x in links]==[c.source_owner_id,c.target_owner_id])
                def ends(obj):
                    return [dict(x.attrib) for x in obj._element.xpath('.//a:headEnd | .//a:tailEnd')]
                arrows.append(ends(objects[c.id])==ends(baseline[c.id]))
            invariants.append(dict(sample=name, arm=arm,
                exact_object_ids=set(objects)=={s.id for s in scene.shapes}|{c.id for c in scene.connectors},
                labels_preserved=all(objects[s.id].text==(s.label or '') for s in scene.shapes),
                owner_pairs_preserved=all(pairs), direction_flags_preserved=all(arrows),
                max_box_error_px=max(abs(a-b) for s in scene.shapes for a,b in zip(
                    (objects[s.id].left/9525,objects[s.id].top/9525,objects[s.id].width/9525,objects[s.id].height/9525),
                    (s.bbox.x0,s.bbox.y0,s.bbox.width,s.bbox.height))),
                raster_count=len(deck.slides[0]._element.xpath('.//p:pic'))))
        panels = []
        for title, path in zip(('SOURCE: target ROI', 'A: frozen R0', 'B: normalized slide', 'C: normalized + native sites'), paths):
            with Image.open(path) as im:
                cropped = im.convert('RGB').crop(pixel_box(sample['roi'], im.size))
            height = max(160, min(700, round(470 * cropped.height / cropped.width)))
            panel = Image.new('RGB', (490, height + 50), '#f1f5f9')
            thumb = ImageOps.contain(cropped, (470, height))
            panel.paste(thumb, ((490-thumb.width)//2, 40+(height-thumb.height)//2))
            ImageDraw.Draw(panel).text((10, 8), title, fill='black', font=font)
            panels.append(panel)
        sheet = Image.new('RGB', (1960, max(p.height for p in panels)), 'white')
        for i, panel in enumerate(panels):
            sheet.paste(panel, (i*490, 0))
        path = output / f'{name}-comparison.png'
        sheet.save(path)
        sheets.append(dict(sample=name, path=str(path.relative_to(root)), sha256=digest(path)))

    scene = normalize_scene(load_scene(root / next(i['scene_path'] for i in inputs if i['sample']=='mixed-attention')))
    owners = {s.id: s for s in scene.shapes}
    probe = json.loads((base / 'edit-probe-v1/probe.json').read_text(encoding='utf-8-sig'))
    audits = []
    for record in probe['records']:
        for stage in ('edited', 'reopened', 'rerouted'):
            path = base / f"edit-probe-v1/{record['arm']}--{stage}.pptx"
            deck = Presentation(path)
            shapes = {s.name:s for s in deck.slides[0].shapes}
            by_id = {str(s.shape_id):s.name for s in shapes.values()}
            errors, edges = [], []
            for c in scene.connectors:
                shape = shapes[c.id]
                wanted = []
                for owner, p in ((c.source_owner_id,c.path_points[0]), (c.target_owner_id,c.path_points[-1])):
                    b = owners[owner].bbox
                    u, v = (p.x-b.x0)/b.width, (p.y-b.y0)/b.height
                    # Expected transform is preregistered, not inferred from the emitted connector.
                    x, y, w, h = b.x0*.75, b.y0*.75, b.width*.75, b.height*.75
                    if owner == 'attention':
                        x += 12; y += 18; w *= 1.25; h *= 1.20
                    wanted.extend((x+u*w, y+v*h))
                actual = [shape.begin_x/12700, shape.begin_y/12700, shape.end_x/12700, shape.end_y/12700]
                error = max(abs(a-b) for a,b in zip(actual,wanted))
                errors.append(error)
                bindings = shape._element.xpath('.//a:stCxn | .//a:endCxn')
                pairs = [by_id[b.get('id')] for b in bindings]
                edges.append(dict(id=c.id, expected_pt=wanted, actual_pt=actual, max_error_pt=error,
                    owners=pairs, owner_pair_preserved=pairs==[c.source_owner_id,c.target_owner_id],
                    sites=[b.get('idx') for b in bindings]))
            source_error = max(abs(actual-expected) for s in scene.shapes if s.id!='attention'
                for actual,expected in zip((shapes[s.id].left/12700,shapes[s.id].top/12700,shapes[s.id].width/12700,shapes[s.id].height/12700),
                    (s.bbox.x0*.75,s.bbox.y0*.75,s.bbox.width*.75,s.bbox.height*.75)))
            b = owners['attention'].bbox
            target_error = max(abs(a-b) for a,b in zip(
                (shapes['attention'].left/12700,shapes['attention'].top/12700,shapes['attention'].width/12700,shapes['attention'].height/12700),
                (b.x0*.75+12,b.y0*.75+18,b.width*.75*1.25,b.height*.75*1.20)))
            geometry = shapes['attention']._element.spPr.find(qn('a:custGeom'))
            audits.append(dict(arm=record['arm'], stage=stage, path=str(path.relative_to(root)), sha256=digest(path),
                max_endpoint_error_pt=max(errors), endpoint_tolerance_pass=max(errors)<=1,
                target_transform_error_pt=target_error, target_custom_site_count=0 if geometry is None else len(geometry.find(qn('a:cxnLst'))),
                render_sha256=digest(path.with_suffix('.png')),
                max_untouched_owner_error_pt=source_error, shape_count=len(shapes), expected_shape_count=len(scene.shapes)+len(scene.connectors),
                labels_preserved=all(shapes[s.id].text==(s.label or '') for s in scene.shapes), edges=edges))
    evidence = dict(status='awaiting_visual_review_not_product_acceptance', sheets=sheets, edit_audits=audits,
        fixed_scene_invariants=invariants,
        probe=probe, run_results=json.loads((base/'run1/results.json').read_text()),
        render_manifest_sha256=digest(base/'run1-renders/powerpoint_render_manifest.json'),
        probe_sha256=digest(base/'edit-probe-v1/probe.json'),
        audit_code_sha256={p:digest(root/p) for p in ('tools/review_native_emit_experiment.py','tools/probe_native_emit_powerpoint.ps1')})
    if args.visual_review:
        evidence['visual_review'] = json.loads((root/args.visual_review).read_text(encoding='utf-8'))
        evidence['status'] = 'reviewed_partial_emitter_gain_reroute_stress_failed_not_product_acceptance'
    (output/'evidence.json').write_text(json.dumps(evidence, indent=2), encoding='utf-8')
    if args.evidence_output:
        with (root/args.evidence_output).open('x', encoding='utf-8') as f:
            json.dump(evidence, f, indent=2)
    for a in audits:
        print(a['arm'], a['stage'], f"endpoint_error={a['max_endpoint_error_pt']:.4f}pt", 'labels=', a['labels_preserved'])


if __name__ == '__main__':
    main()

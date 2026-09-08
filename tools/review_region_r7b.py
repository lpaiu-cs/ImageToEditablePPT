"""R7b 최종 native 렌더 및 동결 전후 비교 증거."""
import json
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from run_region_r7 import ROOT, BASE, ANNOTATIONS, digest, save
from run_region_r7b import verify, LOCK_B
from image_to_editable_ppt.research.accepted_owners import load_annotation


def main():
    verify()
    results = json.loads((BASE / 'r7b-run/results.json').read_text(encoding='utf-8'))
    manifests = {}
    for directory in ('renders', 'r7b-renders'):
        m = json.loads((BASE / directory / 'powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
        manifests[directory] = {r['id']: r for r in m['records']}
        for r in m['records']:
            assert digest(r['pptx_path']) == r['pptx_sha256']
            assert digest(r['render_path']) == r['render_sha256']
    out = BASE / 'r7b-review'
    out.mkdir(parents=True, exist_ok=False)
    font = ImageFont.truetype('C:/Windows/Fonts/malgun.ttf', 20)
    records = []
    for r in results['records']:
        data, rgb = load_annotation(ANNOTATIONS / f'{r["id"]}.json', ROOT)
        original = Image.fromarray(rgb)
        pictures = [('원본', original.crop(data['scope']))]
        for directory, name, label in [('renders', r['id']+'-local-oracle', 'R7 정답 역할 승인'),
                                       ('r7b-renders', r['id'], 'R7b 차단·소유 채움 (정답 역할 의존)')]:
            im = Image.open(BASE / directory / f'{name}.png').convert('RGB')
            box = tuple(round(v*(im.width/original.width if i % 2 == 0 else im.height/original.height))
                        for i, v in enumerate(data['scope']))
            pictures.append((label, im.crop(box)))
        sheet = Image.new('RGB', (1920, 720), '#edf1f5')
        draw = ImageDraw.Draw(sheet)
        for i, (label, im) in enumerate(pictures):
            draw.text((i*640+10, 12), label, fill='black', font=font)
            thumb = ImageOps.contain(im, (620, 650))
            sheet.paste(thumb, (i*640+(640-thumb.width)//2, 50+(650-thumb.height)//2))
        path = out / f'{r["id"]}.png'
        sheet.save(path)
        pp = BASE / 'r7b-run' / f'{r["id"]}.pptx'
        render_record = manifests['r7b-renders'][r['id']]
        assert digest(pp) == r['pptx_sha256'] == render_record['pptx_sha256']
        deck = Presentation(pp)
        assert len(deck.slides) == 1
        slide = deck.slides[0]
        objects = {s.name: s for s in slide.shapes}
        ids = {str(s.shape_id): s.name for s in slide.shapes}
        scene = r['emitted_scene']
        assert set(objects) == {s['id'] for s in scene['shapes']} | {c['id'] for c in scene['connectors']}
        assert len(objects) == len(slide.shapes)
        for c in scene['connectors']:
            assert [ids[b.get('id')] for b in objects[c['id']]._element.xpath('.//a:stCxn | .//a:endCxn')] == [c['source_owner_id'], c['target_owner_id']]
        assert not slide._element.xpath('.//p:pic')
        # Check the actual emitted colors rather than only the intended style dict.
        for sid, style in r['styles'].items():
            if style['fill'] is not None:
                assert list(objects[sid].fill.fore_color.rgb) == style['fill']
        records.append(dict(id=r['id'], source_sha256=data['source_sha256'],
                            before=r['before'], after_blockers=r['after_blockers'],
                            after_boundary=r['after_boundary'], emitted=r['emitted_metrics'],
                            style_audit=r['style_audit'], rejected_boundary=r['observation']['rejected_boundary'],
                            pptx_sha256=digest(pp), render_sha256=render_record['render_sha256'],
                            sheet_sha256=digest(path), native_ids_bindings_fill_verified=True, raster_count=0))
    save(ROOT / 'benchmarks/region_r7b_evidence.json', dict(schema='r7b_diagnostic_evidence_v1',
         population='consumed_development', promotion=False, lock_sha256=digest(LOCK_B),
         results_sha256=digest(BASE / 'r7b-run/results.json'),
         render_manifest_sha256=digest(BASE / 'r7b-renders/powerpoint_render_manifest.json'), records=records))
    print('Four native R7b decks and emitted fills verified.')


if __name__ == '__main__':
    main()

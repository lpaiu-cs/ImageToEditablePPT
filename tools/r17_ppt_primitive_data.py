"""R17 fresh calibration/eval scenes; frozen R16 occupancy recipe, no training."""
import argparse
from collections import Counter
import json
from pathlib import Path
import random
import zipfile
from xml.etree import ElementTree as ET
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from r10_classifier_pilot import EMU, label
from r10_owner_region_p2_audit import sha

KINDS={1:'rectangle',2:'rounded rectangle',3:'ellipse',4:'diamond',5:'parallelogram',6:'cylinder'}
PRESETS=[MSO_SHAPE.RECTANGLE,MSO_SHAPE.ROUNDED_RECTANGLE,MSO_SHAPE.OVAL,MSO_SHAPE.DIAMOND,MSO_SHAPE.PARALLELOGRAM,MSO_SHAPE.CAN]
PALETTE=[[0,0,0],[255,0,0],[0,255,0],[0,0,255],[255,255,0],[255,0,255],[0,255,255]]
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}


def no_effect(shape):
    shape._element.spPr.append(OxmlElement('a:effectLst'))


def build(out):
    out.mkdir(parents=True,exist_ok=False);(out/'pptx').mkdir()
    records=[]
    groups=[('scene','calibration',12,17022),('scene','eval',12,17023),('negative','calibration',4,17032),('negative','eval',4,17033)]
    for stage,split,count,seed in groups:
        rng=random.Random(seed)
        for i in range(count):
            identity=f'r17-{stage}-{split}-{i:03d}';objects=[]
            kinds=list(KINDS) if stage=='scene' else []
            rng.shuffle(kinds)
            for j,kind in enumerate(kinds):
                w=rng.randint(120,205);h=rng.randint(95,130 if split=='eval' else 170)
                if split=='eval':x=60+(j%2)*380+rng.randint(0,40);y=15+(j//2)*160+rng.randint(0,15)
                else:x=20+(j%3)*260+rng.randint(0,30);y=25+(j//3)*240+rng.randint(0,25)
                objects.append(dict(id=f'object{j}',kind=kind,bbox_source=[x,y,x+w,y+h],bbox=[v*.4 for v in [x,y,x+w,y+h]],
                    fill=[rng.randint(110,245) for _ in range(3)],stroke=[rng.randint(0,80) for _ in range(3)],stroke_pt=rng.choice([1,1.5,2,2.5]),
                    text=rng.choice(['','Read','Build','Check','A','B']),adjustments=[]))
            negative=[]
            if stage=='negative':
                for j in range(3):
                    x=rng.randint(30,200);y=70+j*140+rng.randint(-15,15)
                    negative.append(dict(start=[x,y],end=[rng.randint(400,740),y+rng.randint(-25,25)],text=rng.choice(['Read','Build','Check','Input'])+f' {i}-{j}'))
            records.append(dict(id=identity,split=split,stage=stage,seed=seed,scene_id=identity,
                layout_family=('grid2x3' if split=='eval' else 'grid3x2') if stage=='scene' else 'open_lines_text',
                lineage_id=identity,objects=objects,negative_elements=negative))
    prereg=dict(records=records,classmapping=KINDS,palette=PALETTE,builder_sha256=sha(__file__),
        helper_sha256=sha(Path(__file__).with_name('r10_classifier_pilot.py')),
        contract='Frozen R16 white COM occupancy recipe. Fresh calibration grid3x2 and eval grid2x3; eval heights95..130 versus train/cal95..170. Negative layout family shared. No training or prediction-based targets. Eval visual outputs not inspected or shared with root before evaluation.')
    (out/'preregister.json').write_text(json.dumps(prereg,indent=2))
    for r in records:
        geometry=[]
        for mask in [False,True]:
            p=Presentation();p.slide_width=Emu(800*EMU);p.slide_height=Emu(500*EMU)
            s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(*(PALETTE[0] if mask else [255]*3))
            for obj in r['objects']:
                x,y,a,b=obj['bbox_source'];v=s.shapes.add_shape(PRESETS[obj['kind']-1],Emu(x*EMU),Emu(y*EMU),Emu((a-x)*EMU),Emu((b-y)*EMU));v.name=obj['id'];no_effect(v)
                v.fill.solid();v.fill.fore_color.rgb=RGBColor(*([255,255,255] if mask else obj['fill']))
                if mask:v.line.fill.background()
                else:v.line.color.rgb=RGBColor(*obj['stroke']);v.line.width=Pt(obj['stroke_pt'])
                adjustments=list(v.adjustments)
                if mask:assert adjustments==obj['adjustments']
                else:obj['adjustments']=adjustments
                if not mask and obj['text']:
                    label(s,obj['text'],x+20,y+20,a-x-35)
            if not mask:
                for e in r['negative_elements']:
                    a,b=e['start'],e['end'];v=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Emu(q*EMU) for q in [*a,*b]])
                    v.line.color.rgb=RGBColor(0,0,0);v.line.width=Pt(1.5);no_effect(v);label(s,e['text'],a[0],a[1]+12,200)
            path=out/'pptx'/f'{r["id"]}-{"mask" if mask else "input"}.pptx';p.save(path)
            with zipfile.ZipFile(path) as z:xml=ET.fromstring(z.read('ppt/slides/slide1.xml'))
            assert not xml.findall('.//p:pic',NS)
            current=[]
            for v in xml.findall('.//p:sp',NS):
                name=v.find('p:nvSpPr/p:cNvPr',NS).get('name')
                if name.startswith('object'):
                    current.append((name,ET.tostring(v.find('p:spPr/a:xfrm',NS)),ET.tostring(v.find('p:spPr/a:prstGeom',NS))))
            if mask:assert current==geometry
            else:geometry=current
            r['mask_pptx' if mask else 'pptx']=str(path.resolve());r['mask_pptx_sha256' if mask else 'pptx_sha256']=sha(path)
    (out/'source-records.json').write_text(json.dumps(records,indent=2))
    assert len(records)==32
    print('32 pairs / 64 source PPTs; prereg',sha(out/'preregister.json'))


def finalize(out):
    records=json.loads((out/'source-records.json').read_text());m=out/'renders/powerpoint_render_manifest.json'
    receipt={r['id']:r for r in json.loads(m.read_text(encoding='utf-8-sig'))['records']}
    assert len(receipt)==64
    (out/'rgb').mkdir();(out/'targets').mkdir();seen={}
    for r in records:
        for suffix,key in [('input','image'),('mask','mask')]:
            q=receipt[r['id']+'-'+suffix];source_key='pptx' if suffix=='input' else 'mask_pptx'
            assert sha(q['pptx_path'])==q['pptx_sha256']==r[source_key+'_sha256']
            assert sha(q['render_path'])==q['render_sha256']
            with Image.open(q['render_path']) as im:
                assert 'transparency' not in im.info
                a=np.array(im.convert('RGB'));assert a.shape==(200,320,3)
            dest=out/'rgb'/f'{r["id"]}-{suffix}.png';Image.fromarray(a).save(dest)
            assert np.array_equal(a,np.array(Image.open(dest)))
            r[key]=str(dest.resolve());r[key+'_sha256']=sha(dest);r[key+'_render_sha256']=q['render_sha256']
            if suffix=='input':
                assert sha(dest) not in seen,(r['id'],seen.get(sha(dest)))
                seen[sha(dest)]=r['id']
            else:
                r['mask_raw']=r['mask'];r['mask_raw_sha256']=r['mask_sha256']
                white=np.all(a==255,axis=2);black=np.all(a==0,axis=2)
                colored=np.full_like(a,128);colored[black]=0
                coverage=np.zeros((200,320),np.uint8)
                for obj in r['objects']:
                    x,y,c,d=obj['bbox'];region=np.zeros((200,320),bool)
                    region[int(np.floor(y)):int(np.ceil(d)),int(np.floor(x)):int(np.ceil(c))]=True
                    assert not np.any(coverage[region])
                    coverage[region]=1
                    pixels=white & region;assert pixels.any()
                    colored[pixels]=PALETTE[obj['kind']]
                assert not np.any(white & (coverage==0))
                a=colored;path=out/'rgb'/f'{r["id"]}-palette.png';Image.fromarray(a).save(path)
                r['mask']=str(path.resolve());r['mask_sha256']=sha(path)
                target=np.full((200,320),255,np.uint8)
                for kind,color in enumerate(PALETTE):target[np.all(a==color,axis=2)]=kind
                expected={v['kind'] for v in r['objects']}
                assert set(np.unique(target))-{0,255}==expected
                if r['stage']=='negative':assert np.all(target==0)
                r['target_pixel_counts']={str(v):int(np.count_nonzero(target==v)) for v in np.unique(target)}
                path=out/'targets'/f'{r["id"]}.png';Image.fromarray(target).save(path);r['target']=str(path.resolve());r['target_sha256']=sha(path)
        for i,a in enumerate(r['objects']):
            x,y,c,d=a['bbox']
            assert 0<x<c<320 and 0<y<d<200
            for b in r['objects'][i+1:]:
                u,v,w,h=b['bbox'];assert c<=u or w<=x or d<=v or h<=y
    counts=Counter((r['stage'],r['split']) for r in records)
    manifest=dict(records=records,classmapping=KINDS,palette=PALETTE,ignore_index=255,
        source_size=[800,500],render_size=[320,200],generator_sha256=sha(__file__),
        preregister_sha256=sha(out/'preregister.json'),com_manifest_sha256=sha(m),
        counts={f'{a}/{b}':v for (a,b),v in counts.items()},unique_input_rgb=32,
        scope='Fresh calibration grid3x2 / eval grid2x3. Eval heights restricted95..130. Shared negative family. No instance/overlap/renderer generalization.')
    (out/'dataset_manifest.json').write_text(json.dumps(manifest,indent=2));print(manifest['counts'],sha(out/'dataset_manifest.json'))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['build','finalize']);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    (build if a.mode=='build' else finalize)(a.out)

"""Read-only artifact audit and scientific comparison sheets for R3-E0/E1."""
import argparse
import json
import math
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from image_to_editable_ppt.research.native_emit import load_scene, normalize_scene
from image_to_editable_ppt.research.preset_emit import preset_vertices, distance, samples
from image_to_editable_ppt.research.region_first import digest, pixel_box


def read(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def audit(path, scene, fits, edited_owner=None):
    deck = Presentation(path)
    objects = {s.name: s for s in deck.slides[0].shapes}
    ids = {str(s.shape_id): s.name for s in objects.values()}
    owners = {s.id: s for s in scene.shapes}
    checks = []
    for c in scene.connectors:
        obj = objects[c.id]
        actual = preset_vertices(obj)
        expected = []
        for owner, p in [(c.source_owner_id, c.path_points[0]), (c.target_owner_id, c.path_points[-1])]:
            x, y = p.x * .75, p.y * .75
            b = owners[owner].bbox
            if owner == edited_owner:
                x = b.x0 * .75 + 12 + (p.x-b.x0) * .75 * 1.25
                y = b.y0 * .75 + 18 + (p.y-b.y0) * .75 * 1.20
            expected.append((x, y))
        links = obj._element.xpath('.//a:stCxn | .//a:endCxn')
        check = dict(connector=c.id, vertices_pt=actual,
            arrow_ends=[dict(e.attrib) for e in obj._element.xpath('.//a:headEnd | .//a:tailEnd')],
            owner_bindings=[ids[l.get('id')] for l in links], site_indices=[int(l.get('idx')) for l in links],
            expected_endpoints_pt=expected,
            max_endpoint_error_pt=max(math.dist(a,b) for a,b in zip([actual[0],actual[-1]],expected)))
        check['owners_preserved'] = check['owner_bindings'] == [c.source_owner_id,c.target_owner_id]
        if edited_owner is None:
            route = [(x*.75,y*.75) for x,y in fits[c.id]['route_px']]
            check['fit_route_error_pt'] = max(distance(samples(route),actual),distance(samples(actual),route))
        checks.append(check)
    box_errors=[]
    for s in scene.shapes:
        b=s.bbox; expected=[b.x0*.75,b.y0*.75,b.width*.75,b.height*.75]
        if s.id==edited_owner:
            expected=[expected[0]+12,expected[1]+18,expected[2]*1.25,expected[3]*1.20]
        obj=objects[s.id]
        box_errors.extend(abs(a-b) for a,b in zip([obj.left/12700,obj.top/12700,obj.width/12700,obj.height/12700],expected))
    return dict(path=str(path),sha256=digest(path),connectors=checks,
        native_ids_exact=set(objects)==set(owners)|{c.id for c in scene.connectors},
        labels_preserved=all(objects[s.id].text==(s.label or '') for s in scene.shapes),
        raster_count=len(deck.slides[0]._element.xpath('.//p:pic')),
        max_box_error_pt=max(box_errors,default=0),
        max_endpoint_error_pt=max((c['max_endpoint_error_pt'] for c in checks),default=0))


def sheet(paths, labels, roi, output):
    panels=[]
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
    for path,label in zip(paths,labels):
        with Image.open(path) as im:
            crop=im.convert('RGB').crop(pixel_box(roi,im.size))
        height=max(150,min(600,round(470*crop.height/crop.width)))
        panel=Image.new('RGB',(490,height+50),'#f1f5f9')
        thumb=ImageOps.contain(crop,(470,height))
        panel.paste(thumb,((490-thumb.width)//2,40+(height-thumb.height)//2))
        ImageDraw.Draw(panel).text((10,10),label,font=font,fill='black')
        panels.append(panel)
    result=Image.new('RGB',(490*len(panels),max(p.height for p in panels)),'white')
    for i,p in enumerate(panels): result.paste(p,(490*i,0))
    result.save(output)
    return dict(path=str(output),sha256=digest(output))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',required=True)
    parser.add_argument('--visual-review')
    parser.add_argument('--evidence-output')
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]
    out=root/args.output_dir; out.mkdir(parents=True,exist_ok=False)
    e1=root/'artifacts/region-first-r3e1'; e0=root/'artifacts/region-first-r3e'
    result=read(e1/'run1/results.json'); render=read(e1/'run1-renders/powerpoint_render_manifest.json')
    inputs={r['sample']:r for r in read(root/'benchmarks/region_first_r3_inputs_v1.json')['inputs']}
    probe=read(e1/'edit-probe-v1/probe.json')
    initial=[]; edits=[]; sheets=[]
    for r in result['records']:
        sample=r['sample']; item=inputs[sample]
        assert digest(root/r['input_scene_path'])==r['input_scene_sha256']
        assert digest(root/r['pptx_path'])==r['pptx_sha256']
        assert digest(root/item['source_path'])==item['source_sha256']
        scene=normalize_scene(load_scene(root/r['input_scene_path']))
        a=audit(root/r['pptx_path'],scene,r['sites']); a['sample']=sample
        assert a['native_ids_exact'] and a['labels_preserved'] and a['raster_count']==0
        assert a['max_endpoint_error_pt']<1 and a['max_box_error_pt']<1
        assert all(c['owners_preserved'] and c['fit_route_error_pt']<1 for c in a['connectors'])
        initial.append(a)
        control=root/f'artifacts/region-first-r3/run1-renders/{sample}--transparent_gap_trace.png'
        native=e1/f'run1-renders/{sample}--preset_path.png'
        sheets.append(sheet([root/item['source_path'],control,native],['Source ROI','R3 trace + old writer','R3 trace + native preset'],item['roi'],out/f'{sample}-initial.png'))
        rec=next((p for p in probe['records'] if p['sample']==sample),None)
        if rec:
            assert rec['input_sha256']==r['pptx_sha256'] and rec['original_unchanged'] and rec['text_edit_succeeded']
            stages=[]
            for stage in ('roundtrip','edited','reopened','rerouted'):
                path=e1/f'edit-probe-v1/{sample}--{stage}.pptx'
                check=audit(path,scene,r['sites'],None if stage=='roundtrip' else rec['edited_owner'])
                check['stage']=stage
                check['sites_preserved']=all(x['site_indices']==y['site_indices'] for x,y in zip(a['connectors'],check['connectors']))
                check['arrow_ends_preserved']=all(x['arrow_ends']==y['arrow_ends'] for x,y in zip(a['connectors'],check['connectors']))
                assert check['native_ids_exact'] and check['labels_preserved'] and check['raster_count']==0
                assert all(c['owners_preserved'] for c in check['connectors'])
                assert check['arrow_ends_preserved'] and check['max_box_error_pt']<1
                if stage!='rerouted':
                    assert check['sites_preserved'] and check['max_endpoint_error_pt']<1
                if stage=='roundtrip':
                    assert all(c['fit_route_error_pt']<1 for c in check['connectors'])
                stages.append(check)
            edits.append(dict(sample=sample,stages=stages))
            # Full slide after movement: cropping to the old ROI would hide the moved owner.
            sheets.append(sheet([e1/f'edit-probe-v1/{sample}--{s}.png' for s in ('roundtrip','reopened','rerouted')],
                ['Native roundtrip','Move + resize + reopen','Explicit reroute stress'],[0,0,1,1],out/f'{sample}-edit.png'))
    assert len(initial)==len(render['records'])==6
    for r in render['records']:
        assert digest(Path(r['pptx_path']))==r['pptx_sha256']
        assert digest(Path(r['render_path']))==r['render_sha256']
    open_probe=read(e0/'open_probe.json')
    for r in open_probe['records']: assert digest(Path(r['path']))==r['sha256']
    files=[e0/'run1/results.json',e0/'open_probe.json',e1/'run1/results.json',e1/'run1-renders/powerpoint_render_manifest.json',e1/'edit-probe-v1/probe.json']
    files+=list((e1/'edit-probe-v1').glob('*.png'))+list((e1/'edit-probe-v1').glob('*.pptx'))
    evidence=dict(status='development_only_not_product_acceptance',e0_results=read(e0/'run1/results.json'),
        e0_open_probe=open_probe,e1_results=result,e1_render_manifest=render,initial_audits=initial,
        edit_audits=edits,native_probe=probe,sheets=sheets,
        provenance={str(p.relative_to(root)):digest(p) for p in files},
        audit_code_sha256=digest(Path(__file__)),probe_code_sha256=digest(root/'tools/probe_preset_paths.ps1'))
    if args.visual_review: evidence['visual_review']=read(root/args.visual_review)
    (out/'evidence.json').write_text(json.dumps(evidence,indent=2),encoding='utf-8')
    if args.evidence_output:
        with (root/args.evidence_output).open('x',encoding='utf-8') as f: json.dump(evidence,f,indent=2)
    print(json.dumps(dict(initial_max_endpoint_pt=max(a['max_endpoint_error_pt'] for a in initial),
        edits=[dict(sample=e['sample'],stages=[dict(stage=s['stage'],endpoint_pt=s['max_endpoint_error_pt'],sites_preserved=s['sites_preserved']) for s in e['stages']]) for e in edits]),indent=2))


if __name__=='__main__': main()

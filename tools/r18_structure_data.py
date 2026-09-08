"""Role-labelled source PPT scenes and hard negatives, no model or union masks."""
import argparse
from collections import Counter
import json
from pathlib import Path
import random
import zipfile
from xml.etree import ElementTree as ET
import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from r16_ppt_primitive_data import KINDS, PRESETS, NS, no_effect
from r10_classifier_pilot import EMU
from r10_owner_region_p2_audit import sha

SCALE=640/1200
FAMILIES=['colored_fill_white_bg','white_card_color_bg','outline_white_bg','nested_container']
WORDS={'development':['Read','Build','Check','Input','Process','Store','Task','Group'],
       'eval':['Observe','Route','Assess','Signal','Archive','Result','Stage','Cluster']}


def build(out):
    out.mkdir(parents=True,exist_ok=False);(out/'pptx').mkdir()
    records=[]
    for split,count,negative_count,seed,negative_seed in [('train',32,8,18011,18021),('calibration',12,4,18012,18022),('eval',12,4,18013,18023)]:
        for stage,n,current_seed in [('scene',count,seed),('negative',negative_count,negative_seed)]:
            rng=random.Random(current_seed);words=WORDS['eval' if split=='eval' else 'development']
            for i in range(n):
                identity=f'r18-{stage}-{split}-{i:03d}';family=FAMILIES[i%4] if stage=='scene' else 'hard_negative'
                background=[rng.randint(155,220) for _ in range(3)] if family=='white_card_color_bg' else [255]*3
                elements=[dict(id='background',role='background',type='background',color=background)]
                objects=[]
                def shape(name,kind,bbox,role,fill,stroke,parent=None):
                    item=dict(id=name,type='shape',kind=kind,role=role,bbox_source=bbox,bbox=[v*SCALE for v in bbox],fill=fill,stroke=stroke,stroke_pt=rng.choice([1.5,2,2.5]),parent_id=parent)
                    elements.append(item)
                    if role=='closed_structure':objects.append(item)
                def text(name,content,bbox,role='decoration',owner=None,font=16):
                    elements.append(dict(id=name,type='text',role=role,text=content,bbox_source=bbox,font_pt=font,owner_id=owner))
                def line(name,a,b):
                    elements.append(dict(id=name,type='line',role='open_structure',start=a,end=b,stroke_pt=1.5))
                shape('banner',1,[20,15,1180,65],'decoration',[rng.randint(190,235) for _ in range(3)],[80]*3)
                text('title',' '.join(rng.sample(words,3)),[40,18,1120,62],font=24)
                text('paragraph-top',' '.join(rng.choices(words,k=14)),[60,80,1140,120],font=14)
                text('paragraph-bottom',' '.join(rng.choices(words,k=14)),[60,712,1140,748],font=14)
                if stage=='scene':
                    parent='container' if family=='nested_container' else None
                    if parent:shape(parent,2,[50,135,1150,705],'closed_structure',None,[65,75,90])
                    kinds=list(KINDS);rng.shuffle(kinds)
                    for j,kind in enumerate(kinds):
                        w=rng.randint(180,250);h=rng.randint(120,155) if split=='eval' else rng.randint(125,180)
                        if split=='eval':x=180+(j%2)*560+rng.randint(0,35);y=165+(j//2)*175+rng.randint(0,15)
                        else:x=80+(j%3)*380+rng.randint(0,25);y=170+(j//3)*270+rng.randint(0,20)
                        fill=None if family=='outline_white_bg' else [255]*3 if family=='white_card_color_bg' else [rng.randint(130,235) for _ in range(3)]
                        node=f'node{j}';shape(node,kind,[x,y,x+w,y+h],'closed_structure',fill,[rng.randint(0,65) for _ in range(3)],parent)
                        text(node+'-label','\n'.join(rng.sample(words,2)),[x+40,y+35,x+w-35,y+h-20],'structural_label',node)
                        if rng.random()<.55:
                            dx=x+w-40;dy=y+12
                            if rng.random()<.5:
                                shape(node+'-glyph',7,[dx,dy,dx+22,dy+22],'decoration',[70,90,130],[50]*3)
                            else:
                                for k in range(2):shape(node+f'-dot{k}',3,[dx+k*14,dy,dx+10+k*14,dy+10],'decoration',[80]*3,[80]*3)
                                line(node+'-glyph-line',[dx,dy+18],[dx+24,dy+18]);elements[-1]['role']='decoration'
                    # Observed open structures, with no asserted connector endpoints or relations.
                    line('open-top',[350,130],[800,128]);line('open-middle',[350,395],[730,390])
                else:
                    for j in range(5):
                        x=rng.randint(60,220);y=150+j*95+rng.randint(0,15)
                        text(f'body{j}',' '.join(rng.choices(words,k=12)),[x,y,1050,y+36],font=16)
                        line(f'open{j}',[x,y+45],[rng.randint(750,1120),y+45+rng.randint(-15,15)])
                        shape(f'glyph{j}',rng.choice([3,7]),[1100,y,1124,y+24],'decoration',[90,110,150],[60]*3)
                records.append(dict(id=identity,split=split,stage=stage,family=family,seed=current_seed,
                    layout_family=('grid2x3' if split=='eval' else 'grid3x2') if stage=='scene' else 'text_open_glyph',
                    vocabulary_group='eval' if split=='eval' else 'development',objects=objects,elements=elements))
    prereg=dict(records=records,classmapping=KINDS,builder_sha256=sha(__file__),
        helper_sha256={name:sha(Path(__file__).with_name(name)) for name in ['r16_ppt_primitive_data.py','r10_classifier_pilot.py','r10_owner_region_p2_audit.py']},
        contract='Objects are closed_structure only; native decorations are not targets. Nested parent separate with child parent_id. No masks, downloads, photos, or model use. Eval vocabulary/layout separated; height distribution also changes.')
    (out/'preregister.json').write_text(json.dumps(prereg,indent=2))
    for r in records:
        p=Presentation();p.slide_width=Emu(1200*EMU);p.slide_height=Emu(750*EMU);s=p.slides.add_slide(p.slide_layouts[6])
        for e in r['elements']:
            if e['type']=='background':
                s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(*e['color']);continue
            if e['type'] in ['shape','text']:
                x,y,a,b=e['bbox_source'];position=[Emu(round(v*EMU)) for v in [x,y,a-x,b-y]]
                if e['type']=='text':
                    v=s.shapes.add_textbox(*position);v.text=e['text']
                    for paragraph in v.text_frame.paragraphs:
                        for run in paragraph.runs:run.font.size=Pt(e['font_pt']);run.font.color.rgb=RGBColor(25,30,40)
                else:
                    v=s.shapes.add_shape(PRESETS[e['kind']-1] if e['kind']<=6 else MSO_SHAPE.STAR_5_POINT,*position)
                    if e['fill'] is None:v.fill.background()
                    else:v.fill.solid();v.fill.fore_color.rgb=RGBColor(*e['fill'])
                    v.line.color.rgb=RGBColor(*e['stroke']);v.line.width=Pt(e['stroke_pt']);e['adjustments']=list(v.adjustments)
            else:
                v=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Emu(round(q*EMU)) for q in [*e['start'],*e['end']]])
                v.line.color.rgb=RGBColor(45,50,60);v.line.width=Pt(e['stroke_pt'])
            v.name=e['id'];no_effect(v);e['source_shape_id']=v.shape_id
        path=out/'pptx'/f'{r["id"]}.pptx';p.save(path)
        with zipfile.ZipFile(path) as z:
            assert len([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')])==1
            xml=ET.fromstring(z.read('ppt/slides/slide1.xml'))
        assert not xml.findall('.//p:pic',NS)
        xml_shapes={v.find('p:nvSpPr/p:cNvPr',NS).get('name'):v for v in xml.findall('.//p:sp',NS)}
        for obj in r['objects']:
            v=xml_shapes[obj['id']];xf=v.find('p:spPr/a:xfrm',NS);off=xf.find('a:off',NS);ext=xf.find('a:ext',NS)
            x,y,a,b=obj['bbox_source'];assert [int(off.get('x')),int(off.get('y')),int(ext.get('cx')),int(ext.get('cy'))]==[round(q*EMU) for q in [x,y,a-x,b-y]]
            assert obj['role']=='closed_structure' and 1<=obj['kind']<=6
        assert len(r['objects'])==(7 if r['family']=='nested_container' else 6 if r['stage']=='scene' else 0)
        r['pptx']=str(path.resolve());r['pptx_sha256']=sha(path)
    assert len(records)==72
    (out/'source-records.json').write_text(json.dumps(records,indent=2));print('72 source PPTs; prereg',sha(out/'preregister.json'))


def finalize(out):
    records=json.loads((out/'source-records.json').read_text());m=out/'renders/powerpoint_render_manifest.json'
    receipts={r['id']:r for r in json.loads(m.read_text(encoding='utf-8-sig'))['records']};assert len(receipts)==72
    (out/'rgb').mkdir();seen=set()
    for r in records:
        receipt=receipts[r['id']];assert sha(receipt['pptx_path'])==receipt['pptx_sha256']==r['pptx_sha256']
        assert sha(receipt['render_path'])==receipt['render_sha256']
        with Image.open(receipt['render_path']) as im:
            assert 'transparency' not in im.info;pixels=np.array(im.convert('RGB'));assert pixels.shape==(400,640,3)
        path=out/'rgb'/f'{r["id"]}.png';Image.fromarray(pixels).save(path);assert np.array_equal(pixels,np.array(Image.open(path)))
        digest=sha(path);assert digest not in seen;seen.add(digest)
        r.update(image=str(path.resolve()),image_sha256=digest,render_sha256=receipt['render_sha256'])
        nodes=[v for v in r['objects'] if v['id']!='container']
        for i,a in enumerate(nodes):
            x,y,c,d=a['bbox'];assert 0<x<c<640 and 0<y<d<400
            for b in nodes[i+1:]:
                u,v,w,h=b['bbox'];assert c<=u or w<=x or d<=v or h<=y
            if a['parent_id']:
                parent=next(v for v in r['objects'] if v['id']==a['parent_id']);u,v,w,h=parent['bbox'];assert u<x<c<w and v<y<d<h
        assert {v['id'] for v in r['objects']}=={v['id'] for v in r['elements'] if v['role']=='closed_structure'}
    counts=Counter((r['split'],r['stage']) for r in records)
    manifest=dict(records=records,classmapping=KINDS,source_size=[1200,750],render_size=[640,400],
        generator_sha256=sha(__file__),preregister_sha256=sha(out/'preregister.json'),com_manifest_sha256=sha(m),
        counts={f'{a}/{b}':n for (a,b),n in counts.items()},scope='Role-labelled source structures; no union masks. Synthetic native decorations only, no natural photos. Eval layout and vocabulary plus height shift.')
    (out/'dataset_manifest.json').write_text(json.dumps(manifest,indent=2));print(manifest['counts'],sha(out/'dataset_manifest.json'))


def qa(out):
    m=json.loads((out/'dataset_manifest.json').read_text());records=m['records']
    for split,n in [('train',8),('calibration',3),('eval',3)]:
        assert Counter(r['family'] for r in records if r['split']==split and r['stage']=='scene')=={f:n for f in FAMILIES}
    assert not set(WORDS['development']) & set(WORDS['eval'])
    selected=[next(r for r in records if r['split']==split and r['family']==family) for split in ['train','calibration'] for family in FAMILIES+['hard_negative']]
    for i,r in enumerate(selected):
        im=Image.open(r['image']);board=Image.new('RGB',(1280,430),'white');board.paste(im,(0,30));board.paste(im,(640,30));draw=ImageDraw.Draw(board)
        draw.text((5,5),r['id']+' | source / closed_structure GT only',fill='black')
        for obj in r['objects']:
            x,y,a,b=obj['bbox'];draw.rectangle([x+640,y+30,a+640,b+30],outline='red',width=2);draw.text((x+642,y+32),f'{obj["id"]}:k{obj["kind"]}',fill='red')
        board.save(out/f'development-qa-{i:02d}.png')
    (out/'qa.json').write_text(json.dumps(dict(records=72,family_balance=True,vocabulary_disjoint=True,eval_visual_opened=False,model_run=False),indent=2))
    print('QA passed; 10 train/calibration boards only')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['build','finalize','qa']);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    {'build':build,'finalize':finalize,'qa':qa}[a.mode](a.out)

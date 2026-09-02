"""Prepare and validate the human PowerPoint edit-operation calibration."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import json
from pathlib import Path
from statistics import median
from typing import Any, Mapping

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from image_to_editable_ppt.ml.phase16.validation import sha256_file


CALIBRATION_SCHEMA = "phase16_operation_time_calibration_v1"


def _load(path: Path) -> dict[str, Any]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ValueError(f"{path} must contain a JSON object")
    return raw


def _add_instruction(slide, title: str, task: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.8))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(24)
    paragraph.font.bold = True
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.0), Inches(0.7))
    paragraph = body.text_frame.paragraphs[0]
    paragraph.text = task
    paragraph.font.size = Pt(17)


def _rect(slide, x: float, y: float, w: float, h: float, *, text: str, color: RGBColor):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = RGBColor(60, 60, 60)
    shape.text = text
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].font.size = Pt(16)
    return shape


def _ghost(slide, x: float, y: float, w: float, h: float):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.background()
    shape.line.color.rgb = RGBColor(150, 150, 150)
    shape.line.width = Pt(1.5)
    return shape


def _connector(slide, x0: float, y0: float, x1: float, y1: float, *, red: bool = False):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1)
    )
    line.line.color.rgb = RGBColor(210, 40, 40) if red else RGBColor(40, 70, 130)
    line.line.width = Pt(2.5)
    return line


def prepare_calibration_deck(
    protocol: Mapping[str, object], target: Path, *, overwrite: bool = False
) -> None:
    if target.exists() and not overwrite:
        raise ValueError(f"refusing to overwrite calibration deck: {target}")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    for operation in protocol["operations"]:
        operation_id = str(operation["id"])
        slide = presentation.slides.add_slide(blank)
        _add_instruction(slide, operation_id, str(operation["task"]))
        blue = RGBColor(205, 225, 248)
        red = RGBColor(250, 190, 190)
        if operation_id == "object_create":
            _ghost(slide, 5.0, 2.4, 2.5, 1.4)
            _rect(slide, 0.8, 5.5, 3.0, 0.7, text="Ghost target is above", color=blue)
        elif operation_id == "object_delete":
            _rect(slide, 5.0, 2.4, 2.5, 1.4, text="DELETE ME", color=red)
        elif operation_id == "geometry_correct":
            _ghost(slide, 7.0, 3.8, 2.8, 1.5)
            _rect(slide, 2.0, 2.2, 1.7, 0.8, text="MOVE + RESIZE", color=blue)
        elif operation_id == "connector_create":
            _rect(slide, 2.0, 2.8, 2.3, 1.2, text="Source", color=blue)
            _rect(slide, 8.5, 2.8, 2.3, 1.2, text="Target", color=blue)
        elif operation_id == "connector_delete":
            _rect(slide, 2.0, 2.8, 2.3, 1.2, text="Source", color=blue)
            _rect(slide, 8.5, 2.8, 2.3, 1.2, text="Target", color=blue)
            _connector(slide, 4.3, 3.4, 8.5, 3.4, red=True)
        elif operation_id == "connector_reroute_direction":
            _rect(slide, 2.0, 4.4, 2.3, 1.2, text="Source", color=blue)
            _rect(slide, 8.5, 2.0, 2.3, 1.2, text="Target", color=blue)
            _connector(slide, 4.3, 5.0, 8.5, 2.6, red=True)
        elif operation_id == "endpoint_rebind":
            _rect(slide, 2.0, 2.8, 2.3, 1.2, text="Source", color=blue)
            _rect(slide, 8.5, 2.8, 2.3, 1.2, text="Target", color=blue)
            _connector(slide, 4.3, 3.4, 7.7, 3.4)
        elif operation_id == "label_retype_delete":
            _rect(slide, 4.2, 2.7, 4.8, 1.5, text="WRONG LABLE", color=blue)
            target_box = slide.shapes.add_textbox(
                Inches(4.2), Inches(4.8), Inches(4.8), Inches(0.6)
            )
            target_box.text = "Target text: Correct Label"
            target_box.text_frame.paragraphs[0].font.size = Pt(18)
        elif operation_id == "redundant_text_delete":
            _rect(slide, 4.2, 2.4, 4.8, 1.5, text="Keep this shape label", color=blue)
            duplicate = slide.shapes.add_textbox(
                Inches(4.2), Inches(4.5), Inches(4.8), Inches(0.8)
            )
            duplicate.text = "DELETE THIS REDUNDANT TEXT"
            duplicate.text_frame.paragraphs[0].font.size = Pt(18)
            duplicate.text_frame.paragraphs[0].font.color.rgb = RGBColor(210, 40, 40)
        else:
            raise ValueError(f"unsupported calibration operation: {operation_id}")
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(target)


def prepare_measurement_csv(
    protocol: Mapping[str, object], target: Path, *, overwrite: bool = False
) -> None:
    if target.exists() and not overwrite:
        raise ValueError(f"refusing to overwrite measurement sheet: {target}")
    repetitions = int(protocol["minimum_repetitions_per_operation"])
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=["operation_id", "repetition", "seconds", "valid", "note"]
        )
        writer.writeheader()
        for operation in protocol["operations"]:
            for repetition in range(1, repetitions + 1):
                writer.writerow(
                    {
                        "operation_id": operation["id"],
                        "repetition": repetition,
                        "seconds": "",
                        "valid": "true",
                        "note": "",
                    }
                )


def summarize_measurements(
    protocol: Mapping[str, object], measurements_path: Path
) -> dict[str, object]:
    minimum = int(protocol["minimum_repetitions_per_operation"])
    expected = {str(item["id"]) for item in protocol["operations"]}
    values: dict[str, list[float]] = {item: [] for item in expected}
    repetitions: dict[str, set[int]] = {item: set() for item in expected}
    with measurements_path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row_number, row in enumerate(csv.DictReader(handle), start=2):
            operation_id = str(row.get("operation_id", ""))
            if operation_id not in expected:
                raise ValueError(f"row {row_number}: unknown operation_id {operation_id!r}")
            valid = str(row.get("valid", "true")).strip().lower()
            if valid not in {"true", "false"}:
                raise ValueError(f"row {row_number}: valid must be true or false")
            try:
                repetition = int(str(row.get("repetition", "")).strip())
            except ValueError as error:
                raise ValueError(f"row {row_number}: repetition must be an integer") from error
            if repetition < 1:
                raise ValueError(f"row {row_number}: repetition must be positive")
            if repetition in repetitions[operation_id]:
                raise ValueError(
                    f"row {row_number}: duplicate repetition {repetition} for {operation_id}"
                )
            repetitions[operation_id].add(repetition)
            if valid == "false":
                continue
            try:
                seconds = float(str(row.get("seconds", "")).strip())
            except ValueError as error:
                raise ValueError(f"row {row_number}: seconds must be numeric") from error
            if seconds <= 0.0:
                raise ValueError(f"row {row_number}: seconds must be positive")
            values[operation_id].append(seconds)
    missing = {item: minimum - len(rows) for item, rows in values.items() if len(rows) < minimum}
    if missing:
        raise ValueError(f"valid calibration trials are incomplete: {missing}")
    return {
        operation_id: {
            "valid_trials": len(rows),
            "median_seconds": round(median(rows), 6),
            "raw_seconds": rows,
        }
        for operation_id, rows in sorted(values.items())
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--protocol",
        type=Path,
        default=Path("benchmarks/phase16_operation_time_calibration_protocol_v1.json"),
    )
    parser.add_argument(
        "--output-dir", type=Path, default=Path("artifacts/phase16-operation-calibration")
    )
    parser.add_argument("--measurements", type=Path)
    parser.add_argument("--participant-id")
    parser.add_argument("--powerpoint-version")
    parser.add_argument("--overwrite-template", action="store_true")
    args = parser.parse_args(argv)
    protocol_path = args.protocol.resolve()
    output_dir = args.output_dir.resolve()
    try:
        protocol = _load(protocol_path)
        if protocol.get("schema_version") != "phase16_operation_time_calibration_protocol_v1":
            raise ValueError("unsupported calibration protocol schema")
        if args.measurements is None:
            deck_path = output_dir / "operation_time_calibration.pptx"
            csv_path = output_dir / "operation_time_measurements.csv"
            prepare_calibration_deck(
                protocol, deck_path, overwrite=args.overwrite_template
            )
            prepare_measurement_csv(
                protocol, csv_path, overwrite=args.overwrite_template
            )
            print(json.dumps({"deck": str(deck_path), "measurements": str(csv_path)}, indent=2))
            return 0
        if not args.participant_id or not args.powerpoint_version:
            raise ValueError(
                "--participant-id and --powerpoint-version are required with --measurements"
            )
        measurements_path = args.measurements.resolve()
        deck_path = output_dir / "operation_time_calibration.pptx"
        if not deck_path.is_file():
            raise ValueError(f"calibration deck is missing: {deck_path}")
        operation_times = summarize_measurements(protocol, measurements_path)
        calibration = {
            "schema_version": CALIBRATION_SCHEMA,
            "status": "completed",
            "completed_at": datetime.now(timezone.utc).isoformat(),
            "participant_id": args.participant_id,
            "powerpoint_version": args.powerpoint_version,
            "protocol_path": str(protocol_path),
            "protocol_sha256": sha256_file(protocol_path),
            "deck_path": str(deck_path),
            "deck_sha256": sha256_file(deck_path),
            "measurements_path": str(measurements_path),
            "measurements_sha256": sha256_file(measurements_path),
            "minimum_repetitions_per_operation": protocol[
                "minimum_repetitions_per_operation"
            ],
            "operation_times": operation_times,
        }
        target = output_dir / "operation_time_calibration.json"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(calibration, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(calibration, indent=2))
        return 0
    except (OSError, ValueError, KeyError, TypeError, json.JSONDecodeError) as error:
        parser.error(str(error))
    return 2


if __name__ == "__main__":
    raise SystemExit(main())

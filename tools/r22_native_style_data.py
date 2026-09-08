"""Native appearance pairs of frozen R19 sources; geometry and text remain fixed."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import random
import zipfile
from xml.etree import ElementTree as ET
import numpy as np
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from r16_ppt_primitive_data import NS
from r10_owner_region_p2_audit import sha


def geometry(path):
    with zipfile.ZipFile(path) as z:root=ET.fromstring(z.read('ppt/slides/slide1.xml'))
    result=[]
    for v in root.findall('.//p:sp',NS)+root.findall('.//p:cxnSp',NS):
        n=v.find('.//p:cNvPr',NS);x=v.find('p:spPr/a:xfrm',NS);p=v.find('p:spPr/a:prstGeom',NS)
        result.append(dict(id=n.get('id'),name=n.get('name'),xfrm=ET.tostring(x).decode() if x is not None else None,
            preset=ET.tostring(p).decode() if p is not None else None,text=[t.text for t in v.findall('.//a:t',NS)]))
    return result


def gradient(fill,colors,angle):
    fill.gradient()
    for stop,color in zip(fill.gradient_stops,colors,strict=True):stop.color.rgb=RGBColor(*color)
    fill.gradient_angle=angle


def build(source,out):
    out.mkdir(parents=True,exist_ok=False);(out/'pptx').mkdir()
    original=json.loads(source.read_text());assert len(original['records'])==384
    (out/'preregister.json').write_text(json.dumps(dict(seed=2201,source_manifest=str(source.resolve()),source_manifest_sha256=sha(source),
        generator_sha256=sha(__file__),record_seed='SHA256 UTF8 2201:<original_id>, first 8 bytes big endian',
        contract='Native font/size/weight/fill/gradient/stroke/shadow only. NoFill preserved, no xfrm/preset/text/role/layout edits. Exposed paired development data, not independent eval.'),indent=2))
    records=[]
    for old in original['records']:
        assert sha(old['pptx'])==old['pptx_sha256']
        seed=int.from_bytes(hashlib.sha256(('2201:'+old['id']).encode()).digest()[:8],'big');rng=random.Random(seed)
        p=Presentation(old['pptx']);assert len(p.slides)==1;s=p.slides[0];changes=[]
        source_elements={e['id']:e for e in old['elements']};bg=source_elements['background']['color']
        if rng.random()<.6:
            colors=[[max(0,min(255,c+delta)) for c in bg] for delta in [-rng.randint(0,25),rng.randint(0,20)]];angle=rng.choice([0,45,90,135])
            gradient(s.background.fill,colors,angle);changes.append(dict(id='background',gradient=colors,angle=angle))
        for v in s.shapes:
            e=source_elements[v.name];change=dict(id=v.name,source_shape_id=v.shape_id)
            if v.has_text_frame:
                fonts=[]
                for paragraph in v.text_frame.paragraphs:
                    for run in paragraph.runs:
                        scale=rng.uniform(1,1.5);before=run.font.size.pt;name=rng.choice(['Arial','Calibri',run.font.name]);bold=rng.choice([True,True,False,run.font.bold])
                        run.font.size=Pt(before*scale);run.font.name=name;run.font.bold=bold
                        fonts.append(dict(original_size_pt=before,size_pt=run.font.size.pt,scale=scale,name=name,bold=bold))
                change['fonts']=fonts
            if e['type']=='shape' and e['fill'] is not None:
                mode=rng.choice(['original','solid','gradient']);change['fill_mode']=mode
                if mode!='original':
                    colors=[[rng.randint(35,115) if rng.random()<.35 else rng.randint(165,250) for _ in range(3)] for _ in range(2)]
                    if mode=='solid':v.fill.solid();v.fill.fore_color.rgb=RGBColor(*colors[0]);change['fill']=colors[0]
                    else:
                        angle=rng.choice([0,45,90,135]);gradient(v.fill,colors,angle);change.update(gradient=colors,angle=angle)
                if e['role']=='closed_structure' and rng.random()<.5:
                    effects=v._element.spPr.find('{'+NS['a']+'}effectLst')
                    if effects is None:effects=OxmlElement('a:effectLst');v._element.spPr.append(effects)
                    for child in list(effects):effects.remove(child)
                    blur=rng.choice([1,2,3]);distance=rng.choice([1,2,3]);alpha=rng.choice([12000,18000,24000]);direction=rng.choice([2700000,5400000])
                    shadow=OxmlElement('a:outerShdw')
                    for key,value in dict(blurRad=blur*12700,dist=distance*12700,dir=direction,algn='ctr',rotWithShape='0').items():shadow.set(key,str(value))
                    color=OxmlElement('a:srgbClr');color.set('val','000000');opacity=OxmlElement('a:alpha');opacity.set('val',str(alpha));color.append(opacity);shadow.append(color);effects.append(shadow)
                    change['shadow']=dict(blur_pt=blur,distance_pt=distance,alpha=alpha,direction=direction)
            if e['type'] in ['shape','line']:
                width=rng.choice([.25,.5,1,2,3,4]);color=[rng.randint(160,235)]*3 if width<1 else [rng.randint(0,110) for _ in range(3)]
                v.line.width=Pt(width);v.line.color.rgb=RGBColor(*color);change.update(stroke_pt=width,stroke=color)
            changes.append(change)
        identity=old['id'].replace('r19-','r22-',1);path=out/'pptx'/f'{identity}.pptx';p.save(path)
        assert geometry(old['pptx'])==geometry(path),identity
        with zipfile.ZipFile(path) as z:xml=ET.fromstring(z.read('ppt/slides/slide1.xml'))
        shaped={v.find('.//p:cNvPr',NS).get('name'):v for v in xml.findall('.//p:sp',NS)}
        for e in old['elements']:
            if e['type']=='shape' and e['fill'] is None:assert shaped[e['id']].find('p:spPr/a:noFill',NS)is not None
        r=copy.deepcopy(old)
        for key in ['image','image_sha256','render_sha256']:r.pop(key,None)
        r.update(id=identity,original_id=old['id'],original_image=old['image'],original_image_sha256=old['image_sha256'],
            original_pptx=old['pptx'],original_pptx_sha256=old['pptx_sha256'],pptx=str(path.resolve()),pptx_sha256=sha(path),style_seed=seed,style_changes=changes)
        records.append(r)
    assert sha(source)==json.loads((out/'preregister.json').read_text())['source_manifest_sha256']
    (out/'source-records.json').write_text(json.dumps(records,indent=2));print('384 appearance source pairs; prereg',sha(out/'preregister.json'))


def finalize(out):
    prereg=json.loads((out/'preregister.json').read_text());assert sha(prereg['source_manifest'])==prereg['source_manifest_sha256']
    old={r['id']:r for r in json.loads(Path(prereg['source_manifest']).read_text())['records']}
    records=json.loads((out/'source-records.json').read_text());m=out/'renders/powerpoint_render_manifest.json'
    receipts={r['id']:r for r in json.loads(m.read_text(encoding='utf-8-sig'))['records']};assert len(receipts)==384
    (out/'rgb').mkdir();seen=set()
    for r in records:
        base=old[r['original_id']]
        assert r['objects']==base['objects'] and r['elements']==base['elements']
        assert all(r[k]==base[k] for k in ['split','stage','family','layout_family','node_count','container_count'])
        assert sha(r['original_pptx'])==r['original_pptx_sha256'] and sha(r['original_image'])==r['original_image_sha256']
        assert geometry(r['original_pptx'])==geometry(r['pptx'])
        q=receipts[r['id']];assert sha(q['pptx_path'])==q['pptx_sha256']==r['pptx_sha256'];assert sha(q['render_path'])==q['render_sha256']
        with Image.open(q['render_path']) as im:
            assert 'transparency' not in im.info;a=np.array(im.convert('RGB'));assert a.shape==(600,960,3)
        path=out/'rgb'/f'{r["id"]}.png';Image.fromarray(a).save(path);assert np.array_equal(a,np.array(Image.open(path)))
        digest=sha(path);assert digest not in seen;seen.add(digest);r.update(image=str(path.resolve()),image_sha256=digest,render_sha256=q['render_sha256'])
    manifest=dict(records=records,classmapping=json.loads(Path(prereg['source_manifest']).read_text())['classmapping'],source_size=[1200,750],render_size=[960,600],
        source_manifest=prereg['source_manifest'],source_manifest_sha256=prereg['source_manifest_sha256'],generator_sha256=sha(__file__),
        preregister_sha256=sha(out/'preregister.json'),com_manifest_sha256=sha(m),
        scope='R19 geometry/words/roles unchanged appearance pairs. objects/elements preserve original style metadata; style_changes overrides describe actual new native appearance. Exposed paired development, not independent evaluation.')
    (out/'dataset_manifest.json').write_text(json.dumps(manifest,indent=2));print('384 verified appearance RGBs; manifest',sha(out/'dataset_manifest.json'))


def qa(out):
    records=json.loads((out/'dataset_manifest.json').read_text())['records']
    for family in sorted({r['family'] for r in records}):
        r=next(r for r in records if r['split']=='train' and r['family']==family)
        board=Image.new('RGB',(1920,625),'white');draw=ImageDraw.Draw(board);draw.text((4,3),r['original_id']+' | R19 / R22 appearance',fill='black')
        board.paste(Image.open(r['original_image']),(0,25));board.paste(Image.open(r['image']),(960,25));board.save(out/f'train-pair-{family}.png')
    print('6 train pair boards; no eval visuals')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['build','finalize','qa']);p.add_argument('--source',type=Path);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    if a.mode=='build':build(a.source,a.out)
    elif a.mode=='finalize':finalize(a.out)
    else:qa(a.out)

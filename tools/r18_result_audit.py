"""Verify frozen evidence and diagnose the exposed original only after RGB inference."""
import argparse
from collections import Counter
import json
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from r18_candidate_validity import ROOT, OCC, digest, save, verify, geometry, rgb, propose, load_model


def run(output):
    output.mkdir(parents=True,exist_ok=False);base=ROOT/'artifacts/region-image-only-r18'
    results=json.loads((base/'eval-v1/results.json').read_text());verify(json.loads((base/'eval-v1/lock.json').read_text()))
    assert results['lock_sha256']==digest(base/'eval-v1/lock.json')
    native=json.loads((base/'com-v1/audit.json').read_text())
    assert native['results_sha256']==digest(base/'eval-v1/results.json') and native['checked_pptx']==64
    assert native['code_sha256']==digest(ROOT/'tools/r18_visual_audit.py')
    assert native['receipt_sha256']==digest(base/'com-v1/renders/powerpoint_render_manifest.json')
    for board in native['boards']:assert digest(board['path'])==board['sha256']
    for batch in ('com-v1','original-com-v1'):
        receipts=json.loads((base/batch/'renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))['records']
        for r in receipts:
            assert digest(r['pptx_path'])==r['pptx_sha256'] and digest(r['render_path'])==r['render_sha256']
    single=json.loads((base/'rgb-only-v1/result.json').read_text());verify(single)
    for arm,row in single['arms'].items():
        expected=next(r for r in results['arms'][arm]['rows'] if r['id']=='r18-scene-eval-000')
        assert row['objects']==expected['objects'] and row['all_candidates']==expected['all_candidates']
    original=json.loads((base/'original-v1/result.json').read_text());verify(original)
    choices=json.loads((base/'eval-v1/calibration.json').read_text())
    assert original['thresholds']==single['thresholds']=={v:r['threshold'] for v,r in choices.items()}
    source=ROOT/'artifacts/region-image-only-r10/complex-inputs/original-rgb.png'
    image=rgb(source);size=(image.shape[1],image.shape[0]);views=propose(load_model(OCC),image)
    # Manual geometry is posthoc diagnosis, never an inference or threshold input.
    ledger_path=ROOT/'docs/r15_visual_loss_ledger.json';ledger=json.loads(ledger_path.read_text())
    expected=[dict(kind=1,bbox=[v*size[i%2]/ledger['source_display_size'][i%2] for i,v in enumerate(n['box'])]) for n in ledger['nodes']]
    receipts={r['id']:r for r in receipts};assert len(receipts)==4
    board=Image.new('RGB',(1800,734),'white');draw=ImageDraw.Draw(board)
    board.paste(Image.fromarray(image).resize((600,327)),(0,40));draw.text((4,4),'Exposed original RGB',fill='black')
    diagnostic={}
    for index,(arm,row) in enumerate(original['arms'].items(),1):
        receipt=receipts[arm];assert receipt['pptx_sha256']==row['pptx_sha256']==digest(base/'original-v1'/arm/'result.pptx')
        deck=Presentation(receipt['pptx_path']);assert len(deck.slides)==1 and len(deck.slides[0].shapes)==len(row['objects'])
        assert all(s.auto_shape_type is not None for s in deck.slides[0].shapes)
        candidates=row['all_candidates'];view=arm.split('_')[0]
        actual,crops=views[view];assert [(o['id'],o['bbox']) for o in actual]==[(o['id'],o['bbox']) for o in candidates]
        kept=[o for o in candidates if arm.endswith('_raw') or o['validity']>=original['thresholds'][view]]
        diagnostic[arm]=dict(candidates=len(candidates),emitted=len(row['objects']),kinds=Counter(o['kind'] for o in row['objects']),
            candidate_geometry=geometry(candidates,expected),validity_geometry=geometry(kept,expected),final_geometry=geometry(row['objects'],expected))
        x=index%3*600;y=index//3*367
        with Image.open(receipt['render_path']) as im:board.paste(im.convert('RGB').resize((600,327)),(x,y+40))
        draw.text((x+4,y+4),f'{arm}: {len(row["objects"])} native shapes',fill='black')
        if arm.endswith('_raw'):
            sheet=Image.new('RGB',(768,((len(crops)+5)//6)*90),'white');labels=ImageDraw.Draw(sheet)
            for i,(o,crop) in enumerate(zip(candidates,crops,strict=True)):
                cx=i%6*128;cy=i//6*90;labels.text((cx+2,cy),f'{o["id"]} v={o["validity"]:.3f}',fill='black')
                for ch in range(2):sheet.paste(Image.fromarray((crop[ch].numpy()*255).astype('uint8')).convert('RGB'),(cx+ch*64,cy+24))
            sheet.save(output/(view+'-candidate-crops.png'))
    board.save(output/'original-comparison.png')
    gates={a:bool(r['calibration']['gate_pass'] and r['precision']==1 and r['recall']>=.8 and r['negative_fp']==0
        and all(v['recall']>=.5 for f,v in r['per_family'].items() if f!='hard_negative')) for a,r in results['arms'].items()}
    verify(original);verify(json.loads((base/'eval-v1/lock.json').read_text()))
    save(output/'audit.json',dict(research_gate=gates,production_approved=False,rgb_only_equal=True,original=diagnostic,
        source_results_sha256=digest(base/'eval-v1/results.json'),com_audit_sha256=digest(base/'com-v1/audit.json'),
        original_prediction_sha256=digest(base/'original-v1/result.json'),original_receipt_sha256=digest(base/'original-com-v1/renders/powerpoint_render_manifest.json'),
        ledger_sha256=digest(ledger_path),code_sha256=digest(__file__),
        scope='Manual 17-node geometry is exposed development diagnosis, not independent GT; no post-eval tuning.'))
    print(json.dumps(dict(research_gate=gates,original=diagnostic),indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,required=True);run(parser.parse_args().output)

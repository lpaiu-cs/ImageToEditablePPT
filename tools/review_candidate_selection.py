"""R2 native-render comparison and frozen output audit; no candidate decisions."""
import argparse
import json
from pathlib import Path

from PIL import Image,ImageDraw,ImageFont,ImageOps
from pptx import Presentation

from image_to_editable_ppt.research.native_emit import load_scene,normalize_scene
from image_to_editable_ppt.research.region_first import digest,pixel_box


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir',default='artifacts/region-first-r2/run1')
    parser.add_argument('--render-dir',default='artifacts/region-first-r2/run1-renders')
    parser.add_argument('--output-dir',required=True)
    parser.add_argument('--visual-review')
    parser.add_argument('--evidence-output')
    parser.add_argument('--diagnostic-manifest')
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]
    output=root/args.output_dir
    output.mkdir(parents=True,exist_ok=False)
    run=root/args.run_dir
    renders=root/args.render_dir
    result=json.loads((run/'results.json').read_text(encoding='utf-8'))
    manifest_path=root/'benchmarks/region_first_r2_inputs_v1.json'
    assert digest(manifest_path)==result['manifest_sha256']
    inputs=json.loads(manifest_path.read_text())['inputs']
    rm=json.loads((renders/'powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
    rm={r['id']:r for r in rm['records']}
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
    sheets,audits=[],[]
    for item in inputs:
        name=item['sample']
        source=root/item['source_path']
        assert digest(source)==item['source_sha256']
        original=load_scene(root/item['scene_path'])
        assert digest(root/item['scene_path'])==item['scene_sha256']
        panels=[]
        for arm in ('source','all_candidates','deduplicate','evidence_select'):
            title={'source':'SOURCE: fixed ROI','all_candidates':'A: all candidates','deduplicate':'B: duplicate removal','evidence_select':'C: source-evidence filter'}[arm]
            path=source if arm=='source' else renders/f'{name}--{arm}.png'
            subtitle=''
            if arm!='source':
                key=f'{name}--{arm}'
                r=next(r for r in result['records'] if r['sample']==name and r['arm']==arm)
                assert digest(path)==rm[key]['render_sha256']
                assert digest(root/r['pptx_path'])==r['pptx_sha256']==rm[key]['pptx_sha256']
                assert digest(root/r['scene_path'])==r['scene_sha256']
                assert digest(root/r['overlay_path'])==r['overlay_sha256']
                selected=load_scene(root/r['scene_path'])
                normalized=normalize_scene(selected)
                deck=Presentation(root/r['pptx_path'])
                objects={s.name:s for s in deck.slides[0].shapes}
                owner_by_id={str(s.shape_id):s.name for s in objects.values()}
                bindings=[]
                for c in selected.connectors:
                    links=objects[c.id]._element.xpath('.//a:stCxn | .//a:endCxn')
                    bindings.append([owner_by_id[b.get('id')] for b in links]==[c.source_owner_id,c.target_owner_id])
                check=dict(sample=name,arm=arm,source_fields_preserved=all(s in original.shapes for s in selected.shapes),
                    connectors_are_exact_subset=all(c in original.connectors for c in selected.connectors),
                    pptx_shape_ids_exact=set(objects)=={s.id for s in selected.shapes}|{c.id for c in selected.connectors},
                    labels_preserved=all(objects[s.id].text==(s.label or '') for s in selected.shapes),
                    bound_owner_pairs_preserved=all(bindings),raster_count=len(deck.slides[0]._element.xpath('.//p:pic')),
                    max_box_error_px=max((abs(a-b) for s in normalized.shapes for a,b in zip(
                        (objects[s.id].left/9525,objects[s.id].top/9525,objects[s.id].width/9525,objects[s.id].height/9525),
                        (s.bbox.x0,s.bbox.y0,s.bbox.width,s.bbox.height))),default=0))
                assert all(check[k] for k in ('source_fields_preserved','connectors_are_exact_subset','pptx_shape_ids_exact','labels_preserved','bound_owner_pairs_preserved'))
                audits.append(check)
                m=r['metrics']
                subtitle=f"Objects {m['matched_objects']}/{m['target_objects']}, extra {len(m['unmatched_local_objects'])}, edges {m['matched_directed_edges']}/{m['target_edges']}"
            with Image.open(path) as im:
                crop=im.convert('RGB').crop(pixel_box(item['roi'],im.size))
            height=max(160,min(700,round(470*crop.height/crop.width)))
            panel=Image.new('RGB',(490,height+80),'#f1f5f9')
            thumb=ImageOps.contain(crop,(470,height))
            panel.paste(thumb,((490-thumb.width)//2,65+(height-thumb.height)//2))
            draw=ImageDraw.Draw(panel)
            draw.text((10,8),title,font=font,fill='black')
            draw.text((10,35),subtitle,font=font,fill='#334155')
            panels.append(panel)
        sheet=Image.new('RGB',(1960,max(p.height for p in panels)),'white')
        for i,p in enumerate(panels): sheet.paste(p,(490*i,0))
        path=output/f'{name}-comparison.png'
        sheet.save(path)
        sheets.append(dict(sample=name,path=str(path.relative_to(root)),sha256=digest(path)))
    evidence=dict(status='awaiting_direct_visual_review',results=result,render_records=list(rm.values()),
        run_sha256=digest(run/'results.json'),render_manifest_sha256=digest(renders/'powerpoint_render_manifest.json'),
        review_code_sha256=digest(Path(__file__)),sheets=sheets,invariants=audits)
    if args.visual_review:
        evidence['visual_review']=json.loads((root/args.visual_review).read_text(encoding='utf-8'))
        evidence['status']='reviewed_B_partial_gain_C_failed_not_product_acceptance'
    if args.diagnostic_manifest:
        dp=root/args.diagnostic_manifest
        diagnostic=json.loads(dp.read_text())
        for r in diagnostic['records']:
            assert digest(root/r['path'])==r['sha256']
        evidence['post_result_diagnostic']=diagnostic
        evidence['diagnostic_manifest_sha256']=digest(dp)
        evidence['diagnostic_code_sha256']=digest(root/'tools/inspect_r2_rejections.py')
    (output/'evidence.json').write_text(json.dumps(evidence,indent=2,ensure_ascii=False),encoding='utf-8')
    if args.evidence_output:
        with (root/args.evidence_output).open('x',encoding='utf-8') as f:
            json.dump(evidence,f,indent=2,ensure_ascii=False)
    print('Verified 18 native outputs; built six comparisons')


if __name__=='__main__': main()

"""Audit R3 path/shape evidence and compare source with four native-output arms."""
import argparse
import json
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont,ImageOps
from pptx import Presentation
from image_to_editable_ppt.research.native_emit import load_scene,normalize_scene
from image_to_editable_ppt.research.region_first import digest,pixel_box


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir',default='artifacts/region-first-r3/run1')
    parser.add_argument('--render-dir',default='artifacts/region-first-r3/run1-renders')
    parser.add_argument('--output-dir',required=True)
    parser.add_argument('--visual-review')
    parser.add_argument('--evidence-output')
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]
    run,renders,output=(root/p for p in (args.run_dir,args.render_dir,args.output_dir))
    output.mkdir(parents=True,exist_ok=False)
    result=json.loads((run/'results.json').read_text(encoding='utf-8'))
    inp=root/'benchmarks/region_first_r3_inputs_v1.json'
    assert digest(inp)==result['manifest_sha256']
    inputs={i['sample']:i for i in json.loads(inp.read_text())['inputs']}
    rm=json.loads((renders/'powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
    rm={r['id']:r for r in rm['records']}
    assert len(result['records'])==len(rm)==24
    checks=[]
    for r in result['records']:
        item=inputs[r['sample']]
        for kind in ('selected','oracle','source'):
            assert digest(root/item[kind+'_path'])==item[kind+'_sha256']
        original=load_scene(root/item['oracle_path' if r['arm']=='oracle_trace' else 'selected_path'])
        assert digest(root/r['scene_path'])==r['scene_sha256']
        assert digest(root/r['overlay_path'])==r['overlay_sha256']
        if r.get('mask_path'):
            assert digest(root/r['mask_path'])==r['mask_sha256']
        key=f"{r['sample']}--{r['arm']}"
        assert digest(root/r['pptx_path'])==r['pptx_sha256']==rm[key]['pptx_sha256']
        assert digest(renders/f'{key}.png')==rm[key]['render_sha256']
        selected=load_scene(root/r['scene_path'])
        normalized=normalize_scene(selected)
        deck=Presentation(root/r['pptx_path'])
        objects={s.name:s for s in deck.slides[0].shapes}
        by_id={str(s.shape_id):s.name for s in objects.values()}
        bindings=[]
        for c in selected.connectors:
            links=objects[c.id]._element.xpath('.//a:stCxn | .//a:endCxn')
            bindings.append([by_id[b.get('id')] for b in links]==[c.source_owner_id,c.target_owner_id])
        check=dict(sample=r['sample'],arm=r['arm'],primitives_preserved=all(s in original.shapes for s in selected.shapes),
            connectors_preserved=all(c.source_owner_id in {s.id for s in selected.shapes} and c.target_owner_id in {s.id for s in selected.shapes} for c in selected.connectors),
            native_ids_exact=set(objects)=={s.id for s in selected.shapes}|{c.id for c in selected.connectors},
            labels_preserved=all(objects[s.id].text==(s.label or '') for s in selected.shapes),
            owner_bindings_preserved=all(bindings),raster_count=len(deck.slides[0]._element.xpath('.//p:pic')),
            max_box_error_px=max((abs(a-b) for s in normalized.shapes for a,b in zip(
                (objects[s.id].left/9525,objects[s.id].top/9525,objects[s.id].width/9525,objects[s.id].height/9525),
                (s.bbox.x0,s.bbox.y0,s.bbox.width,s.bbox.height))),default=0))
        assert all(check[k] for k in ('primitives_preserved','connectors_preserved','native_ids_exact','labels_preserved','owner_bindings_preserved'))
        assert check['raster_count']==0 and check['max_box_error_px']<.001
        checks.append(check)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
    groups=dict(trace=('source','baseline','strict_trace','transparent_gap_trace','oracle_trace'))
    sheets=[]
    for sample,item in inputs.items():
        for group,arms in groups.items():
            panels=[]
            for arm in arms:
                path=root/item['source_path'] if arm=='source' else renders/f'{sample}--{arm}.png'
                with Image.open(path) as im:
                    crop=im.convert('RGB').crop(pixel_box(item['roi'],im.size))
                height=max(150,min(600,round(470*crop.height/crop.width)))
                panel=Image.new('RGB',(490,height+75),'#f1f5f9')
                thumb=ImageOps.contain(crop,(470,height))
                panel.paste(thumb,((490-thumb.width)//2,60+(height-thumb.height)//2))
                draw=ImageDraw.Draw(panel)
                draw.text((10,8),arm,font=font,fill='black')
                if arm!='source':
                    m=next(r['metrics'] for r in result['records'] if r['sample']==sample and r['arm']==arm)
                    draw.text((10,34),f"Objects {m['matched_objects']}/{m['target_objects']}, extra {len(m['unmatched_local_objects'])}, edges {m['matched_directed_edges']}/{m['target_edges']}",font=font,fill='#334155')
                panels.append(panel)
            sheet=Image.new('RGB',(490*len(panels),max(p.height for p in panels)),'white')
            for i,p in enumerate(panels): sheet.paste(p,(490*i,0))
            path=output/f'{sample}-{group}.png'
            sheet.save(path)
            sheets.append(dict(sample=sample,group=group,path=str(path.relative_to(root)),sha256=digest(path)))
    evidence=dict(status='awaiting_visual_review',results=result,render_records=list(rm.values()),invariants=checks,sheets=sheets,
        results_sha256=digest(run/'results.json'),render_manifest_sha256=digest(renders/'powerpoint_render_manifest.json'),
        review_code_sha256=digest(Path(__file__)))
    if args.visual_review:
        evidence['visual_review']=json.loads((root/args.visual_review).read_text(encoding='utf-8'))
        evidence['status']='reviewed_development_gain_not_product_acceptance'
    (output/'evidence.json').write_text(json.dumps(evidence,indent=2),encoding='utf-8')
    if args.evidence_output:
        with (root/args.evidence_output).open('x',encoding='utf-8') as f: json.dump(evidence,f,indent=2)
    print('24 PPTX audited; 6 comparison sheets built')


if __name__=='__main__': main()

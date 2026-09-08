"""격리된 RGB 추론, 별도 GT 평가, 제한된 detector pilot 연구 도구."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
sys.path.insert(0, str(ROOT / 'tools'))

# a00d61c label order; 3ae7caf inserted ELLIPSE at current index 3.
LEGACY_TO_CURRENT = {1:1, 2:2, 3:4, 4:5, 5:6, 6:7}


def map_legacy_labels(labels):
    return labels.clone().apply_(lambda value: LEGACY_TO_CURRENT[value])


def map_current_targets(labels):
    reverse = {v:k for k,v in LEGACY_TO_CURRENT.items()}
    if any(int(value) not in reverse for value in labels):
        raise ValueError('Legacy checkpoint cannot learn ellipse or unknown labels')
    return labels.clone().apply_(lambda value: reverse[value])


def install_gt_guard():
    def deny_gt(event, values):
        if event == 'open' and isinstance(values[0], (str, bytes)):
            name = str(values[0]).lower().replace('\\', '/')
            if ('workbench-ml/ds-' in name or '/benchmarks/' in name) and not name.endswith('.png'):
                raise PermissionError('GT access forbidden in inference: ' + name)
    sys.addaudithook(deny_gt)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError(path)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding='utf-8')


def prepare(args):
    from PIL import Image
    import numpy as np
    root = args.data
    manifest = root / 'dataset_manifest.json'
    doc = json.loads(manifest.read_text())
    groups = {name: defaultdict(list) for name in ('rgb', 'geometry')}
    rows, tiny = [], []
    for s in doc['samples']:
        image = (root / s['image']).resolve()
        annotation = (root / s['annotation']).resolve()
        gt = json.loads(annotation.read_text())
        geom = {k: sorted([(x['kind'], x['bbox']) for x in gt['primitive_scene'].get(k, [])], key=str) for k in ('nodes', 'containers')}
        with Image.open(image) as im:
            rgb = im.convert('RGB')
            rh = hashlib.sha256(str(rgb.size).encode() + rgb.tobytes()).hexdigest()
            tiny.append(np.asarray(rgb.resize((32, 18)), dtype=np.float32).reshape(-1) / 255)
        gh = hashlib.sha256(json.dumps(geom, sort_keys=True).encode()).hexdigest()
        row = dict(id=s['id'], split=s['split'], family=s['family'], image=str(image), annotation=str(annotation), image_sha256=sha(image), annotation_sha256=sha(annotation), rgb_sha256=rh, geometry_sha256=gh)
        if s.get('pptx'):
            row['pptx_sha256'] = sha(root / s['pptx'])
        rows.append(row)
        for name, digest in [('rgb', rh), ('geometry', gh)]:
            groups[name][digest].append(row)
    collisions = {name: [[{'id': r['id'], 'split': r['split']} for r in rs] for rs in group.values() if len({r['split'] for r in rs}) > 1] for name, group in groups.items()}
    near = []
    for i, row in enumerate(rows):
        for j in range(i):
            if row['split'] != rows[j]['split']:
                distance = float(np.mean(np.abs(tiny[i] - tiny[j])))
                if distance < .01:
                    near.append(dict(first=row['id'], second=rows[j]['id'], mae=distance))
    source_checks = []
    for s in doc['generator']['sources']:
        actual = sha(root / s['manifest'])
        source_checks.append(dict(source=s, actual_sha256=actual, matches=actual == s['sha256']))
    save(args.out / 'audit.json', dict(manifest_sha256=sha(manifest), sources=source_checks, split_counts=dict(Counter(r['split'] for r in rows)), pptx_count=sum('pptx_sha256' in r for r in rows), cross_split_exact=collisions, near_duplicate_candidates=near, rows=rows))
    # Validation only: test and Phase16 holdout remain untouched.
    by_id = {r['id']:r for r in rows}
    excluded = {s for pair in near if any(by_id[pair[k]]['split'] == 'train' for k in ('first','second')) for s in (pair['first'],pair['second'])}
    excluded.update(r['id'] for rs in groups['geometry'].values() if any(r['split'] == 'train' for r in rs) for r in rs)
    selected = sorted([r for r in rows if r['split'] == 'val' and 'pptx_sha256' in r and r['id'] not in excluded], key=lambda r:r['id'])
    import shutil
    rgb_dir = args.out / 'rgb'
    rgb_dir.mkdir()
    rgb_inputs = []
    for r in selected:
        destination = rgb_dir / (r['id'] + '.png')
        shutil.copyfile(r['image'], destination)
        rgb_inputs.append(dict(id=r['id'], image=str(destination.resolve()), image_sha256=r['image_sha256']))
    save(args.out / 'rgb_inputs.json', rgb_inputs)
    save(args.out / 'evaluator_inputs.json', selected)
    print(json.dumps(dict(samples=len(rows), selected=len(selected), exact={k:len(v) for k,v in collisions.items()}, near=len(near))))


def infer(args):
    import torch
    import numpy as np
    from PIL import Image
    from image_to_editable_ppt.ml.lightning_module import DetectorLightningModule
    from image_to_editable_ppt.ml.dataset import to_rgb_chw_tensor
    torch.set_num_threads(8)
    inputs = json.loads(args.inputs.read_text())
    assert all(set(r) == {'id', 'image', 'image_sha256'} for r in inputs)
    install_gt_guard()
    model = DetectorLightningModule.load_from_checkpoint(args.checkpoint, map_location='cpu').eval()
    assert model.hparams.num_classes == 7, 'This pilot requires audited v13 legacy class schema'
    start = time.monotonic()
    predictions = []
    with torch.inference_mode():
        for r in inputs:
            assert sha(r['image']) == r['image_sha256']
            with Image.open(r['image']) as im:
                tensor = to_rgb_chw_tensor(np.asarray(im.convert('RGB')))
            p = model([tensor])[0]
            p['labels'] = map_legacy_labels(p['labels'])
            predictions.append(dict(id=r['id'], source_sha256=r['image_sha256'], image_size=[tensor.shape[2], tensor.shape[1]], coordinate_system='original_image_pixels', prediction={k:v.tolist() for k,v in p.items()}))
            print(r['id'], flush=True)
    save(args.out, dict(checkpoint_sha256=sha(args.checkpoint), inputs_sha256=sha(args.inputs), label_adapter=LEGACY_TO_CURRENT, seconds=time.monotonic()-start, predictions=predictions))


def evaluate(args):
    import torch
    from image_to_editable_ppt.ml.dataset import load_annotation_document, document_to_detection_target
    from evaluate_detector_checkpoint import _category_counts, _metrics
    inputs = {r['id']:r for r in json.loads(args.inputs.read_text())}
    prediction = json.loads(args.predictions.read_text())
    assert prediction.get('label_adapter') == {str(k):v for k,v in LEGACY_TO_CURRENT.items()}
    totals = {'nodes':[0,0,0], 'containers':[0,0,0]}
    per_sample = []
    for p in prediction['predictions']:
        row = inputs[p['id']]
        assert sha(row['annotation']) == row['annotation_sha256']
        gt = load_annotation_document(Path(row['annotation']))
        boxes, labels = document_to_detection_target(gt)
        pred = {k:torch.tensor(v, dtype=torch.int64 if k == 'labels' else torch.float32) for k,v in p['prediction'].items()}
        metrics = {}
        for category in totals:
            counts = _category_counts(pred, dict(boxes=boxes, labels=labels), score_threshold=.5, iou_threshold=.5, cross_class_nms_iou=.7, node_containment_threshold=.85, node_category=category == 'nodes')
            totals[category] = [a+b for a,b in zip(totals[category], counts)]
            metrics[category] = _metrics(true_positive=counts[0], prediction_count=counts[1], reference_count=counts[2])
        per_sample.append(dict(id=p['id'], metrics=metrics))
    save(args.out, dict(prediction_sha256=sha(args.predictions), evaluator_inputs_sha256=sha(args.inputs), settings=dict(score=.5,iou=.5,nms=.7,containment=.85), metrics={k:_metrics(true_positive=v[0], prediction_count=v[1], reference_count=v[2]) for k,v in totals.items()}, samples=per_sample, labels='not evaluated by detector', connections='not evaluated by detector'))
    print(args.out.read_text())


def pilot(args):
    import torch
    from image_to_editable_ppt.ml.dataset import DetectorTorchDataset
    from image_to_editable_ppt.ml.lightning_module import DetectorLightningModule
    torch.set_num_threads(8)
    torch.manual_seed(1309)
    model = DetectorLightningModule.load_from_checkpoint(args.checkpoint, map_location='cpu')
    assert model.hparams.num_classes == 7
    dataset = DetectorTorchDataset(args.data, split='train')
    optimizer = torch.optim.AdamW(model.parameters(), lr=.00003)
    indices = torch.randperm(len(dataset)).tolist()[:12]
    start, losses, used = time.monotonic(), [], []
    for i in indices:
        if time.monotonic()-start > 180:
            break
        model.model.train()
        image, target = dataset[i]
        target['labels'] = map_current_targets(target['labels'])
        optimizer.zero_grad()
        loss = sum(model.model([image], [target]).values())
        loss.backward()
        optimizer.step()
        losses.append(float(loss.detach()))
        used.append(dataset.samples[i].sample_id)
        print(len(losses), losses[-1], flush=True)
    args.out.mkdir(parents=True, exist_ok=False)
    destination = args.out / 'pilot.ckpt'
    import lightning
    torch.save({'state_dict':model.state_dict(), 'hyper_parameters':dict(model.hparams), 'pytorch-lightning_version':lightning.__version__}, destination)
    save(args.out / 'run.json', dict(parent_sha256=sha(args.checkpoint), checkpoint_sha256=sha(destination), data_manifest_sha256=sha(args.data / 'dataset_manifest.json'), seed=1309, max_steps=12, time_limit_seconds=180, seconds=time.monotonic()-start, learning_rate=.00003, batch_size=1, sample_ids=used, losses=losses, optimizer='fresh AdamW; weights-only continuation', device='cpu', selection='fixed final step; no validation selection'))


def native(args):
    import torch
    import cv2
    from image_to_editable_ppt.ml.slide_ir_provider import MLSlideIRProvider, _DETECTOR_CACHE
    from image_to_editable_ppt.ml.convert_to_pptx import convert_one
    torch.set_num_threads(8)
    cv2.setNumThreads(1)
    rows = json.loads(args.inputs.read_text())
    doc = json.loads(args.predictions.read_text())
    assert doc['inputs_sha256'] == sha(args.inputs)
    assert doc['label_adapter'] == {str(k):v for k,v in LEGACY_TO_CURRENT.items()}
    predictions = {p['id']:p for p in doc['predictions']}
    install_gt_guard()
    args.out.mkdir(parents=True, exist_ok=False)
    records = []
    for index in (0, 6, 12):
        row = rows[index]
        assert sha(row['image']) == row['image_sha256']
        pred = {k:torch.tensor(v, dtype=torch.int64 if k == 'labels' else torch.float32) for k,v in predictions[row['id']]['prediction'].items()}
        _DETECTOR_CACHE['frozen_rgb_predictions'] = lambda images: [pred]
        provider = MLSlideIRProvider(detector_checkpoint='frozen_rgb_predictions', image_id='rgb')
        config = argparse.Namespace(width=0, no_ocr=False, no_style=False, region=None)
        output = args.out / f'{index:03d}.pptx'
        result = convert_one(Path(row['image']), output, args=config, provider=provider)
        records.append(dict(id=row['id'], source_sha256=row['image_sha256'], output_sha256=sha(output), **result))
    save(args.out / 'native.json', dict(predictions_sha256=sha(args.predictions), settings='fixed RGB predictions, width0, no OOD gate, no learned family, morphological connectors, OCR and style enabled; research diagnostic', records=records))


def contact(args):
    from PIL import Image, ImageDraw
    rows = json.loads((args.data / 'rgb_inputs.json').read_text())
    sheet = Image.new('RGB', (1920,1140), '#dddddd')
    draw = ImageDraw.Draw(sheet)
    for row, index in enumerate((0,6,12)):
        for col, variant in enumerate(('source','native-baseline','native-candidate')):
            path = rows[index]['image'] if variant == 'source' else args.data / variant / 'render-v2' / f'{index:03d}.png'
            with Image.open(path) as im:
                sheet.paste(im.convert('RGB').resize((640,360)), (col*640,row*380+20))
            draw.text((col*640+8,row*380+3), variant + ' ' + rows[index]['id'], fill='black')
    if args.out.exists():
        raise FileExistsError(args.out)
    sheet.save(args.out)


def freeze(args):
    from pptx import Presentation
    root = args.data
    audit = json.loads((root / 'audit.json').read_text())
    file_names = ['rgb_inputs.json','evaluator_inputs.json','audit.json','baseline_legacy_predictions.json','baseline_legacy_metrics.json','candidate_predictions.json','candidate_metrics.json','pilot-legacy/run.json','native-render.json','native_contact.png']
    native = []
    for variant in ('native-baseline','native-candidate'):
        for path in sorted((root / variant).glob('*.pptx')):
            shapes = list(Presentation(path).slides[0].shapes)
            native.append(dict(path=str(path), sha256=sha(path), shapes=len(shapes), picture_shapes=sum(int(s.shape_type) == 13 for s in shapes), texts=[s.text for s in shapes if s.has_text_frame and s.text.strip()], empty_text_frames=sum(s.has_text_frame and not s.text.strip() for s in shapes)))
    save(args.out, dict(schema='ml_image_only_pilot_v1', base_git='d5373b1', code_sha256=sha(__file__), class_map=LEGACY_TO_CURRENT, class_map_sha256=hashlib.sha256(json.dumps(LEGACY_TO_CURRENT,sort_keys=True).encode()).hexdigest(), class_map_provenance='inferred from a00d61c enum and 3ae7caf insertion; checkpoint lacks explicit class names/code SHA', artifact_root=str(root.resolve()), artifacts={name:sha(root / name) for name in file_names}, sources=audit['sources'], split_counts=audit['split_counts'], exact_cross_split=audit['cross_split_exact'], near_candidate_count=len(audit['near_duplicate_candidates']), baseline=json.loads((root/'baseline_legacy_metrics.json').read_text())['metrics'], candidate=json.loads((root/'candidate_metrics.json').read_text())['metrics'], native=native, verdict='candidate not adopted; synthetic validation diagnosis only', invalid_outputs=['baseline_predictions.json','baseline_metrics.json']))


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('stage', choices=['prepare', 'infer', 'evaluate', 'pilot', 'native', 'contact', 'freeze'])
    for name in ['data','inputs','checkpoint','predictions','out']:
        p.add_argument('--'+name, type=Path)
    args = p.parse_args()
    globals()[args.stage](args)

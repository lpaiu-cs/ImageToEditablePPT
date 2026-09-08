"""동결 R7의 실제 PowerPoint 렌더와 추적 단계 오류를 분리하여 검토한다."""
import json
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from run_region_r7 import ROOT, BASE, LOCK, ANNOTATIONS, digest, save, verify_lock
from image_to_editable_ppt.research.accepted_owners import load_annotation


def main():
    verify_lock()
    results = json.loads((BASE / 'run/results.json').read_text(encoding='utf-8'))
    native = json.loads((BASE / 'renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
    native_by_id = {r['id']: r for r in native['records']}
    for record in native['records']:
        assert digest(record['pptx_path']) == record['pptx_sha256']
        assert digest(record['render_path']) == record['render_sha256']
    out = BASE / 'review'
    out.mkdir(parents=True, exist_ok=False)
    font = ImageFont.truetype('C:/Windows/Fonts/malgun.ttf', 20)
    evidence = []
    for record in results['records']:
        data, rgb = load_annotation(ANNOTATIONS / f'{record["id"]}.json', ROOT)
        scope = data['scope']
        original = Image.fromarray(rgb)
        sheet = Image.new('RGB', (1920, 840), '#edf1f5')
        draw = ImageDraw.Draw(sheet)
        panels = [('원본 (명시한 부분 영역)', original.crop(scope))]
        for arm, label in [('r6', 'R6 같은 영역'), ('local-oracle', '정답 역할 승인 (자동 인식 아님)'),
                           ('geometry-oracle', '정답 기하 공급 (자동 인식 아님)')]:
            rendered = Image.open(BASE / 'renders' / f'{record["id"]}-{arm}.png').convert('RGB')
            render_scope = tuple(round(v * (rendered.width / original.width if i % 2 == 0 else rendered.height / original.height))
                                 for i, v in enumerate(scope))
            panels.append((label, rendered.crop(render_scope)))
        for i, (label, picture) in enumerate(panels):
            x, y = (i % 2) * 960, (i // 2) * 420
            draw.text((x+12, y+8), label, fill='black', font=font)
            thumbnail = ImageOps.contain(picture, (936, 370))
            sheet.paste(thumbnail, (x + (960-thumbnail.width)//2, y+42+(370-thumbnail.height)//2))
        sheet_path = out / f'{record["id"]}.png'
        sheet.save(sheet_path)
        graph = original.copy()
        graph_draw = ImageDraw.Draw(graph)
        truth = {(r['source'], r['target']) for r in data['relations']}
        for c in record['arms']['geometry-oracle']['source_scene']['connectors']:
            points = [(p['x'], p['y']) for p in c['path_points']]
            correct = (c['source_owner_id'], c['target_owner_id']) in truth and c['arrowhead_end']
            graph_draw.line(points, fill='#00aa44' if correct else '#ff0033', width=4)
            graph_draw.text(points[len(points)//2], c['source_owner_id']+' > '+c['target_owner_id'],
                            fill='#004400' if correct else '#cc0022', stroke_width=1, stroke_fill='white')
        graph_path = out / f'{record["id"]}-trace.png'
        graph.crop(scope).save(graph_path)
        checks = {}
        for arm, r in record['arms'].items():
            name = f'{record["id"]}-{arm}'
            pp = BASE / 'run' / f'{name}.pptx'
            assert digest(pp) == r['pptx_sha256'] == native_by_id[name]['pptx_sha256']
            deck = Presentation(pp)
            assert len(deck.slides) == 1
            slide = deck.slides[0]
            objects = {s.name: s for s in slide.shapes}
            ids = {str(s.shape_id): s.name for s in slide.shapes}
            scene = r['emitted_scene']
            exact = set(objects) == {s['id'] for s in scene['shapes']} | {c['id'] for c in scene['connectors']}
            bindings = all([ids[b.get('id')] for b in objects[c['id']]._element.xpath('.//a:stCxn | .//a:endCxn')]
                           == [c['source_owner_id'], c['target_owner_id']] for c in scene['connectors'])
            raster_count = len(slide._element.xpath('.//p:pic'))
            assert exact and bindings and raster_count == 0 and len(objects) == len(slide.shapes)
            checks[arm] = dict(ids_exact=exact, owner_bindings=bindings, raster_count=raster_count,
                               pptx_sha256=digest(pp), render_sha256=native_by_id[name]['render_sha256'],
                               traced=r['metrics'], emitted=r['emitted_metrics'])
        evidence.append(dict(id=record['id'], source_sha256=data['source_sha256'],
                             annotation_sha256=record['annotation_sha256'],
                             global_proposals=record['global_proposals'], local_proposals=record['local_proposals'],
                             native=checks, sheet_sha256=digest(sheet_path), trace_sheet_sha256=digest(graph_path)))
    save(out / 'audit.json', evidence)
    save(ROOT / 'benchmarks/region_r7_evidence.json', dict(
        schema='r7_oracle_evidence_v1', population='consumed_development', promotion=False,
        lock_sha256=digest(LOCK), results_sha256=digest(BASE / 'run/results.json'),
        render_manifest_sha256=digest(BASE / 'renders/powerpoint_render_manifest.json'),
        records=evidence))
    print('16 native decks verified; four comparison and four pre-emission trace sheets generated.')


if __name__ == '__main__':
    main()

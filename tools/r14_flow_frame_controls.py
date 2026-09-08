"""Pre-register source PPT flow-frame controls before COM rendering; no inference."""
import argparse
import json
from pathlib import Path
import shutil
import zipfile
from xml.etree import ElementTree as ET
from PIL import Image, ImageDraw, ImageOps
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from r10_classifier_pilot import deck, rect, label, EMU
from r10_owner_region_p2_audit import sha


def build(out):
    out.mkdir(parents=True, exist_ok=False)
    for name in ('assets', 'pptx'):
        (out/name).mkdir()
    source = Path('C:/vs/ImageToEditablePPT/.venv/Lib/site-packages/sklearn/datasets/images')
    for name in ('china.jpg', 'flower.jpg', 'README.txt'):
        shutil.copyfile(source/name, out/'assets'/name)
    icon = Image.new('RGB', (240,160), 'white')
    draw = ImageDraw.Draw(icon)
    draw.ellipse((92,8,148,64), fill=(40,80,160))
    draw.rounded_rectangle((54,72,186,154), radius=24, fill=(40,80,160))
    icon.save(out/'assets/icon.png')
    ImageOps.grayscale(Image.open(source/'flower.jpg')).convert('RGB').save(out/'assets/flower-gray.png')
    specs = [('china-flow','china.jpg','flow'), ('flower-flow','flower.jpg','flow'),
             ('gray-flow','flower-gray.png','flow'), ('icon-flow','icon.png','flow'),
             ('standalone','china.jpg','standalone'), ('one-way','flower.jpg','one-way'),
             ('both-in','china.jpg','both-in'), ('both-out','flower.jpg','both-out'),
             ('same-port','china.jpg','same-port'), ('failed-chain','flower.jpg','failed-chain'),
             ('no-arrows','china.jpg','no-arrows')]
    records = []
    for identity, asset, mode in specs:
        cards = [[280,125,520,325]] if mode!='failed-chain' else [[240,135,370,255],[430,135,560,255]]
        seeds = [] if mode=='standalone' else [dict(id='read',bbox=[40,175,160,235],label='Read'),dict(id='check',bbox=[640,175,760,235],label='Check')]
        edges = []
        def edge(a,b,start,end,arrow=True):
            edges.append(dict(source=a,target=b,start=start,end=end,arrow=arrow))
        if mode in ('flow','one-way','no-arrows'):
            edge('read','card0',[160,205],[280,205],mode!='no-arrows')
            if mode!='one-way':edge('card0','check',[520,205],[640,205],mode!='no-arrows')
        elif mode=='both-in':
            edge('read','card0',[160,205],[280,205]);edge('check','card0',[640,205],[520,205])
        elif mode=='both-out':
            edge('card0','read',[280,205],[160,205]);edge('card0','check',[520,205],[640,205])
        elif mode=='same-port':
            edge('read','card0',[160,205],[280,205]);edge('card0','read',[280,205],[160,205])
        elif mode=='failed-chain':
            edge('read','card0',[160,205],[240,205]);edge('card0','card1',[370,205],[430,205]);edge('card1','check',[560,205],[640,205])
        records.append(dict(id=identity,asset=asset,asset_sha256=sha(out/'assets'/asset),mode=mode,
            seeds=seeds,cards=[dict(id=f'card{i}',bbox=b,caption='Photo' if asset!='icon.png' else 'Input') for i,b in enumerate(cards)],edges=edges,
            expected_frame_ids=['card0'] if mode=='flow' else [],expected_interior_objects=0,
            interpretation='Policy target for observed frame; geometry/OCR/texture/arrow abstentions remain separate measured failures.'))
    lock=dict(records=records,builder_sha256=sha(__file__),helper_sha256=sha(Path(__file__).with_name('r10_classifier_pilot.py')),
        contract='Exposed development controls; no train/eval split, no tuning. Source object roles and flow are evaluator-only. Runtime receives RGB only. Negative expectation applies to candidate cards, not valid labeled seeds. Chain never promotes rescued frames to seeds. Icon texture eligibility is measured, not assumed.')
    (out/'preregister.json').write_text(json.dumps(lock,indent=2))
    receipts=[]
    for r in records:
        p,s=deck()
        for seed in r['seeds']:
            b=seed['bbox'];rect(s,b,(225,235,245),seed['id']);label(s,seed['label'],b[0]+12,b[1]+15,b[2]-b[0]-20)
        for card in r['cards']:
            b=card['bbox'];rect(s,b,name=card['id']);label(s,card['caption'],b[0]+12,b[1]+6,b[2]-b[0]-20)
            pic=s.shapes.add_picture(str(out/'assets'/r['asset']),Emu((b[0]+3)*EMU),Emu((b[1]+38)*EMU),Emu((b[2]-b[0]-6)*EMU),Emu((b[3]-b[1]-41)*EMU))
            pic.name=card['id']+'-nonstructural-interior'
        for i,e in enumerate(r['edges']):
            a,b=e['start'],e['end'];c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Emu(a[0]*EMU),Emu(a[1]*EMU),Emu(b[0]*EMU),Emu(b[1]*EMU))
            c.name=f'flow{i}:{e["source"]}->{e["target"]}';c.line.color.rgb=RGBColor(0,0,0);c.line.width=Pt(2)
            if e['arrow']:
                head=OxmlElement('a:tailEnd');head.set('type','triangle');head.set('w','lg');head.set('len','lg');c._element.spPr.get_or_add_ln().append(head)
        path=out/'pptx'/(r['id']+'.pptx');p.save(path)
        with zipfile.ZipFile(path) as z:root=ET.fromstring(z.read('ppt/slides/slide1.xml'))
        ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        assert len(root.findall('.//p:pic',ns))==len(r['cards'])
        assert len(root.findall('.//p:cxnSp',ns))==len(r['edges'])
        assert len(root.findall('.//a:tailEnd',ns))==sum(e['arrow'] for e in r['edges'])
        receipts.append(dict(id=r['id'],pptx=str(path.resolve()),pptx_sha256=sha(path),picture_count=len(r['cards']),edges=r['edges'],
            objects=[dict(id=n.get('id'),name=n.get('name')) for n in root.findall('.//p:cNvPr',ns)]))
    (out/'source-receipts.json').write_text(json.dumps(receipts,indent=2))
    print('prepared',len(records),'preregister',sha(out/'preregister.json'))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);a=p.parse_args();build(a.out)

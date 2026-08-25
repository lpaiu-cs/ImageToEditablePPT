"""Measure editable-PPTX structure against an optional GT annotation.

This is intentionally an edit-cost proxy, not a pixel-similarity score.  It
counts native objects, real PowerPoint endpoint bindings, recoverable labels,
and the minimum structural additions/deletions a user would still need.
"""
from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


EMU_PER_PIXEL = 9525


def _bbox(shape) -> tuple[float, float, float, float]:
    x0 = shape.left / EMU_PER_PIXEL
    y0 = shape.top / EMU_PER_PIXEL
    return x0, y0, x0 + shape.width / EMU_PER_PIXEL, y0 + shape.height / EMU_PER_PIXEL


def _iou(first, second) -> float:
    ix0 = max(first[0], second[0])
    iy0 = max(first[1], second[1])
    ix1 = min(first[2], second[2])
    iy1 = min(first[3], second[3])
    intersection = max(0.0, ix1 - ix0) * max(0.0, iy1 - iy0)
    first_area = max(0.0, first[2] - first[0]) * max(0.0, first[3] - first[1])
    second_area = max(0.0, second[2] - second[0]) * max(0.0, second[3] - second[1])
    union = first_area + second_area - intersection
    return intersection / union if union > 0 else 0.0


def _greedy_matches(predicted, reference, *, threshold: float = 0.5):
    candidates = sorted(
        (
            (_iou(pred["bbox"], ref["bbox"]), pred_index, ref_index)
            for pred_index, pred in enumerate(predicted)
            for ref_index, ref in enumerate(reference)
        ),
        reverse=True,
    )
    matches: dict[int, int] = {}
    used_reference: set[int] = set()
    for score, pred_index, ref_index in candidates:
        if score < threshold:
            break
        if pred_index in matches or ref_index in used_reference:
            continue
        matches[pred_index] = ref_index
        used_reference.add(ref_index)
    return matches


def _f1_counts(predicted_count: int, reference_count: int, matches: int) -> dict[str, float | int]:
    false_positive = predicted_count - matches
    false_negative = reference_count - matches
    precision = matches / predicted_count if predicted_count else (1.0 if reference_count == 0 else 0.0)
    recall = matches / reference_count if reference_count else (1.0 if predicted_count == 0 else 0.0)
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "tp": matches,
        "fp": false_positive,
        "fn": false_negative,
        "precision": round(precision, 4),
        "recall": round(recall, 4),
        "f1": round(f1, 4),
    }


def _normalize_text(value: str | None) -> str:
    return " ".join((value or "").casefold().split())


def _connected_shape_id(line_shape, tag: str) -> str | None:
    items = line_shape._element.xpath(f".//a:{tag}")
    return items[0].get("id") if items else None


def _has_arrowhead(line_shape, tag: str) -> bool:
    items = line_shape._element.xpath(f".//a:{tag}")
    return any(item.get("type", "none") != "none" for item in items)


def _edge_semantics(source: str, target: str, *, arrowhead_start: bool, arrowhead_end: bool) -> tuple[str, str, str]:
    if arrowhead_start and not arrowhead_end:
        return "directed", target, source
    if arrowhead_end and not arrowhead_start:
        return "directed", source, target
    first, second = sorted((source, target))
    return ("bidirectional" if arrowhead_start and arrowhead_end else "undirected"), first, second


def evaluate(pptx_path: Path, annotation_path: Path | None = None) -> dict[str, object]:
    presentation = Presentation(str(pptx_path))
    if len(presentation.slides) != 1:
        raise ValueError("editability evaluation currently expects exactly one slide")
    slide = presentation.slides[0]
    shapes = list(slide.shapes)
    shape_id_to_shape = {str(shape.shape_id): shape for shape in shapes}

    node_shapes = [
        shape
        for shape in shapes
        if shape.shape_type != MSO_SHAPE_TYPE.LINE
        and (shape.name.startswith("node:") or ":node:" in shape.name)
    ]
    container_shapes = [
        shape
        for shape in shapes
        if shape.shape_type != MSO_SHAPE_TYPE.LINE
        and (shape.name.startswith("container:") or ":container:" in shape.name)
    ]
    line_shapes = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    standalone_text_shapes = [
        shape
        for shape in shapes
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        and not (shape.name.startswith("node:") or ":node:" in shape.name)
    ]
    redundant_standalone_textboxes = 0
    for textbox in standalone_text_shapes:
        if not _normalize_text(textbox.text):
            continue
        text_bbox = _bbox(textbox)
        center = ((text_bbox[0] + text_bbox[2]) / 2, (text_bbox[1] + text_bbox[3]) / 2)
        if any(
            _normalize_text(node.text) == _normalize_text(textbox.text)
            and node_bbox[0] <= center[0] <= node_bbox[2]
            and node_bbox[1] <= center[1] <= node_bbox[3]
            for node in node_shapes
            for node_bbox in (_bbox(node),)
        ):
            redundant_standalone_textboxes += 1
    connector_groups: dict[str, list] = {}
    for shape in line_shapes:
        connector_groups.setdefault(shape.name, []).append(shape)

    attached_starts = sum(
        any(_connected_shape_id(segment, "stCxn") is not None for segment in segments)
        for segments in connector_groups.values()
    )
    attached_ends = sum(
        any(_connected_shape_id(segment, "endCxn") is not None for segment in segments)
        for segments in connector_groups.values()
    )
    report: dict[str, object] = {
        "pptx": str(pptx_path),
        "slide_pixels": [round(presentation.slide_width / EMU_PER_PIXEL), round(presentation.slide_height / EMU_PER_PIXEL)],
        "native_objects": {
            "total": len(shapes),
            "nodes": len(node_shapes),
            "containers": len(container_shapes),
            "connector_segments": len(line_shapes),
            "semantic_connectors": len(connector_groups),
            "standalone_textboxes": len(standalone_text_shapes),
            "redundant_standalone_textboxes": redundant_standalone_textboxes,
            "objects_with_text": sum(bool(getattr(shape, "has_text_frame", False) and shape.text.strip()) for shape in shapes),
        },
        "endpoint_bindings": {
            "attached_starts": attached_starts,
            "attached_ends": attached_ends,
            "total_connectors": len(connector_groups),
            "fully_attached": sum(
                any(_connected_shape_id(segment, "stCxn") is not None for segment in segments)
                and any(_connected_shape_id(segment, "endCxn") is not None for segment in segments)
                for segments in connector_groups.values()
            ),
        },
    }
    if annotation_path is None:
        return report

    annotation = json.loads(annotation_path.read_text(encoding="utf-8"))
    primitive_scene = annotation["primitive_scene"]
    annotation_size = annotation["image_size"]
    scale_x = (presentation.slide_width / EMU_PER_PIXEL) / annotation_size["width"]
    scale_y = (presentation.slide_height / EMU_PER_PIXEL) / annotation_size["height"]

    def reference_items(key: str):
        items = []
        for item in primitive_scene[key]:
            bbox = item["bbox"]
            items.append(
                {
                    "id": item["id"],
                    "bbox": (
                        bbox["x0"] * scale_x,
                        bbox["y0"] * scale_y,
                        bbox["x1"] * scale_x,
                        bbox["y1"] * scale_y,
                    ),
                    "label": item.get("label"),
                }
            )
        return items

    predicted_nodes = [{"shape": shape, "bbox": _bbox(shape)} for shape in node_shapes]
    predicted_containers = [{"shape": shape, "bbox": _bbox(shape)} for shape in container_shapes]
    reference_nodes = reference_items("nodes")
    reference_containers = reference_items("containers")
    node_matches = _greedy_matches(predicted_nodes, reference_nodes)
    container_matches = _greedy_matches(predicted_containers, reference_containers)

    predicted_shape_to_reference: dict[str, str] = {}
    for predicted, reference in node_matches.items():
        predicted_shape_to_reference[str(predicted_nodes[predicted]["shape"].shape_id)] = reference_nodes[reference]["id"]
    for predicted, reference in container_matches.items():
        predicted_shape_to_reference[str(predicted_containers[predicted]["shape"].shape_id)] = reference_containers[reference]["id"]

    predicted_edges: Counter[tuple[str, str]] = Counter()
    predicted_semantics: Counter[tuple[str, str, str]] = Counter()
    unattached_endpoint_count = 0
    for segments in connector_groups.values():
        source_id = next((_connected_shape_id(segment, "stCxn") for segment in segments if _connected_shape_id(segment, "stCxn")), None)
        target_id = next((_connected_shape_id(segment, "endCxn") for segment in segments if _connected_shape_id(segment, "endCxn")), None)
        if source_id is None:
            unattached_endpoint_count += 1
        if target_id is None:
            unattached_endpoint_count += 1
        source_ref = predicted_shape_to_reference.get(source_id or "")
        target_ref = predicted_shape_to_reference.get(target_id or "")
        if source_ref and target_ref:
            predicted_edges[tuple(sorted((source_ref, target_ref)))] += 1
            predicted_semantics[
                _edge_semantics(
                    source_ref,
                    target_ref,
                    arrowhead_start=any(_has_arrowhead(segment, "headEnd") for segment in segments),
                    arrowhead_end=any(_has_arrowhead(segment, "tailEnd") for segment in segments),
                )
            ] += 1

    reference_edges: Counter[tuple[str, str]] = Counter()
    reference_semantics: Counter[tuple[str, str, str]] = Counter()
    for connector in primitive_scene["connector_candidates"]:
        source = connector["start_endpoint"]["owner_id"]
        target = connector["end_endpoint"]["owner_id"]
        reference_edges[tuple(sorted((source, target)))] += 1
        reference_semantics[
            _edge_semantics(
                source,
                target,
                arrowhead_start=connector.get("arrowhead_start", False),
                arrowhead_end=connector.get("arrowhead_end", False),
            )
        ] += 1
    edge_true_positive = sum((predicted_edges & reference_edges).values())
    semantic_true_positive = sum((predicted_semantics & reference_semantics).values())

    correct_labels = 0
    for predicted, reference in node_matches.items():
        predicted_text = predicted_nodes[predicted]["shape"].text
        reference_text = reference_nodes[reference]["label"]
        if reference_text and _normalize_text(predicted_text) == _normalize_text(reference_text):
            correct_labels += 1

    node_metrics = _f1_counts(len(predicted_nodes), len(reference_nodes), len(node_matches))
    container_metrics = _f1_counts(len(predicted_containers), len(reference_containers), len(container_matches))
    connector_metrics = _f1_counts(len(connector_groups), sum(reference_edges.values()), edge_true_positive)
    connector_semantic_metrics = _f1_counts(
        len(connector_groups), sum(reference_semantics.values()), semantic_true_positive
    )
    direction_corrections = edge_true_positive - semantic_true_positive
    missing_labels = sum(bool(item["label"]) for item in reference_nodes) - correct_labels
    edit_operations_lower_bound = (
        node_metrics["fp"]
        + node_metrics["fn"]
        + container_metrics["fp"]
        + container_metrics["fn"]
        + connector_metrics["fp"]
        + connector_metrics["fn"]
        + direction_corrections
        + missing_labels
        + unattached_endpoint_count
        + redundant_standalone_textboxes
    )
    report["ground_truth"] = {
        "nodes": node_metrics,
        "containers": container_metrics,
        "connector_topology": connector_metrics,
        "connector_semantics": connector_semantic_metrics,
        "connector_direction_corrections": direction_corrections,
        "labels": {
            "reference": sum(bool(item["label"]) for item in reference_nodes),
            "exact_on_matched_node": correct_labels,
            "missing_or_incorrect": missing_labels,
        },
        "unattached_connector_endpoints": unattached_endpoint_count,
        "edit_operations_lower_bound": edit_operations_lower_bound,
    }
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--annotation", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = evaluate(args.pptx, args.annotation)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

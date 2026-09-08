"""Bind frozen numeric results to actual COM and posthoc development-only diagnosis."""
import argparse
import json
from pathlib import Path
from PIL import Image,ImageDraw
from pptx import Presentation
from r17_object_primitive import ROOT,digest,save,verify,load_model,candidates,object_counts


def run(output):
    output.mkdir(parents=True,exist_ok=False);base=ROOT/'artifacts/region-image-only-r17'
    result_path=base/'eval-v1/results.json';results=json.loads(result_path.read_text());verify(json.loads((base/'eval-v1/lock.json').read_text()))
    native=json.loads((base/'com-v1/audit.json').read_text())
    assert digest(base/'com-v1/renders/powerpoint_render_manifest.json')==native['receipt_sha256']
    assert native['checked_pptx']==48
    receipts=json.loads((base/'com-v1/renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))['records']
    for row in receipts:
        assert digest(row['pptx_path'])==row['pptx_sha256'] and digest(row['render_path'])==row['render_sha256']
    for row in native['boards']:assert digest(row['path'])==row['sha256']
    single=json.loads((base/'rgb-only-v1/result.json').read_text());verify(single)
    expected=next(r for r in results['arms']['scene300_scene300']['rows'] if r['id']=='r17-scene-eval-000')
    assert single['objects']==expected['objects']
    original=base/'original-diagnostic-v1';diagnostic=json.loads((original/'result.json').read_text());verify(diagnostic)
    receipt=json.loads((original/'renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))['records'][0]
    assert digest(original/'result.pptx')==diagnostic['pptx_sha256']==receipt['pptx_sha256']
    assert digest(receipt['render_path'])==receipt['render_sha256']
    deck=Presentation(original/'result.pptx');assert len(deck.slides[0].shapes)==len(diagnostic['objects'])==24
    source=ROOT/'artifacts/region-image-only-r10/complex-inputs/original-rgb.png'
    proposals,crops,size=candidates(load_model(ROOT/'artifacts/region-image-only-r16/primitive-training-v1/scene80_scene80.pt'),source)
    # Manual ledger is consumed after RGB-only inference and never supplies candidate geometry.
    ledger_path=ROOT/'docs/r15_visual_loss_ledger.json';ledger=json.loads(ledger_path.read_text())
    expected=[dict(kind=1,bbox=[v*size[i%2]/ledger['source_display_size'][i%2] for i,v in enumerate(n['box'])]) for n in ledger['nodes']]
    coverage=object_counts([dict(o,kind=1) for o in proposals],expected)
    board=Image.new('RGB',(2400,694),'white');draw=ImageDraw.Draw(board)
    with Image.open(source) as im:board.paste(im.convert('RGB').resize((1200,654)),(0,40))
    with Image.open(receipt['render_path']) as im:board.paste(im.convert('RGB'),(1200,40))
    draw.text((4,4),'Development original: source RGB',fill='black');draw.text((1204,4),'Actual native PowerPoint: 24 incorrect diamonds',fill='black')
    board.save(output/'original-comparison.png')
    board=Image.new('RGB',(6*128,6*88),'white');draw=ImageDraw.Draw(board)
    emitted={o['id']:o for o in diagnostic['objects']}
    for i,(obj,crop) in enumerate(zip(proposals,crops,strict=True)):
        x=(i%6)*128;y=(i//6)*88
        draw.text((x+2,y),obj['id']+(' diamond' if obj['id'] in emitted else ' rejected'),fill='black')
        for j in range(2):board.paste(Image.fromarray((crop[j].numpy()*255).astype('uint8')).convert('RGB'),(x+j*64,y+22))
    board.save(output/'original-candidate-crops.png')
    gates={a:bool(r['gate_pass'] and r['precision']==1 and r['recall']>=.8 and r['negative_fp']==0 and all(v['recall']>=.5 for v in r['per_kind'].values())) for a,r in results['arms'].items()}
    save(output/'audit.json',dict(research_gate=gates,pretraining_advantage=False,production_approved=False,
        source_results_sha256=digest(result_path),com_audit_sha256=digest(base/'com-v1/audit.json'),rgb_only_equal=True,
        original=dict(proposals=len(proposals),emitted=24,manual_17_class_agnostic=coverage,ledger_sha256=digest(ledger_path),
            source_sha256=digest(source),prediction_sha256=digest(original/'result.json'),render_sha256=receipt['render_sha256'],
            scope='Posthoc exposed development manual ledger; not independent GT. Visual verdict: all emitted diamonds incorrect.'),
        code_sha256=digest(__file__),scope='Synthetic gate and real-diagram failure are separate verdicts.'))
    print(json.dumps(dict(gates=gates,original_candidates=len(proposals),manual_coverage=coverage)))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);run(p.parse_args().output)

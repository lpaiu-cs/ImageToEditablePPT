"""Small local COM-paired classifier development diagnostic, no held-out datasets."""
import argparse
import json
from pathlib import Path
import shutil
import numpy as np
from PIL import Image,ImageOps
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.dml.color import RGBColor
from r10_owner_region_p2_audit import sha,snapshot

BOX=[180,100,560,320]
EMU=9525


def deck():
    p=Presentation();p.slide_width=Emu(800*EMU);p.slide_height=Emu(450*EMU)
    return p,p.slides.add_slide(p.slide_layouts[6])


def rect(s,b,color=(255,255,255),name='structure'):
    x,y,c,d=b;r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Emu(x*EMU),Emu(y*EMU),Emu((c-x)*EMU),Emu((d-y)*EMU))
    r.name=name;r.fill.solid();r.fill.fore_color.rgb=RGBColor(*color);r.line.color.rgb=RGBColor(0,0,0);r.line.width=Pt(1)
    return r


def label(s,text,x,y,w):
    r=s.shapes.add_textbox(Emu(x*EMU),Emu(y*EMU),Emu(w*EMU),Emu(25*EMU));r.text=text
    for run in r.text_frame.paragraphs[0].runs:run.font.size=Pt(12)


def prepare(out,source):
    out.mkdir(parents=True,exist_ok=False);(out/'native').mkdir();(out/'derived').mkdir();(out/'assets').mkdir()
    _,ps=snapshot(source/'closed_rect_proposals.py',out,'proposal_snapshot')
    photo=Path('C:/vs/ImageToEditablePPT/.venv/Lib/site-packages/sklearn/datasets/images')
    shutil.copyfile(photo/'README.txt',out/'assets/README-source.txt')
    for name in ['china.jpg','flower.jpg']:shutil.copyfile(photo/name,out/'assets'/name)
    records=[]
    for split_index,split in enumerate(['train','calibration','eval']):
        for i in range(2):
            base=f'{split}-base{i}';p,s=deck();rect(s,BOX,name='parent');label(s,'Process overview',200,110,300)
            for j in range(3):
                x=200+j*112;y=160+split_index*24+i*8;rect(s,[x,y,x+90,y+60],(225,235,245),f'node{j}');label(s,['Read','Build','Check'][j],x+6,y+15,75)
                if j<2:s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Emu((x+90)*EMU),Emu((y+30)*EMU),Emu((x+112)*EMU),Emu((y+30)*EMU))
            p.save(out/'native'/f'{base}-native.pptx')
            for variant in ['native','diagram_picture']:
                records.append(dict(id=f'{base}-{variant}',base=base,split=split,variant=variant,label=1,asset_group=base,target_bbox=BOX))
            if split!='calibration':
                asset='china.jpg' if split=='train' else 'flower.jpg'
                for variant in ['photo','caption','gray_photo','gray_caption']:
                    records.append(dict(id=f'{base}-{variant}',base=base,split=split,variant=variant,label=0,asset_group=asset,target_bbox=BOX))
    lock=dict(seed=1909,epochs=120,learning_rate=.003,threshold=.5,input_size=[96,64],
              policy='RGB crop only; classifier-only fixed proposal; calibration positive-only, no threshold tuning; 2 photos cannot establish generalization.',
              source_sha256=ps,photo_sha256={n:sha(out/'assets'/n) for n in ['china.jpg','flower.jpg']},
              builder_sha256=sha(__file__),records=records)
    (out/'preregister.json').write_text(json.dumps(lock,indent=2));print(sha(out/'preregister.json'))


def derive(out):
    lock=json.loads((out/'preregister.json').read_text())
    for r in lock['records']:
        if r['variant']=='native':continue
        p,s=deck();variant=r['variant']
        if variant=='diagram_picture':
            im=Image.open(out/'native-renders'/f"{r['base']}-native.png").crop(BOX)
        else:
            im=Image.open(out/'assets'/r['asset_group']).convert('RGB')
            if variant.startswith('gray'):im=ImageOps.grayscale(im).convert('RGB')
        asset=out/'assets'/f"{r['id']}.png";im.save(asset)
        rect(s,BOX,name='parent')
        caption='caption' in variant
        if caption:label(s,'Photo reference',200,108,300)
        s.shapes.add_picture(str(asset),Emu(182*EMU),Emu((140 if caption else 102)*EMU),Emu(376*EMU),Emu((178 if caption else 216)*EMU))
        p.save(out/'derived'/f"{r['id']}.pptx")


def evaluate(out):
    import torch
    from torch import nn
    lock=json.loads((out/'preregister.json').read_text());torch.manual_seed(lock['seed']);torch.set_num_threads(4)
    spec=__import__('importlib.util',fromlist=['x']);ss=spec.spec_from_file_location('frozen_proposal',out/'proposal_snapshot.py');m=spec.module_from_spec(ss);ss.loader.exec_module(m)
    receipts={}
    for folder in ['native-renders','derived-renders']:
        manifest=json.loads((out/folder/'powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
        for r in manifest['records']:
            assert sha(Path(r['pptx_path']))==r['pptx_sha256'] and sha(Path(r['render_path']))==r['render_sha256'];receipts[r['id']]=r
    data=[]
    for r in lock['records']:
        receipt=receipts[r['id']];rgb=np.array(Image.open(receipt['render_path']).convert('RGB'));selected,audit=m.propose_union(rgb)
        c=max(selected,key=lambda v:m._iou(v['bbox'],r['target_bbox']),default=None)
        if c is None or m._iou(c['bbox'],r['target_bbox'])<.8:
            data.append(dict(**r,missing=True));continue
        x,y,a,b=c['bbox'];crop=Image.fromarray(rgb[y:b,x:a]);path=out/'assets'/f"{r['id']}-crop.png";crop.save(path)
        data.append(dict(**r,missing=False,crop=str(path.resolve()),crop_sha256=sha(path),candidate_bbox=c['bbox'],receipt=receipt))
    # This persisted dataset owns paths/targets; forward receives only numeric RGB tensors.
    (out/'candidate_manifest.json').write_text(json.dumps(data,indent=2))
    available=[r for r in data if not r['missing']]
    X=torch.stack([torch.tensor(np.array(Image.open(r['crop']).resize(tuple(lock['input_size']),Image.Resampling.BILINEAR)),dtype=torch.float32).permute(2,0,1)/255 for r in available])
    Y=torch.tensor([r['label'] for r in available],dtype=torch.float32)
    train=torch.tensor([r['split']=='train' for r in available])
    model=nn.Sequential(nn.Conv2d(3,8,3,padding=1),nn.ReLU(),nn.AvgPool2d(2),nn.Conv2d(8,16,3,padding=1),nn.ReLU(),nn.AdaptiveAvgPool2d((4,4)),nn.Flatten(),nn.Linear(256,1))
    opt=torch.optim.Adam(model.parameters(),lr=lock['learning_rate']);losses=[]
    for epoch in range(lock['epochs']):
        opt.zero_grad();loss=nn.functional.binary_cross_entropy_with_logits(model(X[train]).flatten(),Y[train]);loss.backward();opt.step();losses.append(float(loss.detach()))
    torch.save(model.state_dict(),out/'model.pt');model.eval()
    with torch.no_grad():scores=torch.sigmoid(model(X).flatten()).tolist()
    for r,score in zip(available,scores):r['score']=score;r['predicted']=int(score>=lock['threshold'])
    summary={}
    for split in ['train','calibration','eval']:
        rs=[r for r in available if r['split']==split];summary[split]=dict(n=len(rs),positive=sum(r['label'] for r in rs),
            false_accept=sum(r['label']==0 and r['predicted']==1 for r in rs),false_reject=sum(r['label']==1 and r['predicted']==0 for r in rs))
    result=dict(preregister_sha256=sha(out/'preregister.json'),model_sha256=sha(out/'model.pt'),candidate_manifest_sha256=sha(out/'candidate_manifest.json'),
                missing=sum(r['missing'] for r in data),summary=summary,loss_start=losses[0],loss_end=losses[-1],records=available)
    (out/'results.json').write_text(json.dumps(result,indent=2));print(json.dumps({k:v for k,v in result.items() if k!='records'}))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['prepare','derive','evaluate']);p.add_argument('--out',type=Path,required=True);p.add_argument('--source',type=Path);a=p.parse_args()
    if a.mode=='prepare':prepare(a.out,a.source)
    elif a.mode=='derive':derive(a.out)
    else:evaluate(a.out)

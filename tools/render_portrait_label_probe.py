"""Render-only adapter for frozen RGB observations; not a fresh inference claim."""
import argparse
import json
import shutil
from pathlib import Path
from PIL import ImageFont
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR,MSO_AUTO_SIZE,PP_ALIGN
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
from run_region_rgb_baseline import digest,save


def physical_bbox(box,width,height,orientation):
    a,b,c,d=box
    if orientation=='clockwise':return [width-d,a,width-b,c]
    if orientation=='counterclockwise':return [b,height-c,d,height-a]
    raise ValueError('unknown text orientation')


def apply_label(shape,row,source_width,source_height,scale):
    observed=row['selected'];line=observed['spatial']['accepted'][0]
    a,b,c,d=physical_bbox(line['bbox'],source_width,source_height,observed['orientation'])
    if not (0<=a<c<=source_width and 0<=b<d<=source_height):raise ValueError('text outside owner')
    frame=shape.text_frame;frame.text=observed['label']
    frame._txBody.bodyPr.set('vert','vert' if observed['orientation']=='clockwise' else 'vert270')
    emu=lambda v:Emu(round(v*scale*9525))
    frame.margin_left=emu(a);frame.margin_right=emu(source_width-c)
    frame.margin_top=emu(b);frame.margin_bottom=emu(source_height-d)
    frame.vertical_anchor=MSO_ANCHOR.MIDDLE;frame.auto_size=MSO_AUTO_SIZE.NONE;frame.word_wrap=False
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',100)
    available_width=d-b;available_height=c-a
    size=min(available_height*.85,available_width*95/max(1,font.getlength(observed['label'])))*scale*.75
    paragraph=frame.paragraphs[0];paragraph.alignment=PP_ALIGN.CENTER
    paragraph.space_before=Pt(0);paragraph.space_after=Pt(0)
    # Assigning frame.text discards the original runs and their explicit color.
    # Theme defaults can otherwise make recovered labels white on white.
    fill=shape.fill.fore_color.rgb
    dark_background=sum(float(v)*w for v,w in zip(fill,(.2126,.7152,.0722)))<128
    for run in paragraph.runs:
        run.font.name='Arial';run.font.size=Pt(max(1,size))
        run.font.color.rgb=RGBColor(255,255,255) if dark_background else RGBColor(31,41,55)
    return dict(id=shape.name,label=frame.text,orientation=observed['orientation'],source_bbox=[a,b,c,d],font_size_pt=size)


def run(native,observations,output):
    output.mkdir(parents=True,exist_ok=False)
    data=json.loads(observations.read_text(encoding='utf-8'));result=json.loads((native/'result.json').read_text(encoding='utf-8'))
    assert data['source_sha256']==result['source_sha256']
    pptx=native/'result.pptx';original=digest(pptx);target=output/'result.pptx'
    shutil.copy2(pptx,target);deck=Presentation(target);shapes={s.name:s for s in deck.slides[0].shapes}
    owners={s['id']:s for s in result['source_scene']['shapes']};changes=[]
    for row in data['rows']:
        if row['selected'] is None or row['id'] not in owners:continue
        owner=owners[row['id']];box=owner['bbox'];width=box['x1']-box['x0'];height=box['y1']-box['y0']
        shape=shapes[row['id']];scale=shape.width/9525/width
        changes.append(apply_label(shape,row,width,height,scale))
    deck.save(target);assert digest(pptx)==original
    save(output/'render-adapter.json',dict(source_sha256=data['source_sha256'],input_pptx_sha256=original,
        observations_sha256=digest(observations),runner_sha256=digest(Path(__file__)),pptx_sha256=digest(target),changes=changes,
        scope='Frozen-observation presentation probe only; not fresh integrated inference.'))
    print(changes)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--native',type=Path,required=True);p.add_argument('--observations',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    a=p.parse_args();run(a.native,a.observations,a.output)

"""Audit already-completed fixed-ROI intervention; manual GT is posthoc only."""
import argparse
import json
from pathlib import Path
from PIL import Image,ImageDraw
from pptx import Presentation
from r19_outcome_audit import ROOT,digest,save,verify,geometry,object_counts


def run(output):
    output.mkdir(parents=True,exist_ok=False);base=ROOT/'artifacts/region-image-only-r20'
    result=json.loads((base/'probe-v1/results.json').read_text());verify(result)
    original=json.loads((ROOT/'artifacts/region-image-only-r19/original-v1/result.json').read_text())
    assert result['candidates']==original['all_candidates'];candidates={o['id']:o for o in result['candidates']}
    receipts={r['id']:r for r in json.loads((base/'com-v1/renders/powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))['records']};assert len(receipts)==4
    ledger_path=ROOT/'docs/r15_visual_loss_ledger.json';ledger=json.loads(ledger_path.read_text())
    source=ROOT/'artifacts/region-image-only-r10/complex-inputs/original-rgb.png'
    with Image.open(source) as im:image=im.convert('RGB')
    expected=[dict(kind=6 if n.get('kind')=='cylinder' else 2,bbox=[v*image.size[i%2]/ledger['source_display_size'][i%2] for i,v in enumerate(n['box'])]) for n in ledger['nodes']]
    board=Image.new('RGB',(2070,812),'white');draw=ImageDraw.Draw(board);board.paste(image.resize((690,376)),(0,30));draw.text((4,4),'Exposed RGB source',fill='black');counts={}
    for i,(arm,row) in enumerate(result['arms'].items(),1):
        receipt=receipts[arm];assert receipt['pptx_sha256']==row['pptx_sha256']==digest(receipt['pptx_path'])
        assert receipt['render_sha256']==digest(receipt['render_path'])
        assert len(Presentation(receipt['pptx_path']).slides[0].shapes)==len(row['objects'])
        assert all(candidates[o['id']]['bbox']==o['bbox'] for o in row['objects'])
        counts[arm]=dict(kind=object_counts(row['objects'],expected),geometry=geometry(row['objects'],expected))
        x=i%3*690;y=i//3*406
        with Image.open(receipt['render_path']) as im:board.paste(im.convert('RGB').resize((690,376)),(x,y+30))
        c=counts[arm]['kind'];draw.text((x+4,y+4),f'{arm} manual TP{c["tp"]} FP{c["fp"]} FN{c["fn"]}',fill='black')
    board.save(output/'comparison.png');save(output/'audit.json',dict(counts=counts,fixed_boxes_verified=True,
        prediction_sha256=digest(base/'probe-v1/results.json'),receipts_sha256=digest(base/'com-v1/renders/powerpoint_render_manifest.json'),
        ledger_sha256=digest(ledger_path),code_sha256=digest(__file__),scope='Exposed development manual matching; no independent performance claim.'))
    print(json.dumps(counts,indent=2))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);run(p.parse_args().output)

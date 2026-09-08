"""Variable-layout role-labelled source PPTs; no predictions, GPU or union masks."""
import argparse
from collections import Counter
import json
import math
from pathlib import Path
import random
import zipfile
from xml.etree import ElementTree as ET
import numpy as np
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.dml.color import RGBColor
from r16_ppt_primitive_data import KINDS,PRESETS,NS,no_effect
from r10_classifier_pilot import EMU
from r10_owner_region_p2_audit import sha

SCALE=.8
FAMILIES=['colored_fill_white_bg','white_card_color_bg','outline_white_bg','nested_container','mixed_dense','hard_negative']
WORDS={'development':['Read','Build','Check','Input','Process','Store','Task','Group','Plan','State','Memory','Output'],
       'eval':['Observe','Route','Assess','Signal','Archive','Result','Stage','Cluster','Decide','Cache','Merge','Review']}


def plan():
    records=[]
    for split,n,base in [('train',48,190100),('calibration',8,190200),('eval',8,190300)]:
        for fi,family in enumerate(FAMILIES):
            seed=base+fi;rng=random.Random(seed);words=WORDS['eval' if split=='eval' else 'development']
            lo,hi=(9,15) if family=='mixed_dense' else (3,9)
            counts=[rng.randint(lo,hi) for _ in range(n)]
            counts[-1]=next(k for k in range(lo,hi+1) if (sum(counts[:-1])+k)%6==0)
            kind_pool=list(range(1,7))*(sum(counts)//6);rng.shuffle(kind_pool);cursor=0
            parents=[1,2]*(n//2);rng.shuffle(parents)
            for i,count in enumerate(counts):
                identity=f'r19-{family}-{split}-{i:03d}';negative=family=='hard_negative'
                background=[rng.randint(140,225) for _ in range(3)] if family=='white_card_color_bg' or (family=='mixed_dense' and i%2) else [255]*3
                elements=[dict(id='background',role='background',type='background',color=background)];objects=[]
                def shape(name,kind,bbox,role,fill,parent=None):
                    e=dict(id=name,type='shape',kind=kind,role=role,bbox_source=bbox,bbox=[v*SCALE for v in bbox],fill=fill,
                        stroke=[rng.randint(0,65) for _ in range(3)],stroke_pt=rng.choice([1,1.5,2,2.5,3]),parent_id=parent)
                    elements.append(e)
                    if role=='closed_structure':objects.append(e)
                def text(name,content,bbox,role='decoration',owner=None,font=14):
                    elements.append(dict(id=name,type='text',role=role,text=content,bbox_source=bbox,font_pt=font,owner_id=owner))
                def line(name,a,b,role='open_structure'):
                    elements.append(dict(id=name,type='line',role=role,start=a,end=b,stroke_pt=rng.choice([1,1.5,2])))
                shape('banner',1,[20,12,1180,62],'decoration',[rng.randint(180,235) for _ in range(3)])
                text('title',' '.join(rng.sample(words,3)),[35,16,1160,60],font=24)
                text('paragraph-top',' '.join(rng.choices(words,k=16)),[40,76,1160,110],font=13)
                text('paragraph-bottom',' '.join(rng.choices(words,k=16)),[40,718,1160,748],font=13)
                if not negative:
                    parent_count=parents[i] if family in ['nested_container','mixed_dense'] else 0
                    if parent_count==2:regions=[[45,125,585,705],[615,125,1155,705]]
                    else:regions=[[45,125,1155,705]]
                    for pi,region in enumerate(regions):
                        if parent_count:shape(f'container{pi}',2,region,'closed_structure',None)
                    groups=[count//len(regions)+(1 if p<count%len(regions) else 0) for p in range(len(regions))]
                    node_index=0
                    for pi,(region,num) in enumerate(zip(regions,groups)):
                        a,b,c,d=region
                        cols=(3 if split=='eval' else 2) if parent_count==2 else (4 if split=='eval' else 5 if family=='mixed_dense' else 3)
                        cols=min(cols,num);rows=math.ceil(num/cols);cw=(c-a-30)/cols;ch=(d-b-30)/rows
                        slots=list(range(num));rng.shuffle(slots)
                        for slot in slots:
                            col=slot%cols;row=slot//cols
                            # Separate eval position family: stagger alternate rows as well as column count.
                            shift=(.07*cw if row%2 else 0) if split=='eval' else 0
                            w=round(cw*rng.uniform(.53,.80));h=round(ch*rng.uniform(.44,.76))
                            x=round(a+15+col*cw+cw*rng.uniform(.03,.10)+shift);y=round(b+15+row*ch+ch*rng.uniform(.03,.10))
                            kind=kind_pool[cursor];cursor+=1;name=f'node{node_index}';node_index+=1
                            style=('filled','outline','whitecard')[node_index%3] if family in ['nested_container','mixed_dense'] else 'outline' if family=='outline_white_bg' else 'whitecard' if family=='white_card_color_bg' else 'filled'
                            fill=None if style=='outline' else [255]*3 if style=='whitecard' else [rng.randint(115,235) for _ in range(3)]
                            shape(name,kind,[x,y,x+w,y+h],'closed_structure',fill,f'container{pi}' if parent_count else None);objects[-1]['style_group']=style
                            lines=rng.choice([2,3]);content='\n'.join(' '.join(rng.choices(words,k=rng.choice([1,2]))) for _ in range(lines))
                            text(name+'-label',content,[x+w*.19,y+h*.19,x+w*.82,y+h*.87],'structural_label',name,font=max(7,min(15,int(h/lines/2))))
                            if rng.random()<.65:
                                size=max(9,min(23,round(min(w,h)*.18)));dx=x+w-size-5;dy=y+4
                                shape(name+'-glyph',rng.choice([3,7]),[dx,dy,dx+size,dy+size],'decoration',[70,90,125])
                    # These may cross structure/text; they remain open_structure, never closed GT.
                    for k in range(rng.randint(2,5)):
                        line(f'crossing{k}',[rng.randint(70,400),rng.randint(160,650)],[rng.randint(760,1120),rng.randint(160,650)])
                else:
                    for j in range(rng.randint(5,9)):
                        y=130+j*58;x=rng.randint(40,160)
                        text(f'body{j}',' '.join(rng.choices(words,k=14)),[x,y,1020,y+28],font=12)
                        line(f'open{j}',[x,y+32],[rng.randint(720,1050),y+32+rng.randint(-18,18)])
                        shape(f'glyph{j}',rng.choice([3,7]),[1090,y,1110,y+20],'decoration',[90,110,150])
                records.append(dict(id=identity,split=split,stage='negative' if negative else 'scene',family=family,seed=seed,
                    layout_family='staggered_eval' if split=='eval' else 'grid_development',vocabulary_group='eval' if split=='eval' else 'development',
                    objects=objects,elements=elements,node_count=0 if negative else count,container_count=0 if negative else parent_count))
            if not family=='hard_negative':assert cursor==len(kind_pool)
    return records


def build(out):
    out.mkdir(parents=True,exist_ok=False);(out/'pptx').mkdir();records=plan()
    (out/'preregister.json').write_text(json.dumps(dict(records=records,classmapping=KINDS,generator_sha256=sha(__file__),
        contract='Variable nodes and separate containers; roles distinguish native decoration from closed structure. Eval layout/vocabulary separated. Balanced node kinds per family/split; additional rounded containers reported separately. No masks/models/manual CogBot ledger.'),indent=2))
    for r in records:
        p=Presentation();p.slide_width=Emu(1200*EMU);p.slide_height=Emu(750*EMU);s=p.slides.add_slide(p.slide_layouts[6])
        for e in r['elements']:
            if e['type']=='background':s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(*e['color']);continue
            if e['type'] in ['shape','text']:
                x,y,a,b=e['bbox_source'];position=[Emu(round(v*EMU)) for v in [x,y,a-x,b-y]]
                if e['type']=='text':
                    v=s.shapes.add_textbox(*position);v.text=e['text']
                    for para in v.text_frame.paragraphs:
                        for run in para.runs:run.font.size=Pt(e['font_pt']);run.font.color.rgb=RGBColor(25,30,40)
                else:
                    v=s.shapes.add_shape(PRESETS[e['kind']-1] if e['kind']<=6 else MSO_SHAPE.STAR_5_POINT,*position)
                    if e['fill'] is None:v.fill.background()
                    else:v.fill.solid();v.fill.fore_color.rgb=RGBColor(*e['fill'])
                    v.line.color.rgb=RGBColor(*e['stroke']);v.line.width=Pt(e['stroke_pt']);e['adjustments']=list(v.adjustments)
            else:
                v=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Emu(round(q*EMU)) for q in [*e['start'],*e['end']]])
                v.line.color.rgb=RGBColor(45,50,60);v.line.width=Pt(e['stroke_pt'])
            v.name=e['id'];no_effect(v);e['source_shape_id']=v.shape_id
        path=out/'pptx'/f'{r["id"]}.pptx';p.save(path);r['pptx']=str(path.resolve());r['pptx_sha256']=sha(path)
    assert len(records)==384
    (out/'source-records.json').write_text(json.dumps(records,indent=2));print('384 source PPTs; prereg',sha(out/'preregister.json'))


def audit_sources(records):
    presets={1:'rect',2:'roundRect',3:'ellipse',4:'diamond',5:'parallelogram',6:'can',7:'star5'}
    for r in records:
        assert sha(r['pptx'])==r['pptx_sha256']
        with zipfile.ZipFile(r['pptx']) as z:
            assert len([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')])==1
            xml=ET.fromstring(z.read('ppt/slides/slide1.xml'))
        assert not xml.findall('.//p:pic',NS)
        shapes={v.find('.//p:cNvPr',NS).get('name'):v for v in xml.findall('.//p:sp',NS)+xml.findall('.//p:cxnSp',NS)}
        assert len(shapes)==len(r['elements'])-1
        for e in r['elements']:
            if e['type']=='background':continue
            v=shapes[e['id']];assert int(v.find('.//p:cNvPr',NS).get('id'))==e['source_shape_id']
            if e['type']=='shape':
                assert v.find('p:spPr/a:prstGeom',NS).get('prst')==presets[e['kind']]
                if e['fill'] is None:assert v.find('p:spPr/a:noFill',NS)is not None
                xf=v.find('p:spPr/a:xfrm',NS);a=xf.find('a:off',NS);b=xf.find('a:ext',NS);x,y,c,d=e['bbox_source']
                assert [int(a.get('x')),int(a.get('y')),int(b.get('cx')),int(b.get('cy'))]==[round(q*EMU) for q in [x,y,c-x,d-y]]
        assert {v['id'] for v in r['objects']}=={v['id'] for v in r['elements'] if v['role']=='closed_structure'}
        nodes=[v for v in r['objects'] if v['id'].startswith('node')]
        assert len(nodes)==r['node_count'] and len(r['objects'])==r['node_count']+r['container_count']
        for i,a in enumerate(nodes):
            x,y,c,d=a['bbox'];assert 0<x<c<960 and 0<y<d<600
            for b in nodes[i+1:]:
                u,v,w,h=b['bbox'];assert c<=u or w<=x or d<=v or h<=y
            if a['parent_id']:
                parent=next(v for v in r['objects'] if v['id']==a['parent_id']);u,v,w,h=parent['bbox'];assert u<x<c<w and v<y<d<h
    for split in ['train','calibration','eval']:
        for family in FAMILIES[:-1]:
            counts=Counter(v['kind'] for r in records if r['split']==split and r['family']==family for v in r['objects'] if v['id'].startswith('node'))
            assert len(counts)==6 and len(set(counts.values()))==1,(split,family,counts)


def finalize(out):
    records=json.loads((out/'source-records.json').read_text());audit_sources(records)
    m=out/'renders/powerpoint_render_manifest.json';receipts={r['id']:r for r in json.loads(m.read_text(encoding='utf-8-sig'))['records']};assert len(receipts)==384
    (out/'rgb').mkdir();seen=set()
    for r in records:
        q=receipts[r['id']];assert sha(q['pptx_path'])==q['pptx_sha256']==r['pptx_sha256'];assert sha(q['render_path'])==q['render_sha256']
        with Image.open(q['render_path']) as im:
            assert 'transparency' not in im.info;a=np.array(im.convert('RGB'));assert a.shape==(600,960,3)
        path=out/'rgb'/f'{r["id"]}.png';Image.fromarray(a).save(path);assert np.array_equal(a,np.array(Image.open(path)))
        digest=sha(path);assert digest not in seen;seen.add(digest);r.update(image=str(path.resolve()),image_sha256=digest,render_sha256=q['render_sha256'])
    counts=Counter((r['split'],r['family']) for r in records)
    manifest=dict(records=records,classmapping=KINDS,source_size=[1200,750],render_size=[960,600],generator_sha256=sha(__file__),
        preregister_sha256=sha(out/'preregister.json'),com_manifest_sha256=sha(m),counts={f'{a}/{b}':n for (a,b),n in counts.items()},
        scope='Native synthetic structures and hard negatives; node kinds balanced, rounded containers additional. No union masks/photos/models. Eval vocabulary/layout separated.')
    (out/'dataset_manifest.json').write_text(json.dumps(manifest,indent=2));print('384 RGB/source audited; manifest',sha(out/'dataset_manifest.json'))


def qa(out):
    m=json.loads((out/'dataset_manifest.json').read_text());records=m['records'];audit_sources(records)
    assert not set(WORDS['development']) & set(WORDS['eval'])
    for family in FAMILIES:
        r=next(r for r in records if r['split']=='train' and r['family']==family)
        im=Image.open(r['image']);board=Image.new('RGB',(1920,630),'white');board.paste(im,(0,30));board.paste(im,(960,30));draw=ImageDraw.Draw(board)
        draw.text((5,5),r['id']+' | source / closed_structure GT only',fill='black')
        for obj in r['objects']:
            x,y,a,b=obj['bbox'];draw.rectangle([x+960,y+30,a+960,b+30],outline='red',width=2);draw.text((x+962,y+32),f'{obj["id"]}:k{obj["kind"]}',fill='red')
        board.save(out/f'train-qa-{family}.png')
    (out/'qa.json').write_text(json.dumps(dict(records=384,source_checks_pass=True,node_kind_balance=True,vocabulary_disjoint=True,eval_visual_opened=False,model_run=False),indent=2));print('QA passed; train-only six boards')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('mode',choices=['build','finalize','qa']);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    {'build':build,'finalize':finalize,'qa':qa}[a.mode](a.out)

"""Source/baseline/candidate native comparison with serialization audits."""
import argparse
import json
from pathlib import Path
from PIL import Image,ImageOps,ImageDraw,ImageFont
from pptx import Presentation
from image_to_editable_ppt.research.region_first import digest,pixel_box

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'artifacts/region-generalization-r6'

def main():
    p=argparse.ArgumentParser(); p.add_argument('cohort',choices=['complex','transfer'])
    p.add_argument('--output-name'); args=p.parse_args()
    cohort=args.cohort
    manifest=ROOT/('benchmarks/region_complex_r5_inputs_v2.json' if cohort=='complex' else 'benchmarks/region_r6_transfer_inputs.json')
    samples=json.loads(manifest.read_text())['samples']
    prior=ROOT/'artifacts/region-complex-r5/final-v2' if cohort=='complex' else BASE/'transfer-baseline'
    current=BASE/('complex-candidate2' if cohort=='complex' else 'transfer-candidate')
    prior_run=prior/'run' if cohort=='complex' else prior
    priors={r['id']:r for r in json.loads((prior_run/'results.json').read_text())['records']}
    rows={r['id']:r for r in json.loads((current/'results.json').read_text())['records']}
    for directory in (prior/'renders',current/'renders'):
        native=json.loads((directory/'powerpoint_render_manifest.json').read_text(encoding='utf-8-sig'))
        for r in native['records']:
            assert digest(Path(r['pptx_path']))==r['pptx_sha256']
            assert digest(Path(r['render_path']))==r['render_sha256']
    out=BASE/(args.output_name or f'{cohort}-review'); out.mkdir(parents=True,exist_ok=False)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
    audit=[]
    for sample in samples:
        name=sample.get('id',sample.get('pool_id')); assert digest(ROOT/sample['path'])==sample['sha256']
        sheet=Image.new('RGB',(1920,720),'#eef2f6'); draw=ImageDraw.Draw(sheet)
        for n,(label,source,record) in enumerate([
            ('Original source',ROOT/sample['path'],None),('Approved v0.3',prior/'renders'/f'{name}.png',priors[name]),
            ('R6 frozen candidate',current/'renders'/f'{name}.png',rows[name])]):
            draw.text((n*640+10,12),label,font=font,fill='black')
            if record is not None and record['returncode']!=0:
                draw.text((n*640+20,300),'Resource rejection / no PPT',font=font,fill='black'); continue
            with Image.open(source) as image:
                image=image.convert('RGB')
                if record is None: image=image.crop(pixel_box(sample['roi'],image.size))
                thumb=ImageOps.contain(image,(620,650)); sheet.paste(thumb,(n*640+(640-thumb.width)//2,50+(650-thumb.height)//2))
        path=out/f'{name}.png'; sheet.save(path)
        r=rows[name]; check=dict(id=name,sheet_sha256=digest(path),counts=r.get('report',{}).get('counts'))
        if r['returncode']==0:
            pp=current/f'{name}.pptx'; assert digest(pp)==r['pptx_sha256']
            deck=Presentation(pp); objects={s.name:s for s in deck.slides[0].shapes}; ids={str(s.shape_id):s.name for s in objects.values()}
            scene=r['report']['emitted_scene']
            check.update(ids_exact=set(objects)=={s['id'] for s in scene['shapes']}|{c['id'] for c in scene['connectors']},
                raster_count=len(deck.slides[0]._element.xpath('.//p:pic')),
                owner_bindings=all([ids[b.get('id')] for b in objects[c['id']]._element.xpath('.//a:stCxn | .//a:endCxn')]==[c['source_owner_id'],c['target_owner_id']] for c in scene['connectors']))
            assert check['ids_exact'] and check['owner_bindings'] and check['raster_count']==0
        audit.append(check)
    with (out/'audit.json').open('x',encoding='utf-8',newline='\n') as f: json.dump(audit,f,indent=2)
    print(audit)

if __name__=='__main__': main()

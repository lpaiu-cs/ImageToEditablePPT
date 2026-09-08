"""Bind production artifacts, native edit audits, historical failures and visual sheets."""
import argparse
import json
from pathlib import Path
from PIL import Image
from pptx import Presentation
from image_to_editable_ppt.research.native_emit import load_scene
from image_to_editable_ppt.region.geometry import normalize_scene
from image_to_editable_ppt.research.region_first import digest
from review_preset_campaign import audit,sheet


def read(path): return json.loads(path.read_text(encoding='utf-8-sig'))


def main():
    p=argparse.ArgumentParser(); p.add_argument('--output-dir',required=True); p.add_argument('--visual-review'); p.add_argument('--evidence-output'); args=p.parse_args()
    root=Path(__file__).resolve().parents[1]; base=root/'artifacts/region-native-production'
    out=root/args.output_dir; out.mkdir(parents=True,exist_ok=False)
    results=read(base/'run4/results.json'); probe=read(base/'edit-probe-v2/probe.json')
    inputs={r['sample']:r for r in read(root/'benchmarks/region_first_r3_inputs_v1.json')['inputs']}
    initial=[]; edits=[]; sheets=[]
    for r in results['records']:
        if r.get('status')=='cli_rejected': continue
        sample=r['sample']; report=read(root/r['report_path'])
        assert digest(root/r['pptx_path'])==r['pptx_sha256']==report['output_sha256']
        assert digest(root/r['scene_path'])==r['scene_sha256']
        assert digest(root/r['report_path'])==r['report_sha256']
        assert digest(root/inputs[sample]['source_path'])==r['source_sha256']==report['source_sha256']
        scene,_=normalize_scene(load_scene(root/r['scene_path']))
        check=audit(root/r['pptx_path'],scene,report['emission']['fits']); check['sample']=sample
        assert check['native_ids_exact'] and check['labels_preserved'] and check['raster_count']==0
        assert check['max_endpoint_error_pt']<1 and check['max_box_error_pt']<1
        assert all(c['owners_preserved'] and c['fit_route_error_pt']<1 for c in check['connectors'])
        deck=Presentation(root/r['pptx_path'])
        assert all(s._element.xpath('./p:spPr/a:effectLst') for s in deck.slides[0].shapes)
        initial.append(check)
        item=inputs[sample]
        sheets.append(sheet([root/item['source_path'],root/f'artifacts/region-first-r3e1/run1-renders/{sample}--preset_path.png',base/f'run4-renders/{sample}--region_native.png'],
            ['Source ROI','R3 native preset','Production region v0.2'],item['roi'],out/f'{sample}-initial.png'))
        rec=next((e for e in probe['records'] if e['sample']==sample),None)
        if rec:
            assert rec['input_sha256']==r['pptx_sha256'] and rec['original_unchanged'] and rec['text_edit_succeeded']
            stages=[]
            for stage in ('roundtrip','edited','reopened','rerouted'):
                a=audit(base/f'edit-probe-v2/{sample}--{stage}.pptx',scene,report['emission']['fits'],None if stage=='roundtrip' else rec['edited_owner'])
                a['stage']=stage
                a['sites_preserved']=all(x['site_indices']==y['site_indices'] for x,y in zip(check['connectors'],a['connectors']))
                a['arrow_ends_preserved']=all(x['arrow_ends']==y['arrow_ends'] for x,y in zip(check['connectors'],a['connectors']))
                assert a['native_ids_exact'] and a['labels_preserved'] and a['raster_count']==0 and a['max_box_error_pt']<1
                assert all(c['owners_preserved'] for c in a['connectors']) and a['arrow_ends_preserved']
                if stage!='rerouted': assert a['sites_preserved'] and a['max_endpoint_error_pt']<1
                if stage=='roundtrip': assert all(c['fit_route_error_pt']<1 for c in a['connectors'])
                stages.append(a)
            edits.append(dict(sample=sample,stages=stages))
            # Include movement margin in source-normalized scope.
            roi=list(item['roi']); roi[2]=min(1,roi[2]+.15); roi[3]=min(1,roi[3]+.15)
            sheets.append(sheet([base/f'edit-probe-v2/{sample}--{s}.png' for s in ('roundtrip','reopened','rerouted')],
                ['Native roundtrip','Move / resize / reopen','Explicit reroute stress'],roi,out/f'{sample}-edit.png'))
    assert len(initial)==5 and len(edits)==2
    histories={}
    for run in ('run1','run2','run3','run4','synthetic-v1','synthetic-v2','synthetic-v3'):
        data=read(base/run/'results.json')
        records=[]
        for r in data['records']:
            record={k:v for k,v in r.items() if k not in ('report','command')}
            if 'report' in r:
                record['diagnostics']={k:r['report'].get(k) for k in ('counts','identity','text','emission')}
            records.append(record)
        histories[run]=dict(git_commit=data['git_commit'],code_sha256=data.get('code_sha256'),records=records)
    assert all(r['passed'] for r in histories['synthetic-v3']['records'])
    synthetic_panels=[]
    for r in histories['synthetic-v3']['records']:
        name=r['sample']; path=out/f'{name}-synthetic.png'
        sheets.append(sheet([base/f'synthetic-v3/{name}.png',base/f'synthetic-v3-renders/{name}.png'],
            [name+' source','Native: '+str(r['counts']['nodes'])+' nodes / '+str(r['counts']['connectors'])+' links'],[0,0,1,1],path))
        with Image.open(path) as im: synthetic_panels.append(im.copy())
    # Compact contact sheet; individual sheets remain available for full-size inspection.
    thumbs=[]
    for panel in synthetic_panels:
        panel.thumbnail((720,320)); thumbs.append(panel)
    contact=Image.new('RGB',(1440,320*((len(thumbs)+1)//2)),'white')
    for i,panel in enumerate(thumbs): contact.paste(panel,((i%2)*720,(i//2)*320))
    contact.save(out/'synthetic-contact.png')
    manifests={}
    for folder in ('run4-renders','synthetic-v3-renders'):
        manifest=read(base/folder/'powerpoint_render_manifest.json')
        for r in manifest['records']:
            assert digest(Path(r['pptx_path']))==r['pptx_sha256'] and digest(Path(r['render_path']))==r['render_sha256']
        manifests[folder]=manifest
    paths=[base/run/'results.json' for run in histories]
    paths+=list((base/'edit-probe-v2').glob('*'))
    paths+=[base/folder/'powerpoint_render_manifest.json' for folder in manifests]
    old=root/'artifacts/region-first-r4/run1/results.json'
    evidence=dict(status='production_integrated_explicit_roi_not_general_quality_acceptance',
        histories=histories,r4=read(old),initial_audits=initial,edit_audits=edits,native_probe=probe,
        renders=manifests,sheets=sheets,provenance={str(p.relative_to(root)):digest(p) for p in [*paths,old]},
        auditor_sha256=digest(Path(__file__)),shared_auditor_sha256=digest(root/'tools/review_preset_campaign.py'),
        probe_sha256=digest(root/'tools/probe_region_runtime.ps1'))
    if args.visual_review: evidence['visual_review']=read(root/args.visual_review)
    (out/'evidence.json').write_text(json.dumps(evidence,indent=2),encoding='utf-8')
    if args.evidence_output:
        with (root/args.evidence_output).open('x',encoding='utf-8') as f: json.dump(evidence,f,indent=2)
    print(json.dumps(dict(initial_max_endpoint_pt=max(a['max_endpoint_error_pt'] for a in initial),
        edit=[dict(sample=e['sample'],stages=[dict(stage=s['stage'],error=s['max_endpoint_error_pt'],sites=s['sites_preserved']) for s in e['stages']]) for e in edits]),indent=2))


if __name__=='__main__': main()
